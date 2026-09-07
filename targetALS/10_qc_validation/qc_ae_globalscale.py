"""
Does the per-image min-max normalisation hide the PBS/FBS contrast?

The deployed pipeline normalises EVERY EEM to [0,1] individually before the
encoder sees it (eval_5fold_metrics.load_excel), which destroys absolute
brightness -- the most obvious physical difference between a protein-free
buffer well and a serum well. Every null result so far could therefore mean
either "the model cannot represent the corona" or "the pipeline discarded the
contrast before the model got it". This separates those.

Three questions, in order:
  1. IS the discarded information discriminative? Test PBS vs FBS directly on
     raw, un-normalised image summaries (total / mean / max intensity).
     If this is null too, preprocessing was never the explanation.
  2. Does encoding with a SHARED GLOBAL scale (relative brightness preserved)
     recover a separation the per-image scaling loses?
  3. Head-to-head against the per-image baseline, same wells, same statistics.

CAVEAT carried throughout: the encoder was TRAINED on per-image normalised
inputs, so globally-scaled inputs are out-of-distribution for it -- a null in
(2) is therefore weaker evidence than a positive would be.

Inference only. Outputs: qc_ae_globalscale.csv, FigSX_qc_ae_globalscale.png/.pdf
"""
import os, re, glob, importlib.util, itertools
import numpy as np, pandas as pd
from scipy import stats
from sklearn.decomposition import PCA
from sklearn.neighbors import KNeighborsClassifier
from sklearn.model_selection import LeaveOneOut, cross_val_predict
import torch
from openpyxl import load_workbook
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt

ROOT = "/mnt/c/Users/riccardo-s/Documents/CNT/targetALS/exp5"
HERE = os.path.dirname(os.path.abspath(__file__))
TC   = os.path.join(ROOT, "tech_controls")
MODEL_DIR = os.path.join(ROOT, "5_fold_models_original")
DEVICE, LATENT_DIM, SEED = torch.device("cpu"), 512, 42
BLUE, VERM, GREY, INK, MUT = "#0072B2", "#D55E00", "#9aa3ad", "#222", "#8a8a8a"
torch.manual_seed(SEED); np.random.seed(SEED); torch.set_num_threads(2)

spec = importlib.util.spec_from_file_location(
    "cae", os.path.join(ROOT, "conv_autoencoder_detailed.py"))
cae = importlib.util.module_from_spec(spec); spec.loader.exec_module(cae)
Model = cae.ConvAutoencoderWithAttention

def load_raw(fp):
    """un-normalised EEM, same row/col stripping as the deployed loader."""
    ws = load_workbook(fp, data_only=True).active
    data = []
    for ri, row in enumerate(ws.iter_rows(values_only=True)):
        if ri == 0:
            continue
        rd = [float(c) if isinstance(c, (int, float)) and c is not None else 0.0
              for ci, c in enumerate(row) if ci != 0]
        if rd:
            data.append(rd)
    return np.array(data, dtype=np.float32)

def control_map():
    m = {}
    for tp, pats in [("0h", [os.path.join(TC, "*.xlsx"), os.path.join(TC, "0h", "*.xlsx")]),
                     ("6h", [os.path.join(TC, "6h", "*.xlsx")]),
                     ("24h", [os.path.join(TC, "24h", "*.xlsx")])]:
        for pat in pats:
            for fp in glob.glob(pat):
                mm = re.match(r"^([FP]\d)", os.path.basename(fp))
                if mm:
                    m.setdefault(mm.group(1), {})[tp] = fp
    return {w: d for w, d in sorted(m.items()) if len(d) == 3}

CM = control_map(); WELLS = sorted(CM, key=lambda w: (w[0], w))
GRP = np.array(["PBS" if w[0] == "P" else "FBS" for w in WELLS])
TPS = ["0h", "6h", "24h"]
print(f"controls: {WELLS}  ({(GRP=='PBS').sum()} PBS / {(GRP=='FBS').sum()} FBS)\n")

RAWC = os.path.join(HERE, "raw_controls.npy")
if os.path.exists(RAWC):
    RAW = np.load(RAWC)
else:
    RAW = np.stack([np.stack([load_raw(CM[w][tp]) for tp in TPS]) for w in WELLS])
    np.save(RAWC, RAW)
print(f"raw control EEMs {RAW.shape}  (wells, timepoints, H, W)")

# ---------------------------------------------------------------- statistics
COMBOS = list(itertools.combinations(range(7), 3))
def perm_p_1d(v, g=GRP):
    obs = abs(v[g == "FBS"].mean() - v[g == "PBS"].mean()); null = []
    for c in COMBOS:
        m = np.zeros(7, bool); m[list(c)] = True
        null.append(abs(v[m].mean() - v[~m].mean()))
    return (np.array(null) >= obs).mean()
def auc(v, g=GRP):
    return stats.mannwhitneyu(v[g == "FBS"], v[g == "PBS"],
                              alternative="two-sided").statistic/12.0
def loo_knn(Z, g=GRP, k=3):
    pred = cross_val_predict(KNeighborsClassifier(n_neighbors=k), Z, g, cv=LeaveOneOut())
    return (pred == g).mean()
def perm_knn(Z, g=GRP, k=3):
    obs = loo_knn(Z, g, k); null = []
    for c in COMBOS:
        gg = np.array(["PBS"]*7, dtype=object); gg[list(c)] = "FBS"
        null.append(loo_knn(Z, gg, k))
    return obs, (np.array(null) >= obs).mean()

# ================================================================ Q1: raw signal
print("="*78)
print("Q1  Is the information the normalisation discards actually discriminative?")
print("="*78)
q1 = []
for name, fn in [("mean intensity", lambda a: a.mean()),
                 ("total intensity", lambda a: a.sum()),
                 ("max intensity", lambda a: a.max()),
                 ("dynamic range (max-min)", lambda a: a.max()-a.min())]:
    for ti, tp in enumerate(TPS):
        v = np.array([fn(RAW[i, ti]) for i in range(7)])
        q1.append(dict(quantity=name, tp=tp, auc=auc(v), perm_p=perm_p_1d(v),
                       PBS=v[GRP == "PBS"].mean(), FBS=v[GRP == "FBS"].mean()))
Q1 = pd.DataFrame(q1)
for name in Q1.quantity.unique():
    s = Q1[Q1.quantity == name]
    print(f"  {name:24s} AUC " + "  ".join(f"{t}={a:.2f}(p={p:.3f})"
          for t, a, p in zip(s.tp, s.auc, s.perm_p)))
best_raw = Q1.loc[Q1.perm_p.idxmin()]
print(f"\n  best raw summary: {best_raw.quantity} @ {best_raw.tp}  "
      f"AUC={best_raw.auc:.2f}  p={best_raw.perm_p:.3f}   [floor 0.029]")

# ================================================================ encodings
GMAX = float(RAW.max())
def norm_perimage(a):
    mn, mx = a.min(), a.max()
    return (a-mn)/(mx-mn) if mx > mn else np.zeros_like(a)
def norm_global(a):
    return a/GMAX                                   # relative brightness preserved

def encode(X):
    T = torch.from_numpy(X.astype(np.float32)).unsqueeze(2)
    out = []
    for k in range(1, 6):
        model = Model(in_channels=1, latent_dim=LATENT_DIM, num_classes=2).to(DEVICE)
        with torch.no_grad():
            _ = model(T[:1])
        model.load_state_dict(torch.load(os.path.join(MODEL_DIR, f"best_model_fold{k}.pth"),
                                         map_location=DEVICE), strict=True)
        model.eval()
        with torch.no_grad():
            zs = torch.stack([model.encoder(T[:, t]) for t in range(3)], dim=1)
            za, _ = model.attention(zs)
            p = torch.softmax(model.classifier(za), 1)[:, 1]
        out.append((za.cpu().numpy(), p.cpu().numpy()))
    return out

X_per = np.stack([[norm_perimage(RAW[i, t]) for t in range(3)] for i in range(7)])
X_glo = np.stack([[norm_global(RAW[i, t])   for t in range(3)] for i in range(7)])
print(f"\ninput ranges -- per-image: every image spans [0,1]; "
      f"global: [{X_glo.min():.3f}, {X_glo.max():.3f}], "
      f"per-image mean of maxima = {np.mean([X_glo[i,t].max() for i in range(7) for t in range(3)]):.3f}")
print("  (the global variant has lower dynamic range than anything the encoder saw in training)")

print("\n" + "="*78)
print("Q2/Q3  PBS vs FBS in z_agg, per-image vs global scaling")
print("="*78)
rows = []
for tag, X in [("per-image (deployed)", X_per), ("global scale", X_glo)]:
    res = encode(X)
    for k, (Z, P) in enumerate(res, start=1):
        pca = PCA(n_components=6, random_state=SEED); S = pca.fit_transform(Z)
        a_k, p_k = perm_knn(Z)
        rows.append(dict(scaling=tag, fold=k,
                         knn=a_k, knn_p=p_k,
                         pc1_auc=auc(S[:, 0]), pc1_p=perm_p_1d(S[:, 0]),
                         pals_auc=auc(P), pals_p=perm_p_1d(P),
                         P_PBS=P[GRP == "PBS"].mean(), P_FBS=P[GRP == "FBS"].mean()))
R = pd.DataFrame(rows)
R.to_csv(os.path.join(HERE, "qc_ae_globalscale.csv"), index=False)
Q1.to_csv(os.path.join(HERE, "qc_ae_globalscale_rawsummaries.csv"), index=False)

for tag in R.scaling.unique():
    s = R[R.scaling == tag]
    print(f"\n  {tag}")
    print(f"    LOO 3-NN acc {s.knn.mean():.2f} (chance 0.57)  perm p {s.knn_p.min():.3f}-{s.knn_p.max():.3f}")
    print(f"    PC1 AUC      {s.pc1_auc.mean():.2f}            perm p {s.pc1_p.min():.3f}-{s.pc1_p.max():.3f}")
    print(f"    P_ALS  PBS={s.P_PBS.mean():.3f} FBS={s.P_FBS.mean():.3f}  "
          f"AUC {s.pals_auc.mean():.2f}  perm p {s.pals_p.min():.3f}-{s.pals_p.max():.3f}")

# ---------------------------------------------------------------- figure
plt.rcParams.update({"font.size": 9, "axes.titlesize": 9.5, "axes.labelsize": 9,
                     "axes.spines.top": False, "axes.spines.right": False, "figure.dpi": 150})
fig, axs = plt.subplots(1, 4, figsize=(14.6, 3.5))
fig.subplots_adjust(left=0.055, right=0.995, bottom=0.19, top=0.82, wspace=0.4)

a = axs[0]
v = np.array([RAW[i, 2].mean() for i in range(7)])
for g, col in [("PBS", BLUE), ("FBS", VERM)]:
    m = GRP == g
    a.scatter(np.where(m)[0], v[m], s=55, c=col, label=g, zorder=3)
for i, w in enumerate(WELLS):
    a.annotate(w, (i, v[i]), fontsize=6, color=MUT, xytext=(4, 3), textcoords="offset points")
s24 = Q1[(Q1.quantity == "mean intensity") & (Q1.tp == "24h")].iloc[0]
a.set_xticks(range(7)); a.set_xticklabels(WELLS, fontsize=7)
a.set_ylabel("raw mean intensity @24h"); a.set_xlabel("control well")
a.set_title(f"(a) Raw brightness, un-normalised\n AUC={s24.auc:.2f} p={s24.perm_p:.3f}", loc="left")
a.legend(frameon=False, fontsize=7.5)

b = axs[1]
w = 0.36
for i, (tag, col) in enumerate([("per-image (deployed)", GREY), ("global scale", BLUE)]):
    s = R[R.scaling == tag]
    b.bar(np.arange(5)+(i-0.5)*w, s.knn, w, color=col, alpha=.9, label=tag)
b.axhline(0.57, color=INK, ls="--", lw=1)
b.text(-0.45, 0.59, "chance 0.57", fontsize=6.5, color=INK)
b.set_xticks(range(5)); b.set_xticklabels([f"f{k}" for k in range(1, 6)])
b.set_xlabel("fold model"); b.set_ylabel("LOO 3-NN accuracy"); b.set_ylim(0, 1.05)
b.set_title("(b) Separation in z_agg", loc="left")
b.legend(frameon=False, fontsize=6.6)

c = axs[2]
for i, (tag, col) in enumerate([("per-image (deployed)", GREY), ("global scale", BLUE)]):
    s = R[R.scaling == tag]
    c.bar(np.arange(5)+(i-0.5)*w, s.knn_p, w, color=col, alpha=.9, label=tag)
c.axhline(0.05, color=VERM, ls="--", lw=1)
c.axhline(1/35, color=MUT, ls=":", lw=1)
c.text(-0.45, 1/35*1.6, "floor 0.029", fontsize=6.5, color=MUT)
c.set_xticks(range(5)); c.set_xticklabels([f"f{k}" for k in range(1, 6)])
c.set_xlabel("fold model"); c.set_ylabel("exact permutation p"); c.set_ylim(0, 1.05)
c.set_title("(c) Significance", loc="left")

d = axs[3]
for i, (tag, col) in enumerate([("per-image (deployed)", GREY), ("global scale", BLUE)]):
    s = R[R.scaling == tag]
    d.scatter(np.full(5, i)-0.08, s.P_PBS, s=30, c=BLUE, marker="o", zorder=3,
              label="PBS" if i == 0 else None)
    d.scatter(np.full(5, i)+0.08, s.P_FBS, s=30, c=VERM, marker="^", zorder=3,
              label="FBS" if i == 0 else None)
d.axhline(0.5, color=INK, ls="--", lw=1)
d.set_xticks([0, 1]); d.set_xticklabels(["per-image", "global"], fontsize=8)
d.set_ylabel("classifier P(ALS)"); d.set_ylim(-0.05, 1.05)
d.set_title("(d) Model negative control\n under both scalings", loc="left")
d.legend(frameon=False, fontsize=7.5)

fig.savefig(os.path.join(HERE, "FigSX_qc_ae_globalscale.png"), dpi=300, bbox_inches="tight")
fig.savefig(os.path.join(HERE, "FigSX_qc_ae_globalscale.pdf"), bbox_inches="tight")
print("\nSaved FigSX_qc_ae_globalscale.png/.pdf + qc_ae_globalscale{,_rawsummaries}.csv")
