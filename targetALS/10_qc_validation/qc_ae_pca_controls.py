"""
PCA of the conv-AE latent space computed on the TECHNICAL CONTROLS ONLY
(4 PBS + 3 FBS wells), per saved fold model.

Why controls-only PCA is the right form of this question: PCA is unsupervised,
so its axes are not fitted to the PBS/FBS labels. If a principal axis of the
control-only latent variation lines up with protein-present vs protein-absent,
that is real structure rather than a direction fitted to the answer -- unlike
the difference-of-means axis tested earlier.

Caveat that governs the whole script: n=7 in 512-d. After centering the data
matrix has rank <= 6, so the first few PCs explain ~all variance by
construction; explained variance is NOT evidence. Only the label alignment of
each PC, tested by exact permutation over all C(7,3)=35 relabellings, carries
information -- with a max-over-PC statistic to cover testing several axes.

Also checks what PC1 actually tracks, since "not PBS/FBS" does not mean
"nothing": correlation against the peak-space fingerprint deviation and against
the known first-read (P1) artifact.

Inference only -- no retraining; controls were never in any training split.
Outputs: qc_ae_pca_controls.csv, qc_ae_pca_scores.csv, z_agg_controls.npy,
         FigSX_qc_ae_pca.png/.pdf
"""
import os, re, glob, importlib.util, itertools
import numpy as np, pandas as pd
from scipy import stats
from sklearn.decomposition import PCA
import torch
from openpyxl import load_workbook
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt

ROOT = "/mnt/c/Users/riccardo-s/Documents/CNT/targetALS/exp5"
HERE = os.path.dirname(os.path.abspath(__file__))
TC   = os.path.join(ROOT, "tech_controls")
MODEL_DIR = os.path.join(ROOT, "5_fold_models_original")
DEVICE, LATENT_DIM, SEED = torch.device("cpu"), 512, 42
BLUE, VERM, GREY, INK, MUT = "#0072B2", "#D55E00", "#9aa3ad", "#222", "#8a8a8a"
torch.manual_seed(SEED); np.random.seed(SEED); torch.set_num_threads(2)

spec = importlib.util.spec_from_file_location(
    "cae", os.path.join(ROOT, "conv_autoencoder_detailed.py"))
cae = importlib.util.module_from_spec(spec); spec.loader.exec_module(cae)
Model = cae.ConvAutoencoderWithAttention

def load_excel(fp):
    ws = load_workbook(fp, data_only=True).active
    data = []
    for ri, row in enumerate(ws.iter_rows(values_only=True)):
        if ri == 0:
            continue
        rd = [float(c) if isinstance(c, (int, float)) and c is not None else 0.0
              for ci, c in enumerate(row) if ci != 0]
        if rd:
            data.append(rd)
    a = np.array(data, dtype=np.float32); mn, mx = a.min(), a.max()
    return (a-mn)/(mx-mn) if mx > mn else np.zeros_like(a)

def control_map():
    m = {}
    for tp, pats in [("0h", [os.path.join(TC, "*.xlsx"), os.path.join(TC, "0h", "*.xlsx")]),
                     ("6h", [os.path.join(TC, "6h", "*.xlsx")]),
                     ("24h", [os.path.join(TC, "24h", "*.xlsx")])]:
        for pat in pats:
            for fp in glob.glob(pat):
                mm = re.match(r"^([FP]\d)", os.path.basename(fp))
                if mm:
                    m.setdefault(mm.group(1), {})[tp] = fp
    return {w: d for w, d in sorted(m.items()) if len(d) == 3}

CM = control_map(); WELLS = sorted(CM, key=lambda w: (w[0], w))
GRP = np.array(["PBS" if w[0] == "P" else "FBS" for w in WELLS])
print(f"controls: {WELLS}  ({(GRP=='PBS').sum()} PBS / {(GRP=='FBS').sum()} FBS)")

CACHE = os.path.join(HERE, "z_agg_controls.npy")
if os.path.exists(CACHE):
    Z = np.load(CACHE); print(f"loaded cached embeddings {Z.shape}")
else:
    print("loading control EEMs...")
    Xc = np.stack([np.stack([load_excel(CM[w][tp]) for tp in ["0h", "6h", "24h"]])
                   for w in WELLS]).astype(np.float32)
    Tc = torch.from_numpy(Xc).unsqueeze(2)
    Z = []
    for k in range(1, 6):
        model = Model(in_channels=1, latent_dim=LATENT_DIM, num_classes=2).to(DEVICE)
        with torch.no_grad():
            _ = model(Tc[:1])
        model.load_state_dict(torch.load(os.path.join(MODEL_DIR, f"best_model_fold{k}.pth"),
                                         map_location=DEVICE), strict=True)
        model.eval()
        with torch.no_grad():
            zs = torch.stack([model.encoder(Tc[:, t]) for t in range(3)], dim=1)
            za, _ = model.attention(zs)
        Z.append(za.cpu().numpy())
    Z = np.stack(Z)                       # (5 folds, 7 wells, 512)
    np.save(CACHE, Z); print(f"cached embeddings {Z.shape} -> z_agg_controls.npy")

# ---------------------------------------------------------------- statistics
COMBOS = list(itertools.combinations(range(7), 3))       # 35 -- safe to materialise
def perm_p_1d(v, g):
    obs = abs(v[g == "FBS"].mean() - v[g == "PBS"].mean())
    null = []
    for c in COMBOS:
        m = np.zeros(7, bool); m[list(c)] = True
        null.append(abs(v[m].mean() - v[~m].mean()))
    return (np.array(null) >= obs).mean()

def auc(pos, neg):
    return stats.mannwhitneyu(pos, neg, alternative="two-sided").statistic/(len(pos)*len(neg))

def maxpc_perm_p(S, g, npc):
    """family-wise p for 'ANY of the first npc PCs separates', via max statistic."""
    def stat(gg):
        return max(abs(S[gg == "FBS", j].mean()-S[gg == "PBS", j].mean())/(S[:, j].std()+1e-12)
                   for j in range(npc))
    obs = stat(g); null = []
    for c in COMBOS:
        gg = np.array(["PBS"]*7, dtype=object); gg[list(c)] = "FBS"
        null.append(stat(gg))
    return (np.array(null) >= obs).mean()

# ---------------------------------------------------------------- PCA per fold
NPC = 4
rows, scores = [], []
for k in range(1, 6):
    Zk = Z[k-1]
    p = PCA(n_components=min(6, Zk.shape[0]-1), random_state=SEED)
    S = p.fit_transform(Zk)
    evr = p.explained_variance_ratio_
    for j in range(NPC):
        v = S[:, j]
        rows.append(dict(fold=k, PC=j+1, evr=evr[j],
                         auc=auc(v[GRP == "FBS"], v[GRP == "PBS"]),
                         perm_p=perm_p_1d(v, GRP),
                         mean_FBS=v[GRP == "FBS"].mean(), mean_PBS=v[GRP == "PBS"].mean()))
    fw = maxpc_perm_p(S, GRP, NPC)
    for i, w in enumerate(WELLS):
        scores.append(dict(fold=k, well=w, group=GRP[i], PC1=S[i, 0], PC2=S[i, 1], PC3=S[i, 2]))
    best = min(rows[-NPC:], key=lambda r: r["perm_p"])
    print(f"fold{k}: PC1 evr={evr[0]:.2f}  best PC{best['PC']} "
          f"AUC={best['auc']:.2f} p={best['perm_p']:.3f}   family-wise p(any of {NPC} PCs)={fw:.3f}")
    rows[-1]["fw_p"] = fw

R = pd.DataFrame(rows); SC = pd.DataFrame(scores)
R.to_csv(os.path.join(HERE, "qc_ae_pca_controls.csv"), index=False)
SC.to_csv(os.path.join(HERE, "qc_ae_pca_scores.csv"), index=False)

print("\n" + "="*76)
print("PCA on the 7 control wells only (latent, per fold model)")
print("="*76)
for j in range(1, NPC+1):
    s = R[R.PC == j]
    print(f"  PC{j}: evr {s.evr.mean():.2f}   AUC(FBS vs PBS) {s.auc.mean():.2f} "
          f"(range {s.auc.min():.2f}-{s.auc.max():.2f})   exact perm p "
          f"{s.perm_p.min():.3f}-{s.perm_p.max():.3f}")
print(f"\n  permutation floor with 4 vs 3 wells = {1/len(COMBOS):.3f} "
      f"({len(COMBOS)} relabellings) -- only a PERFECT split can reach it")

# ---------------------------------------------------------------- what is PC1?
try:
    L = pd.read_csv(os.path.join(HERE, "controls_features_long.csv"), dtype={'chirality': str})
    CH = ['6.5','7.5','7.6','8.3','8.4','8.6','8.7','9.4','9.5','10.2','10.3','10.5']
    def fp_dev(w):
        a = L[(L['sample'] == w) & (L.tp == '0h')].set_index('chirality').reindex(CH)['gauss_max'].to_numpy(float)
        return a/np.nanmean(a)
    F = np.array([fp_dev(w) for w in WELLS]); ref = np.median(F, 0)
    dev = np.abs(F-ref).max(1)                       # peak-space fingerprint deviation
    firstread = np.array([1.0 if w in ("P1", "F1") else 0.0 for w in WELLS])
    print("\n  what does PC1 track, if not PBS/FBS?")
    for k in range(1, 6):
        v = SC[SC.fold == k].set_index('well').loc[WELLS, 'PC1'].to_numpy()
        r1, p1 = stats.pearsonr(v, dev); r2, p2 = stats.pearsonr(v, firstread)
        print(f"    fold{k}: r(PC1, peak fingerprint deviation)={r1:+.2f} (p={p1:.2f})   "
              f"r(PC1, first-read well)={r2:+.2f} (p={p2:.2f})")
except Exception as e:
    print("  PC1 provenance check skipped:", e)

# ---------------------------------------------------------------- figure
plt.rcParams.update({"font.size": 9, "axes.titlesize": 9.5, "axes.labelsize": 9,
                     "axes.spines.top": False, "axes.spines.right": False, "figure.dpi": 150})
fig, axs = plt.subplots(1, 6, figsize=(19.5, 3.4))
fig.subplots_adjust(left=0.032, right=0.995, bottom=0.18, top=0.80, wspace=0.34)
for i, k in enumerate(range(1, 6)):
    ax = axs[i]; s = SC[SC.fold == k]
    for g, col in [("PBS", BLUE), ("FBS", VERM)]:
        q = s[s.group == g]
        ax.scatter(q.PC1, q.PC2, s=70, c=col, edgecolor="white", lw=.7, label=g, zorder=3)
        for _, r in q.iterrows():
            ax.annotate(r.well, (r.PC1, r.PC2), fontsize=6, color=MUT,
                        xytext=(4, 3), textcoords="offset points")
    e = R[(R.fold == k)]
    ax.set_xlabel(f"PC1 ({e[e.PC==1].evr.iloc[0]*100:.0f}%)")
    ax.set_ylabel(f"PC2 ({e[e.PC==2].evr.iloc[0]*100:.0f}%)")
    ax.set_title(f"fold {k}  PC1 p={e[e.PC==1].perm_p.iloc[0]:.2f}", loc="left")
    if i == 0:
        ax.legend(frameon=False, fontsize=7.5)
ax = axs[5]
w = 0.15
for j in range(1, NPC+1):
    ax.bar(np.arange(5)+(j-2.5)*w, R[R.PC == j].perm_p, w, label=f"PC{j}", alpha=.9)
ax.axhline(0.05, color=INK, ls="--", lw=1)
ax.axhline(1/len(COMBOS), color=MUT, ls=":", lw=1)
ax.text(-0.45, 1/len(COMBOS)*1.5, "floor 0.029", fontsize=6.5, color=MUT)
ax.set_xticks(range(5)); ax.set_xticklabels([f"f{k}" for k in range(1, 6)])
ax.set_xlabel("fold model"); ax.set_ylabel("exact permutation p"); ax.set_ylim(0, 1.05)
ax.set_title("PBS vs FBS on each PC", loc="left")
ax.legend(frameon=False, fontsize=6.5, ncol=4, loc="lower center")
fig.suptitle("PCA of the conv-AE latent computed on the 7 technical control wells only "
             "(unsupervised axes; n=7 in 512-d, so explained variance is not evidence)",
             x=0.032, ha="left", fontsize=9.5, y=0.955)
fig.savefig(os.path.join(HERE, "FigSX_qc_ae_pca.png"), dpi=300, bbox_inches="tight")
fig.savefig(os.path.join(HERE, "FigSX_qc_ae_pca.pdf"), bbox_inches="tight")
print("\nSaved FigSX_qc_ae_pca.png/.pdf + qc_ae_pca_{controls,scores}.csv, z_agg_controls.npy")
