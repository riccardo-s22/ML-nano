"""
UMAP of the conv-AE latent space (z_agg), with the technical controls embedded
alongside the 39 patients, and a quantified separability check.

Two things are deliberately kept apart:
  * the UMAP is for LOOKING -- it is a nonlinear neighbour embedding and can
    manufacture apparent gaps, so nothing is concluded from it;
  * separability is QUANTIFIED in the original 512-d latent space (LOO k-NN,
    silhouette, exact permutation), where the numbers mean what they say.
Both are reported for two contrasts: PBS vs FBS (the QC question) and
ALS vs CTRL (the sanity check that the space is structured at all).

Reuses the inference recipe from qc_ae_controls.py: saved fold checkpoints,
no retraining, controls fully held out from every split.

Outputs: qc_ae_umap_coords.csv, qc_ae_umap_separability.csv,
         FigSX_qc_ae_umap.png/.pdf
"""
import os, re, glob, importlib.util, itertools, math, warnings
import numpy as np, pandas as pd
import torch
from openpyxl import load_workbook
from sklearn.neighbors import KNeighborsClassifier
from sklearn.model_selection import LeaveOneOut, cross_val_predict
from sklearn.metrics import silhouette_score
import umap
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
warnings.filterwarnings("ignore")

ROOT = "/mnt/c/Users/riccardo-s/Documents/CNT/targetALS/exp5"
HERE = os.path.dirname(os.path.abspath(__file__))
TC   = os.path.join(ROOT, "tech_controls")
MODEL_DIR = os.path.join(ROOT, "5_fold_models_original")
TP_DIRS = [os.path.join(ROOT, d) for d in ["out_0h", "out_6h", "out_24h"]]
DEVICE, LATENT_DIM = torch.device("cpu"), 512
BLUE, VERM, GREEN, GREY, INK, MUT = "#0072B2", "#D55E00", "#009E73", "#9aa3ad", "#222", "#8a8a8a"
SEED = 42
torch.manual_seed(SEED); np.random.seed(SEED)
torch.set_num_threads(2)

spec = importlib.util.spec_from_file_location(
    "cae", os.path.join(ROOT, "conv_autoencoder_detailed.py"))
cae = importlib.util.module_from_spec(spec); spec.loader.exec_module(cae)
Model = cae.ConvAutoencoderWithAttention


def load_excel(fp):
    ws = load_workbook(fp, data_only=True).active
    data = []
    for ri, row in enumerate(ws.iter_rows(values_only=True)):
        if ri == 0:
            continue
        rd = [float(c) if isinstance(c, (int, float)) and c is not None else 0.0
              for ci, c in enumerate(row) if ci != 0]
        if rd:
            data.append(rd)
    a = np.array(data, dtype=np.float32); mn, mx = a.min(), a.max()
    return (a-mn)/(mx-mn) if mx > mn else np.zeros_like(a)


def control_map():
    m = {}
    for tp, pats in [("0h", [os.path.join(TC, "*.xlsx"), os.path.join(TC, "0h", "*.xlsx")]),
                     ("6h", [os.path.join(TC, "6h", "*.xlsx")]),
                     ("24h", [os.path.join(TC, "24h", "*.xlsx")])]:
        for pat in pats:
            for fp in glob.glob(pat):
                mm = re.match(r"^([FP]\d)", os.path.basename(fp))
                if mm:
                    m.setdefault(mm.group(1), {})[tp] = fp
    return {w: d for w, d in sorted(m.items()) if len(d) == 3}

CM = control_map(); CWELLS = sorted(CM, key=lambda w: (w[0], w))
CGRP = np.array(["PBS" if w[0] == "P" else "FBS" for w in CWELLS])

lab = pd.read_csv(os.path.join(ROOT, "sample_labels.csv"))
codes = lab["code"].astype(str).tolist()
ypat = np.array(["ALS" if str(g).strip().lower() == "als" else "CTRL" for g in lab["group"]])

print("Loading EEMs...")
Xp = np.stack([np.stack([load_excel(os.path.join(d, f"{c}.xlsx")) for d in TP_DIRS])
               for c in codes]).astype(np.float32)
Xc = np.stack([np.stack([load_excel(CM[w][tp]) for tp in ["0h", "6h", "24h"]])
               for w in CWELLS]).astype(np.float32)
Tp = torch.from_numpy(Xp).unsqueeze(2); Tc = torch.from_numpy(Xc).unsqueeze(2)

def embed(model, X):
    with torch.no_grad():
        zs = torch.stack([model.encoder(X[:, t]) for t in range(X.shape[1])], dim=1)
        z, _ = model.attention(zs)
        return z.cpu().numpy(), torch.softmax(model.classifier(z), 1)[:, 1].cpu().numpy()

# ---------------------------------------------------------------- separability
def loo_knn(Z, lab_, k=3):
    """leave-one-out k-NN accuracy in the given space."""
    k = min(k, len(Z)-1)
    pred = cross_val_predict(KNeighborsClassifier(n_neighbors=k), Z, lab_, cv=LeaveOneOut())
    return (pred == lab_).mean()

def perm_knn(Z, lab_, k=3, exact_max=2000):
    """Permutation p for LOO k-NN accuracy.

    Enumerate exhaustively only when the number of relabellings is genuinely
    small; otherwise sample. The count is obtained with math.comb -- never by
    materialising the combinations, which for 39 samples is C(39,20) = 6.9e10
    tuples and will exhaust memory.
    """
    obs = loo_knn(Z, lab_, k)
    u = np.unique(lab_); n = len(lab_); n1 = int((lab_ == u[0]).sum())
    total = math.comb(n, n1)
    rng = np.random.RandomState(SEED)
    if total <= exact_max:
        combos = itertools.combinations(range(n), n1)   # lazy
        exact = True
    else:
        combos = (tuple(rng.choice(n, n1, replace=False)) for _ in range(exact_max))
        exact = False
    null = []
    for c in combos:
        ll = np.array([u[1]]*n, dtype=object); ll[list(c)] = u[0]
        null.append(loo_knn(Z, ll, k))
    null = np.array(null)
    return obs, (null >= obs).mean(), len(null), exact

# ---------------------------------------------------------------- run
rows, coords = [], []
for k in range(1, 6):
    model = Model(in_channels=1, latent_dim=LATENT_DIM, num_classes=2).to(DEVICE)
    with torch.no_grad():
        _ = model(Tc[:1])
    model.load_state_dict(torch.load(os.path.join(MODEL_DIR, f"best_model_fold{k}.pth"),
                                     map_location=DEVICE), strict=True)
    model.eval()
    Zc, Pc = embed(model, Tc); Zp, Pp = embed(model, Tp)
    Z = np.vstack([Zp, Zc])
    grp = np.concatenate([ypat, CGRP])
    kind = np.array(["patient"]*len(Zp) + ["control"]*len(Zc))
    name = np.array(codes + CWELLS)

    um = umap.UMAP(n_neighbors=10, min_dist=0.25, metric="euclidean",
                   random_state=SEED).fit_transform(Z)
    for i in range(len(Z)):
        coords.append(dict(fold=k, name=name[i], kind=kind[i], group=grp[i],
                           umap1=um[i, 0], umap2=um[i, 1],
                           P_ALS=(Pp[i] if i < len(Zp) else Pc[i-len(Zp)])))

    # ---- PBS vs FBS (the QC question): 512-d and UMAP-2d
    a_l, p_l, n_l, ex = perm_knn(Zc, CGRP, k=3)
    umc = um[len(Zp):]
    a_u, p_u, _, _ = perm_knn(umc, CGRP, k=3)
    sil_l = silhouette_score(Zc, CGRP) if len(set(CGRP)) > 1 else np.nan
    # ---- ALS vs CTRL (does the space have ANY structure?)
    a_p, p_p, n_p, exp_ = perm_knn(Zp, ypat, k=5, exact_max=1000)
    sil_p = silhouette_score(Zp, ypat)

    rows.append(dict(fold=k,
                     pbsfbs_knn_512d=a_l, pbsfbs_p_512d=p_l, pbsfbs_nperm=n_l,
                     pbsfbs_knn_umap=a_u, pbsfbs_p_umap=p_u, pbsfbs_silhouette=sil_l,
                     als_knn_512d=a_p, als_p_512d=p_p, als_silhouette=sil_p))
    print(f"fold{k}:  PBS/FBS  LOO-kNN 512d={a_l:.2f} (p={p_l:.3f})  UMAP={a_u:.2f} (p={p_u:.3f})"
          f"  sil={sil_l:+.3f}   |   ALS/CTRL 512d={a_p:.2f} (p={p_p:.3f}) sil={sil_p:+.3f}")

R = pd.DataFrame(rows); C = pd.DataFrame(coords)
R.to_csv(os.path.join(HERE, "qc_ae_umap_separability.csv"), index=False)
C.to_csv(os.path.join(HERE, "qc_ae_umap_coords.csv"), index=False)

print("\n" + "="*78)
print("SEPARABILITY  (quantified in the 512-d latent space, not in the UMAP)")
print("="*78)
print(f"  PBS vs FBS   LOO 3-NN accuracy = {R.pbsfbs_knn_512d.mean():.2f} "
      f"(chance {max((CGRP=='PBS').mean(),(CGRP=='FBS').mean()):.2f})   "
      f"exact perm p = {R.pbsfbs_p_512d.min():.3f}-{R.pbsfbs_p_512d.max():.3f}")
print(f"               silhouette = {R.pbsfbs_silhouette.mean():+.3f}  "
      f"(0 = no cluster structure)")
print(f"  ALS vs CTRL  LOO 5-NN accuracy = {R.als_knn_512d.mean():.2f} "
      f"(chance {max((ypat=='ALS').mean(),(ypat=='CTRL').mean()):.2f})   "
      f"perm p = {R.als_p_512d.min():.3f}-{R.als_p_512d.max():.3f}")
print(f"               silhouette = {R.als_silhouette.mean():+.3f}")
print("\n  Read the ALS/CTRL row first: it calibrates how much geometric structure the")
print("  raw latent space carries at all, independent of the classifier head.")

# ---------------------------------------------------------------- figure
plt.rcParams.update({"font.size": 9, "axes.titlesize": 9.5, "axes.labelsize": 9,
                     "axes.spines.top": False, "axes.spines.right": False, "figure.dpi": 150})
fig, axs = plt.subplots(1, 5, figsize=(17.5, 3.6))
fig.subplots_adjust(left=0.035, right=0.995, bottom=0.16, top=0.82, wspace=0.3)
for i, k in enumerate(range(1, 6)):
    ax = axs[i]; s = C[C.fold == k]
    for g, col, mk in [("CTRL", "#B8C4D0", 'o'), ("ALS", "#7F8C9B", 'o')]:
        q = s[(s.kind == "patient") & (s.group == g)]
        ax.scatter(q.umap1, q.umap2, s=26, c=col, marker=mk, lw=0, alpha=.9,
                   label=f"patient {g}", zorder=2)
    for g, col in [("PBS", BLUE), ("FBS", VERM)]:
        q = s[(s.kind == "control") & (s.group == g)]
        ax.scatter(q.umap1, q.umap2, s=95, c=col, marker='*', edgecolor='white',
                   linewidth=.7, label=g, zorder=4)
    ax.set_title(f"fold {k}   PBS/FBS p={R.loc[R.fold==k,'pbsfbs_p_512d'].iloc[0]:.2f}",
                 loc="left")
    ax.set_xticks([]); ax.set_yticks([])
    ax.set_xlabel("UMAP-1");
    if i == 0:
        ax.set_ylabel("UMAP-2"); ax.legend(frameon=False, fontsize=6.6, loc="best")
fig.suptitle("Conv-AE latent space (z_agg, 512-d) — technical controls embedded with the 39 patients."
             "  UMAP is for looking; separability is quantified in the 512-d space.",
             x=0.035, ha="left", fontsize=9.5, y=0.965)
fig.savefig(os.path.join(HERE, "FigSX_qc_ae_umap.png"), dpi=300, bbox_inches="tight")
fig.savefig(os.path.join(HERE, "FigSX_qc_ae_umap.pdf"), bbox_inches="tight")
print("\nSaved FigSX_qc_ae_umap.png/.pdf + qc_ae_umap_{coords,separability}.csv")
