#!/usr/bin/env python3
"""
FUSED PIPELINE: Optimal Percentile Search (Targeted ROIs + Proxy Classification)
================================================================================

This script fuses Step 3 and Step 4:
1. Loads Autoencoder models and computes Gradient Saliency + Linear Probing ONCE.
2. Iterates over a list of mask percentiles (e.g., 75, 80, 85, 90, 92, 95, 97).
3. For each percentile:
    a) Extracts the ROI masks (applying MSE filter).
    b) Reconstructs the Proxy Features via Two-Stage Ridge Regression (LOOCV).
    c) Trains classifiers (SVM, RF, LR) on the proxies.
    d) Calculates statistical power.
4. Outputs an optimization history CSV.

USAGE:
------
python fused_pipeline_opt_percentile.py \
    --tp_dirs ./out_0h,./out_6h,./out_24h \
    --labels_csv sample_labels.csv \
    --step1_dir ./step1_latent_features \
    --step2_dir ./step2_mse_maps \
    --model_dir ./5_fold_models_original \
    --arch_file conv_autoencoder_detailed.py \
    --output_dir ./fused_optimization_results \
    --percentiles "75, 80, 85, 90, 92, 95, 97"
"""

import os
import sys
import argparse
import importlib.util
import inspect
from pathlib import Path
from typing import List, Dict, Tuple, Optional

import numpy as np
import pandas as pd
from openpyxl import load_workbook
import torch
import torch.nn as nn
from sklearn.linear_model import Ridge, LogisticRegression
from sklearn.preprocessing import StandardScaler
from sklearn.model_selection import KFold, LeaveOneOut
from sklearn.svm import SVC
from sklearn.ensemble import RandomForestClassifier
from sklearn.metrics import accuracy_score, roc_auc_score, confusion_matrix
from scipy import ndimage
from scipy.stats import binom

import warnings
warnings.filterwarnings("ignore")

SEED = 42
torch.manual_seed(SEED)
np.random.seed(SEED)

TP_NAMES = ["0h", "6h", "24h"]

# =============================================================================
# CONFIG: The Top 5 Clinically Validated Features
# =============================================================================
TARGET_FEATURES = {
    "best_model_fold1": [103, 179],
    "best_model_fold3": [477, 376, 370]
}

# The explicit feature order we want in the final CSV
TARGET_FEATURES_LIST = [
    "best_model_fold3_dim477",
    "best_model_fold1_dim103",
    "best_model_fold3_dim376",
    "best_model_fold1_dim179",
    "best_model_fold3_dim370"
]

# =============================================================================
# Core Helpers
# =============================================================================
def calculate_power(acc: float, n_samples: int, null_acc: float = 0.5, alpha: float = 0.05) -> float:
    """Calculates Binomial Power for accuracy vs random chance."""
    if acc <= null_acc: return 0.0
    k_crit = 0
    for k in range(n_samples + 1):
        if 1 - binom.cdf(k-1, n_samples, null_acc) <= alpha:
            k_crit = k
            break
    real_power = 1 - binom.cdf(k_crit - 1, n_samples, acc)
    return real_power

def import_arch_module(arch_file: str):
    spec = importlib.util.spec_from_file_location("arch_module", arch_file)
    if spec is None or spec.loader is None: raise RuntimeError(f"Could not import architecture file: {arch_file}")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod

def _load_state_dict_flexible(model: nn.Module, path: str):
    state = torch.load(path, map_location="cpu")
    if "state_dict" in state: state = state["state_dict"]
    new_state = {}
    for k, v in state.items():
        if k.startswith("module."): k = k[7:]
        elif k.startswith("model."): k = k[6:]
        new_state[k] = v
    try:
        model.load_state_dict(new_state, strict=True)
    except Exception:
        model.load_state_dict(new_state, strict=False)

def load_labels(csv_path: str):
    df = pd.read_csv(csv_path)
    codes = df["code"].astype(str).tolist()
    groups = df["group"].values
    y = (groups == "ALS").astype(int)
    return codes, y, groups

def list_sample_codes(tp_dir: str) -> List[str]:
    codes = []
    for f in Path(tp_dir).glob("*.xlsx"):
        if f.name.startswith("~$"): continue
        name = f.stem
        for suffix in ["_0h", "_6h", "_24h"]:
            if name.endswith(suffix):
                name = name[:-len(suffix)]
                break
        codes.append(name)
    return codes

def find_file_for_code(tp_dir: str, code: str) -> Optional[str]:
    tp_path = Path(tp_dir)
    exact = tp_path / f"{code}.xlsx"
    if exact.exists(): return str(exact)
    for f in tp_path.glob("*.xlsx"):
        if f.name.startswith("~$"): continue
        stem = f.stem
        for suffix in ["_0h", "_6h", "_24h"]:
            if stem.endswith(suffix):
                if stem[:-len(suffix)] == code: return str(f)
                if stem[:-len(suffix)].replace('.', '_') == code: return str(f)
    return None

def load_excel_as_array(filepath: str) -> np.ndarray:
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    data = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx == 0: continue
        row_data = []
        for col_idx, cell in enumerate(row):
            if col_idx == 0: continue
            if isinstance(cell, (int, float)) and cell is not None: row_data.append(float(cell))
            else: row_data.append(0.0)
        if row_data: data.append(row_data)
    arr = np.array(data, dtype=np.float32)
    mn, mx = arr.min(), arr.max()
    if mx > mn: arr = (arr - mn) / (mx - mn)
    else: arr = np.zeros_like(arr)
    return arr

def load_all_samples(codes: List[str], tp_dirs: List[str]) -> np.ndarray:
    arrays = []
    for code in codes:
        tp_images = []
        for tp_dir in tp_dirs:
            fpath = find_file_for_code(tp_dir, code)
            if fpath is None: raise FileNotFoundError(f"Missing '{code}' in {tp_dir}")
            tp_images.append(load_excel_as_array(fpath))
        shapes = [a.shape for a in tp_images]
        if len(set(shapes)) > 1:
            min_h = min(s[0] for s in shapes)
            min_w = min(s[1] for s in shapes)
            tp_images = [a[:min_h, :min_w] for a in tp_images]
        arrays.append(np.stack(tp_images, axis=0))
    return np.stack(arrays, axis=0)

# =============================================================================
# Pixel Saliency / Importance
# =============================================================================
def compute_gradient_saliency(model: nn.Module, X: torch.Tensor, dims: List[int]) -> Dict[int, Dict[int, np.ndarray]]:
    N, T, H, W = X.shape
    model.eval()
    grad_maps = {d: {t: np.zeros((H, W), dtype=np.float64) for t in range(T)} for d in dims}
    batch_size = min(8, N)
    n_batches = (N + batch_size - 1) // batch_size
    for b in range(n_batches):
        s = b * batch_size
        e = min(s + batch_size, N)
        B = e - s
        X_batch = X[s:e].clone().requires_grad_(True)
        X_5d = X_batch.unsqueeze(2)
        latents = [model.encoder(X_5d[:, t]) for t in range(T)]
        for d in dims:
            for t in range(T):
                target = latents[t][:, d].sum()
                model.zero_grad()
                if X_batch.grad is not None: X_batch.grad.zero_()
                target.backward(retain_graph=True)
                grads = X_batch.grad.detach().abs()
                grad_maps[d][t] += grads[:, t].mean(dim=0).cpu().numpy() * B
        X_batch = X_batch.detach()
    for d in dims:
        for t in range(T):
            grad_maps[d][t] /= N
    return grad_maps

def compute_linear_probing(X: torch.Tensor, z_per_tp: Dict[str, np.ndarray], dims: List[int]) -> Dict[int, Dict[int, np.ndarray]]:
    X_np = X.detach().cpu().numpy()
    N, T, H, W = X_np.shape
    linear_maps = {d: {t: np.zeros((H, W), dtype=float) for t in range(T)} for d in dims}
    scaler = StandardScaler()
    for t_idx, tp in enumerate(TP_NAMES):
        X_t = X_np[:, t_idx].reshape(N, -1)
        ok_cols = (X_t.var(axis=0) > 1e-12)
        X_t_valid = X_t[:, ok_cols]
        if X_t_valid.shape[1] == 0: continue
        Xs = scaler.fit_transform(X_t_valid)
        for d in dims:
            y = z_per_tp[tp][:, d]
            ridge = Ridge(alpha=100.0, fit_intercept=True)
            ridge.fit(Xs, y)
            coefs = np.abs(ridge.coef_)
            map_1d = np.zeros(H * W, dtype=float)
            map_1d[ok_cols] = coefs
            linear_maps[d][t_idx] = map_1d.reshape(H, W)
    return linear_maps

def normalize_map(m: np.ndarray) -> np.ndarray:
    mn, mx = m.min(), m.max()
    if mx - mn < 1e-12: return np.zeros_like(m)
    return (m - mn) / (mx - mn)

def compute_consensus_maps(grad_maps, linear_maps, dims, T=3):
    consensus = {d: {} for d in dims}
    for d in dims:
        for t in range(T):
            consensus[d][t] = (normalize_map(grad_maps[d][t]) + normalize_map(linear_maps[d][t])) / 2.0
    return consensus

def extract_roi_for_map(importance_map: np.ndarray, mse_map: np.ndarray, 
                        importance_percentile: float, mse_threshold: float, min_area: int = 3) -> np.ndarray:
    imp_thresh = np.percentile(importance_map, importance_percentile)
    roi_mask = (importance_map >= imp_thresh) & (mse_map < mse_threshold)
    labeled, n_comp = ndimage.label(roi_mask)
    filtered_mask = np.zeros_like(roi_mask)
    for cid in range(1, n_comp + 1):
        cmask = labeled == cid
        if cmask.sum() >= min_area:
            filtered_mask |= cmask
    return filtered_mask

# =============================================================================
# Ridge Proxy Builder
# =============================================================================
def fit_predict_ridge(X_tr, y_tr, X_te, alpha):
    mu = X_tr.mean(axis=0, keepdims=True)
    sd = X_tr.std(axis=0, keepdims=True)
    sd[sd < 1e-12] = 1.0
    Xs_tr = (X_tr - mu) / sd
    Xs_te = (X_te - mu) / sd
    reg = Ridge(alpha=float(alpha), fit_intercept=True, random_state=SEED)
    reg.fit(Xs_tr, y_tr)
    return reg.predict(Xs_tr), reg.predict(Xs_te)

def select_alpha_inner_cv(X_tr, y_tr, alpha_grid, inner_folds=3):
    n = X_tr.shape[0]
    if n < 6: return float(alpha_grid[0])
    k = min(int(inner_folds), n - 1)
    kf = KFold(n_splits=k, shuffle=True, random_state=SEED)
    best_alpha = float(alpha_grid[0])
    best_score = -np.inf
    for a in alpha_grid:
        scores = []
        for tr_i, va_i in kf.split(np.arange(n)):
            _, pred_va = fit_predict_ridge(X_tr[tr_i], y_tr[tr_i], X_tr[va_i], alpha=a)
            ss_res = np.sum((y_tr[va_i] - pred_va)**2)
            ss_tot = np.sum((y_tr[va_i] - np.mean(y_tr[va_i]))**2)
            r2 = 1.0 - ss_res/ss_tot if ss_tot > 1e-12 else 0.0
            scores.append(r2)
        if np.mean(scores) > best_score:
            best_score = np.mean(scores)
            best_alpha = float(a)
    return best_alpha

def build_proxy_single_fold(X_np, dims, z_agg, z_per_tp, roi_masks, train_idx, test_idx, alpha_grid):
    n_train = len(train_idx)
    proxies_tr = np.zeros((n_train, len(dims)))
    proxies_te = np.zeros((1, len(dims)))

    for i, d in enumerate(dims):
        s1_tr = np.zeros((n_train, 3))
        s1_te = np.zeros((1, 3))
        
        for t_i, tp in enumerate(TP_NAMES):
            y = z_per_tp[tp][:, d]
            y_tr = y[train_idx]
            mask = roi_masks[d][tp]
            
            pixels = X_np[:, t_i].reshape(X_np.shape[0], -1)
            mask_flat = mask.ravel()
            if not mask_flat.any():
                s1_tr[:, t_i] = np.mean(y_tr)
                s1_te[:, t_i] = np.mean(y_tr)
                continue
                
            X_des = pixels[:, mask_flat]
            X_tr, X_te = X_des[train_idx], X_des[test_idx:test_idx+1]
            
            a = select_alpha_inner_cv(X_tr, y_tr, alpha_grid)
            ptr, pte = fit_predict_ridge(X_tr, y_tr, X_te, a)
            s1_tr[:, t_i] = ptr
            s1_te[:, t_i] = pte[0]

        y_agg = z_agg[:, d]
        y_agg_tr = y_agg[train_idx]
        
        if s1_tr.std() < 1e-12:
            proxies_tr[:, i] = np.mean(y_agg_tr)
            proxies_te[:, i] = np.mean(y_agg_tr)
        else:
            a2 = select_alpha_inner_cv(s1_tr, y_agg_tr, alpha_grid)
            ptr, pte = fit_predict_ridge(s1_tr, y_agg_tr, s1_te, a2)
            proxies_tr[:, i] = ptr
            proxies_te[:, i] = pte[0]

    return proxies_tr, proxies_te


# =============================================================================
# MAIN
# =============================================================================
def main():
    parser = argparse.ArgumentParser(description="Fused Pipeline: Iterative Percentile Optimization")
    parser.add_argument("--tp_dirs", type=str, required=True)
    parser.add_argument("--labels_csv", type=str, required=True)
    parser.add_argument("--step1_dir", type=str, required=True)
    parser.add_argument("--step2_dir", type=str, required=True)
    parser.add_argument("--model_dir", type=str, required=True)
    parser.add_argument("--arch_file", type=str, required=True)
    parser.add_argument("--output_dir", type=str, default="./fused_optimization_results")
    parser.add_argument("--percentiles", type=str, default="75, 80, 85, 90, 92, 95, 97")
    parser.add_argument("--mse_threshold", type=float, default=0.5)
    parser.add_argument("--latent_dim", type=int, default=512)
    args = parser.parse_args()

    outdir = Path(args.output_dir)
    outdir.mkdir(parents=True, exist_ok=True)
    step1_dir = Path(args.step1_dir)
    step2_dir = Path(args.step2_dir)
    model_dir = Path(args.model_dir)
    
    tp_dirs = [d.strip() for d in args.tp_dirs.split(",")]
    percentiles = [float(p.strip()) for p in args.percentiles.split(",")]
    
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    print("=" * 70)
    print("FUSED PIPELINE: OPTIMAL PERCENTILE SEARCH")
    print("=" * 70)
    print(f"Device:           {device}")
    print(f"Target Features:  {TARGET_FEATURES}")
    print(f"Test Percentiles: {percentiles}")
    
    # 1. Load MSE Maps
    print("\n[1/4] Loading MSE maps...")
    mse_ensemble = {}
    for t_name in TP_NAMES:
        mse_path = step2_dir / f"mse_ensemble_{t_name}.npy"
        mse_ensemble[t_name] = np.load(mse_path)

    # 2. Load EEM data
    print("\n[2/4] Loading and aligning samples...")
    codes_full, y_full, _ = load_labels(args.labels_csv)
    codes_per_tp = [set(list_sample_codes(d)) for d in tp_dirs]
    common_codes_set = set(codes_full) & codes_per_tp[0] & codes_per_tp[1] & codes_per_tp[2]
    
    keep_mask = [c in common_codes_set for c in codes_full]
    common_codes = [c for c, k in zip(codes_full, keep_mask) if k]
    y = y_full[np.array(keep_mask)]
    N = len(common_codes)
    
    X_np = load_all_samples(common_codes, tp_dirs)
    _, T, H, W = X_np.shape
    X_tensor = torch.tensor(X_np, dtype=torch.float32).to(device)

    # 3. Load Architecture & Pre-Compute Consensus Maps
    print(f"\n[3/4] Loading architecture and Pre-Computing Importance Maps ONCE...")
    arch_mod = import_arch_module(args.arch_file)
    arch_class = None
    if hasattr(arch_mod, 'ConvAutoencoderWithAttention'):
        arch_class = arch_mod.ConvAutoencoderWithAttention
    else:
        import inspect
        for name, obj in inspect.getmembers(arch_mod, inspect.isclass):
            if issubclass(obj, nn.Module) and obj is not nn.Module:
                arch_class = obj
                break
    
    all_consensus = {}
    all_z_agg = {}
    all_z_tp = {}

    for m_name, target_dims in TARGET_FEATURES.items():
        m_step1_dir = step1_dir / m_name
        possible_paths = [
            model_dir / m_name / "model.pth",      
            model_dir / f"{m_name}.pth",           
            model_dir / f"model_fold{m_name[-1]}.pth" 
        ]
        model_path = next((p for p in possible_paths if p.exists()), None)
        if model_path is None: raise FileNotFoundError(f"Model {m_name} not found.")

        # Load weights
        model = arch_class(in_channels=1, latent_dim=args.latent_dim, num_classes=2).to(device)
        _load_state_dict_flexible(model, str(model_path))
        model.eval()
        
        # Z-Vectors
        z_df = pd.read_csv(m_step1_dir / "z_agg.csv")
        code_idx_map = {str(c): i for i, c in enumerate(z_df["code"].tolist())}
        reindex = [code_idx_map[c] for c in common_codes]
        dim_cols = [f"dim{d}" for d in range(args.latent_dim)]
        all_z_agg[m_name] = z_df[dim_cols].values[reindex]
        
        all_z_tp[m_name] = {}
        for tp in TP_NAMES:
            df_tp = pd.read_csv(m_step1_dir / f"z_{tp}.csv")
            all_z_tp[m_name][tp] = df_tp[dim_cols].values[reindex]

        # Compute Importance
        print(f"  -> Pre-computing gradient saliency & linear probing for {m_name}...")
        grad_maps = compute_gradient_saliency(model, X_tensor, target_dims)
        linear_maps = compute_linear_probing(X_tensor, all_z_tp[m_name], target_dims)
        all_consensus[m_name] = compute_consensus_maps(grad_maps, linear_maps, target_dims, T=3)


    # 4. Iterative Evaluation
    print("\n[4/4] Beginning Iterative Percentile Evaluation...")
    results_history = []

    for pctl in percentiles:
        print(f"\nEvaluating Mask Percentile: {pctl}")
        
        # A. Thresholding (Instantaneous)
        roi_masks = {}
        total_px = 0
        for m_name, target_dims in TARGET_FEATURES.items():
            roi_masks[m_name] = {}
            for d in target_dims:
                roi_masks[m_name][d] = {}
                for t_idx, tp_name in enumerate(TP_NAMES):
                    imp_map = all_consensus[m_name][d][t_idx]
                    mse_map = mse_ensemble[tp_name]
                    mask = extract_roi_for_map(imp_map, mse_map, pctl, args.mse_threshold)
                    roi_masks[m_name][d][tp_name] = mask
                    total_px += mask.sum()
                    
        print(f"  -> Total ROI pixels across all 5 features: {total_px}")

        # B. LOOCV Proxy Reconstruction
        proxy_oof = np.zeros((N, len(TARGET_FEATURES_LIST)))
        classifiers = {
            "SVM": SVC(kernel="linear", probability=True, class_weight="balanced", random_state=SEED),
            "RF": RandomForestClassifier(n_estimators=500, class_weight="balanced_subsample", random_state=SEED),
            "LR": LogisticRegression(class_weight="balanced", random_state=SEED)
        }
        y_preds = {k: np.zeros(N) for k in classifiers}
        y_probs = {k: np.zeros(N) for k in classifiers}

        loo = LeaveOneOut()
        alpha_grid = [0.1, 1.0, 10.0]

        for fold, (train_idx, test_idx) in enumerate(loo.split(np.arange(N))):
            test_i = test_idx[0]
            
            bundle_res = {}
            for m_name, target_dims in TARGET_FEATURES.items():
                tr, te = build_proxy_single_fold(
                    X_np, target_dims, all_z_agg[m_name], all_z_tp[m_name], 
                    roi_masks[m_name], train_idx, test_i, alpha_grid
                )
                bundle_res[m_name] = (tr, te)
                
            fold_X_tr = np.zeros((len(train_idx), len(TARGET_FEATURES_LIST)))
            fold_X_te = np.zeros((1, len(TARGET_FEATURES_LIST)))

            for i, fname in enumerate(TARGET_FEATURES_LIST):
                idx = fname.rfind("_dim")
                mname = fname[:idx]
                dim = int(fname[idx+4:])
                b_tr, b_te = bundle_res[mname]
                col_idx = TARGET_FEATURES[mname].index(dim)
                fold_X_tr[:, i] = b_tr[:, col_idx]
                fold_X_te[:, i] = b_te[:, col_idx]
                
            proxy_oof[test_i] = fold_X_te[0]

            scaler = StandardScaler()
            X_tr_s = scaler.fit_transform(fold_X_tr)
            X_te_s = scaler.transform(fold_X_te)
            
            for name, clf in classifiers.items():
                clf.fit(X_tr_s, y[train_idx])
                if hasattr(clf, "predict_proba"): prob = clf.predict_proba(X_te_s)[0, 1]
                else: prob = clf.decision_function(X_te_s)[0]
                y_probs[name][test_i] = prob
                y_preds[name][test_i] = int(prob >= 0.5)

        # C. Calculate Metrics & Power
        step_results = {"percentile": pctl, "total_pixels": total_px}
        best_acc = 0.0
        
        for name in classifiers:
            acc = accuracy_score(y, y_preds[name])
            auc = roc_auc_score(y, y_probs[name])
            step_results[f"{name}_acc"] = acc
            step_results[f"{name}_auc"] = auc
            if acc > best_acc: best_acc = acc
                
        power = calculate_power(best_acc, N)
        step_results["power"] = power
        step_results["best_acc"] = best_acc
        results_history.append(step_results)

        print(f"  -> Best Accuracy: {best_acc:.3f} | Statistical Power: {power:.3f}")

    # 5. Save Summary
    df_history = pd.DataFrame(results_history)
    df_history.to_csv(outdir / "percentile_optimization_history.csv", index=False)
    
    print("\n" + "="*70)
    print("Optimization Complete!")
    print(df_history[["percentile", "total_pixels", "best_acc", "power"]].to_string(index=False))
    print(f"\nResults saved to: {outdir / 'percentile_optimization_history.csv'}")

if __name__ == "__main__":
    main()