#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
apply_proxy_latents_batch_corrected.py

Compute proxy latents for a NEW dataset using exported parameters from the
"shallow conv + ridge" proxy distillation, WITH batch-effect correction.

The core issue: the frozen input_scaler (mean/scale from the training batch)
produces systematically shifted z-scores when applied to data with a different
intensity distribution (batch effect). This corrupts the ridge predictions
and downstream SVM accuracy.

Three correction strategies are provided:

  --batch_correction none
      No correction. Identical to apply_proxy_latents_to_new_dataset_v2.py.

  --batch_correction fresh_zscore   (RECOMMENDED)
      Replace the frozen input scaler with a fresh StandardScaler computed on
      the new dataset. This exactly removes any location-scale batch effect
      (additive offset + multiplicative gain) because:
        (a*x + b - E[a*x+b]) / std(a*x+b) = (x - E[x]) / std(x)
      The ridge coefficients remain valid since they expect z-scored inputs.

  --batch_correction align_to_training
      Two-step: (1) z-score the new data with its own statistics, then
      (2) "un-z-score" back into the training feature space using the frozen
      scaler statistics. Then apply the frozen scaler as usual.
      Algebraically equivalent to fresh_zscore for the ridge step, but the
      intermediate aligned features are in the original training scale,
      which can be useful for diagnostics and QC.

  --batch_correction combat_proxy
      After computing proxy latents (with fresh_zscore), apply a simple
      location-scale alignment on the final proxy values so their distribution
      matches the training proxy distribution (using the frozen output_scaler
      statistics as the target). Use this if residual drift remains after
      feature-level correction.

Inputs:
  - deterministic_feature_spec.json  (per-dim ROI, input scaler, ridge coef/intercept, etc.)
  - optional proxy_latent_scalers.json (output standardization)
  - dataset folder with timepoint subfolders

Outputs:
  - proxy_latents.csv
  - batch_correction_diagnostics.json  (feature-space shift statistics for QC)

Based on: apply_proxy_latents_to_new_dataset_v2.py
"""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
from typing import Dict, List, Tuple, Optional

import numpy as np
import pandas as pd
from openpyxl import load_workbook


# ─── Fixed kernels (must match training script) ──────────────────────────────

def _kernels() -> Dict[str, np.ndarray]:
    sobel_x = np.array([[ 1, 0,-1],
                        [ 2, 0,-2],
                        [ 1, 0,-1]], dtype=float)
    sobel_y = np.array([[ 1, 2, 1],
                        [ 0, 0, 0],
                        [-1,-2,-1]], dtype=float)
    laplace_4 = np.array([[ 0, 1, 0],
                          [ 1,-4, 1],
                          [ 0, 1, 0]], dtype=float)
    diag_45 = np.array([[ 2, 1, 0],
                        [ 1, 0,-1],
                        [ 0,-1,-2]], dtype=float)
    diag_135 = np.array([[ 0, 1, 2],
                         [-1, 0, 1],
                         [-2,-1, 0]], dtype=float)
    blur = np.array([[1,2,1],
                     [2,4,2],
                     [1,2,1]], dtype=float) / 16.0
    return {
        "sobel_x": sobel_x,
        "sobel_y": sobel_y,
        "laplace4": laplace_4,
        "diag_45": diag_45,
        "diag_135": diag_135,
        "blur": blur,
    }


STAT_NAMES = ("mean", "std", "min", "max", "p05", "p50", "p95", "mean_abs", "energy")


def _conv2_same_symm(img: np.ndarray, k: np.ndarray) -> np.ndarray:
    try:
        from scipy.signal import convolve2d
        return convolve2d(img, k, mode="same", boundary="symm")
    except Exception:
        kh, kw = k.shape
        ph, pw = kh // 2, kw // 2
        padded = np.pad(img, ((ph, ph), (pw, pw)), mode="symmetric")
        out = np.zeros_like(img, dtype=float)
        kf = np.flipud(np.fliplr(k))
        for i in range(out.shape[0]):
            for j in range(out.shape[1]):
                patch = padded[i:i+kh, j:j+kw]
                out[i, j] = float(np.sum(patch * kf))
        return out


def _summ_stats(x: np.ndarray) -> Dict[str, float]:
    x = np.asarray(x, dtype=float)
    return {
        "mean": float(np.mean(x)),
        "std": float(np.std(x)),
        "min": float(np.min(x)),
        "max": float(np.max(x)),
        "p05": float(np.percentile(x, 5)),
        "p50": float(np.percentile(x, 50)),
        "p95": float(np.percentile(x, 95)),
        "mean_abs": float(np.mean(np.abs(x))),
        "energy": float(np.mean(x * x)),
    }


def _parse_float(v, allow_comma_decimal: bool) -> float:
    if v is None:
        return 0.0
    if isinstance(v, (int, float, np.integer, np.floating)):
        if np.isnan(v):
            return 0.0
        return float(v)
    s = str(v).strip()
    if s == "":
        return 0.0
    if allow_comma_decimal and "," in s and "." not in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except Exception:
        return 0.0


def read_eem_excel(filepath: Path, allow_comma_decimal: bool = False) -> np.ndarray:
    wb = load_workbook(str(filepath), data_only=True, read_only=True)
    ws = wb.active
    max_row = ws.max_row
    max_col = ws.max_column
    if max_row < 2 or max_col < 2:
        raise ValueError(f"File {filepath} is too small: rows={max_row} cols={max_col}")
    H = max_row - 1
    W = max_col - 1
    mat = np.zeros((H, W), dtype=float)
    for i, r in enumerate(range(2, max_row + 1)):
        row_cells = ws[r]
        for j, cell in enumerate(row_cells[1:max_col]):
            mat[i, j] = _parse_float(cell.value, allow_comma_decimal=allow_comma_decimal)
    return mat


# ─── Helpers ──────────────────────────────────────────────────────────────────

def _load_labels(labels_csv: Optional[str]) -> Tuple[Optional[pd.DataFrame], Optional[str], Optional[str]]:
    if not labels_csv:
        return None, None, None
    df = pd.read_csv(labels_csv)
    cols = [c.lower() for c in df.columns]
    subj_candidates = ["subject", "sample", "id", "code", "filename"]
    grp_candidates = ["group", "label", "class", "condition"]
    subject_col = None
    group_col = None
    for cand in subj_candidates:
        if cand in cols:
            subject_col = df.columns[cols.index(cand)]
            break
    if subject_col is None:
        subject_col = df.columns[0]
    for cand in grp_candidates:
        if cand in cols:
            group_col = df.columns[cols.index(cand)]
            break
    return df, subject_col, group_col


def _infer_subjects_from_dir(tp_dir: Path) -> List[str]:
    return [p.stem for p in sorted(tp_dir.glob("*.xlsx"))]


def _build_mean_over_time(subject: str, data_dir: Path, tp_dirs: List[str],
                          allow_comma_decimal: bool) -> np.ndarray:
    mats = []
    for tp in tp_dirs:
        f = data_dir / tp / f"{subject}.xlsx"
        if not f.exists():
            raise FileNotFoundError(f"Missing file for subject={subject}: {f}")
        mats.append(read_eem_excel(f, allow_comma_decimal=allow_comma_decimal))
    shapes = {m.shape for m in mats}
    if len(shapes) != 1:
        raise ValueError(f"Shape mismatch across timepoints for {subject}: {shapes}")
    return np.mean(np.stack(mats, axis=0), axis=0)


def _extract_features_for_dim(mean_mat: np.ndarray, dim_spec: dict,
                               kernels: Dict[str, np.ndarray]) -> np.ndarray:
    roi = dim_spec["roi"]
    y0, y1 = int(roi["y0"]), int(roi["y1"])
    x0, x1 = int(roi["x0"]), int(roi["x1"])
    patch = mean_mat[y0:y1+1, x0:x1+1].astype(float, copy=False)
    if patch.size == 0:
        raise ValueError(f"Empty ROI: roi={roi}, mat_shape={mean_mat.shape}")

    responses: Dict[str, np.ndarray] = {"raw": patch}
    for kname, k in kernels.items():
        responses[kname] = _conv2_same_symm(patch, k)

    resp_stats: Dict[Tuple[str, str], float] = {}
    for kname, arr in responses.items():
        s = _summ_stats(arr.ravel())
        for stat_name, val in s.items():
            resp_stats[(kname, stat_name)] = val

    feats = []
    for fname in dim_spec["features"]:
        kname, stat = fname.split(":")
        if (kname, stat) not in resp_stats:
            raise KeyError(f"Missing feature {fname}")
        feats.append(resp_stats[(kname, stat)])
    return np.asarray(feats, dtype=float)


def _standardize(x: np.ndarray, mean: np.ndarray, scale: np.ndarray) -> np.ndarray:
    mean = np.asarray(mean, dtype=float)
    scale = np.asarray(scale, dtype=float)
    scale_safe = np.where(scale == 0.0, 1.0, scale)
    return (x - mean) / scale_safe


# ─── BATCH CORRECTION LOGIC ──────────────────────────────────────────────────

def _two_pass_proxy_latents(
    subjects: List[str],
    group_map: Dict[str, str],
    data_dir: Path,
    tp_dirs: List[str],
    dims_spec: List[dict],
    kernels: Dict[str, np.ndarray],
    allow_comma_decimal: bool,
    batch_correction: str,
    apply_output_scaler: bool,
    out_scalers: Optional[dict],
    strict: bool,
) -> Tuple[pd.DataFrame, dict]:
    """
    Two-pass approach:
      Pass 1: Extract raw features for ALL subjects, per dimension.
      Pass 2: Compute batch-corrected standardization, then ridge predict.

    Returns (proxy_df, diagnostics_dict).
    """

    diagnostics = {"batch_correction": batch_correction, "dims": {}}

    # ── Pass 1: Extract raw features ──────────────────────────────────────────
    # mean_mats[subj] = mean-over-timepoints EEM matrix
    mean_mats: Dict[str, np.ndarray] = {}
    valid_subjects = []
    skipped = 0

    for subj in subjects:
        try:
            mean_mats[subj] = _build_mean_over_time(
                subj, data_dir, tp_dirs, allow_comma_decimal
            )
            valid_subjects.append(subj)
        except Exception as e:
            skipped += 1
            print(f"[WARN] skipping subject={subj}: {e}")
            if strict:
                raise

    # raw_features[dim_idx] = np.ndarray of shape (n_valid_subjects, n_features)
    raw_features: Dict[int, np.ndarray] = {}
    for di, dspec in enumerate(dims_spec):
        dim = int(dspec["dim"])
        X_rows = []
        for subj in valid_subjects:
            x = _extract_features_for_dim(mean_mats[subj], dspec, kernels)
            X_rows.append(x)
        raw_features[dim] = np.vstack(X_rows) if X_rows else np.empty((0, len(dspec["features"])))

    # ── Pass 2: Batch-corrected standardization + ridge ───────────────────────
    dims_list = [int(d["dim"]) for d in dims_spec]
    out_cols = [f"proxy_dim_{d}" for d in dims_list]

    proxy_values: Dict[str, Dict[str, float]] = {subj: {} for subj in valid_subjects}

    for di, dspec in enumerate(dims_spec):
        dim = int(dspec["dim"])
        X_raw = raw_features[dim]  # (n_subjects, n_features)
        n_subj, n_feat = X_raw.shape

        # Frozen training scaler stats
        train_mean = np.asarray(dspec["input_scaler_mean"], dtype=float)
        train_scale = np.asarray(dspec["input_scaler_scale"], dtype=float)

        # Compute new-batch statistics
        if n_subj > 1:
            new_mean = X_raw.mean(axis=0)
            new_scale = X_raw.std(axis=0)
        else:
            # Single sample: fall back to frozen scaler
            new_mean = train_mean
            new_scale = train_scale

        new_scale_safe = np.where(new_scale == 0.0, 1.0, new_scale)
        train_scale_safe = np.where(train_scale == 0.0, 1.0, train_scale)

        # ── Choose standardization strategy ───────────────────────────────────
        if batch_correction == "none":
            # Original behavior: use frozen scaler
            X_std = _standardize(X_raw, train_mean, train_scale)

        elif batch_correction == "fresh_zscore":
            # Replace frozen scaler with fresh z-score on the new batch.
            # This removes location-scale batch effects exactly.
            X_std = _standardize(X_raw, new_mean, new_scale)

        elif batch_correction == "align_to_training":
            # Step 1: z-score with new-batch stats
            # Step 2: un-z-score into training space
            # Step 3: apply frozen scaler (net effect same as fresh_zscore
            #         for the ridge, but intermediate values are in training scale)
            X_aligned = ((X_raw - new_mean) / new_scale_safe) * train_scale_safe + train_mean
            X_std = _standardize(X_aligned, train_mean, train_scale)

        elif batch_correction == "combat_proxy":
            # Feature-level: use fresh_zscore
            X_std = _standardize(X_raw, new_mean, new_scale)

        else:
            raise ValueError(f"Unknown batch_correction mode: {batch_correction}")

        # ── Ridge prediction ──────────────────────────────────────────────────
        coef = np.asarray(dspec["ridge_coef"], dtype=float)
        intercept = float(dspec["ridge_intercept"])
        yhat = X_std @ coef + intercept  # (n_subjects,)

        # ── Output scaling ────────────────────────────────────────────────────
        if apply_output_scaler:
            out_sc = dspec.get("output_scaler", None)
            if isinstance(out_sc, dict) and "mean" in out_sc:
                mu = float(out_sc["mean"])
                # Handle both "std" and "scale" keys
                sd_key = "std" if "std" in out_sc else "scale"
                sd = float(out_sc.get(sd_key, 1.0))
                if sd == 0.0:
                    sd = 1.0

                if batch_correction == "combat_proxy":
                    # Additional proxy-level alignment: force the output
                    # distribution to match training statistics
                    yhat_mean = float(np.mean(yhat))
                    yhat_std = float(np.std(yhat))
                    if yhat_std == 0.0:
                        yhat_std = 1.0
                    # Align to training proxy distribution, then standardize
                    yhat_aligned = (yhat - yhat_mean) / yhat_std * sd + mu
                    yhat = (yhat_aligned - mu) / sd
                else:
                    yhat = (yhat - mu) / sd

            elif out_scalers is not None:
                key = f"proxy_dim_{dim}"
                if key in out_scalers:
                    mu = float(out_scalers[key].get("mean", 0.0))
                    sd = float(out_scalers[key].get("std", 1.0))
                    if sd == 0.0:
                        sd = 1.0
                    yhat = (yhat - mu) / sd

        # ── Store results ─────────────────────────────────────────────────────
        for si, subj in enumerate(valid_subjects):
            proxy_values[subj][f"proxy_dim_{dim}"] = float(yhat[si])

        # ── Diagnostics ───────────────────────────────────────────────────────
        # Compute how much the batch statistics differ from training
        mean_shift = float(np.mean(np.abs(new_mean - train_mean) / train_scale_safe))
        scale_ratio = float(np.mean(new_scale_safe / train_scale_safe))
        diagnostics["dims"][str(dim)] = {
            "mean_shift_in_training_sigmas": mean_shift,
            "scale_ratio_new_over_training": scale_ratio,
            "n_features": n_feat,
            "n_subjects": n_subj,
        }

    # ── Build output DataFrame ────────────────────────────────────────────────
    rows = []
    for subj in valid_subjects:
        row = {"subject": subj}
        if subj in group_map:
            row["group"] = group_map[subj]
        row.update(proxy_values[subj])
        rows.append(row)

    df_out = pd.DataFrame(rows)
    base_cols = ["subject"] + (["group"] if "group" in df_out.columns else [])
    df_out = df_out[base_cols + out_cols]

    diagnostics["skipped_subjects"] = skipped
    diagnostics["valid_subjects"] = len(valid_subjects)

    return df_out, diagnostics


# ─── Main ─────────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(
        description="Apply proxy latents to a new dataset WITH batch-effect correction."
    )
    ap.add_argument("--feature_spec_json", required=True,
                    help="Path to deterministic_feature_spec.json")
    ap.add_argument("--proxy_scalers_json", default=None,
                    help="Optional path to proxy_latent_scalers.json.")
    ap.add_argument("--data_dir", required=True,
                    help="Root directory containing timepoint subfolders.")
    ap.add_argument("--tp_dirs", default=None,
                    help="Comma-separated timepoint directory names.")
    ap.add_argument("--labels_csv", default=None,
                    help="Optional labels CSV.")
    ap.add_argument("--out_csv", required=True,
                    help="Where to write proxy_latents.csv")
    ap.add_argument("--allow_comma_decimal", action="store_true")
    ap.add_argument("--apply_output_scaler", action="store_true",
                    help="Apply per-dimension output scaling (recommended).")
    ap.add_argument("--strict", action="store_true",
                    help="Fail on any missing subject file.")

    # ── Batch correction arguments ────────────────────────────────────────────
    ap.add_argument("--batch_correction", default="fresh_zscore",
                    choices=["none", "fresh_zscore", "align_to_training", "combat_proxy"],
                    help="Batch correction strategy (default: fresh_zscore). "
                         "'none': use frozen scaler (original behavior). "
                         "'fresh_zscore': re-standardize with new-batch stats (recommended). "
                         "'align_to_training': align feature distributions to training space. "
                         "'combat_proxy': fresh_zscore + proxy-level distribution matching.")
    ap.add_argument("--save_diagnostics", action="store_true",
                    help="Save batch_correction_diagnostics.json for QC.")

    args = ap.parse_args()

    spec = json.load(open(args.feature_spec_json, "r", encoding="utf-8"))
    data_dir = Path(args.data_dir)

    tp_dirs = spec.get("tp_dirs", [])
    if args.tp_dirs:
        tp_dirs = [t.strip() for t in args.tp_dirs.split(",") if t.strip()]
    if not tp_dirs:
        raise ValueError("No tp_dirs found.")

    for tp in tp_dirs:
        d = data_dir / tp
        if not d.exists():
            raise FileNotFoundError(f"Timepoint directory not found: {d}")

    out_scalers = None
    if args.proxy_scalers_json:
        out_scalers = json.load(open(args.proxy_scalers_json, "r", encoding="utf-8"))

    labels_df, subj_col, grp_col = _load_labels(args.labels_csv)
    if labels_df is not None:
        subjects = [str(s) for s in labels_df[subj_col].tolist()]
        group_map = {}
        if grp_col is not None:
            for _, row in labels_df.iterrows():
                group_map[str(row[subj_col])] = row[grp_col]
    else:
        subjects = _infer_subjects_from_dir(data_dir / tp_dirs[0])
        group_map = {}

    kernels = _kernels()
    dims_spec = spec["dims"]

    # ── Two-pass extraction with batch correction ─────────────────────────────
    df_out, diagnostics = _two_pass_proxy_latents(
        subjects=subjects,
        group_map=group_map,
        data_dir=data_dir,
        tp_dirs=tp_dirs,
        dims_spec=dims_spec,
        kernels=kernels,
        allow_comma_decimal=args.allow_comma_decimal,
        batch_correction=args.batch_correction,
        apply_output_scaler=args.apply_output_scaler,
        out_scalers=out_scalers,
        strict=args.strict,
    )

    out_path = Path(args.out_csv)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    df_out.to_csv(out_path, index=False)
    print(f"Wrote: {out_path}")
    print(f"  batch_correction = {args.batch_correction}")
    print(f"  valid subjects   = {diagnostics['valid_subjects']}")
    print(f"  skipped subjects = {diagnostics['skipped_subjects']}")

    # Print summary of batch shift per dimension
    print("\n=== Batch shift diagnostics (per dimension) ===")
    for dim_str, info in diagnostics["dims"].items():
        shift = info["mean_shift_in_training_sigmas"]
        ratio = info["scale_ratio_new_over_training"]
        flag = " *** LARGE SHIFT ***" if shift > 2.0 or abs(ratio - 1.0) > 0.5 else ""
        print(f"  dim {dim_str}: mean_shift={shift:.3f}σ  scale_ratio={ratio:.3f}{flag}")

    if args.save_diagnostics:
        diag_path = out_path.parent / "batch_correction_diagnostics.json"
        with open(diag_path, "w", encoding="utf-8") as f:
            json.dump(diagnostics, f, indent=2)
        print(f"\nWrote: {diag_path}")


if __name__ == "__main__":
    main()
