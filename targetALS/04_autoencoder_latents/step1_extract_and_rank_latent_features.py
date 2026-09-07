#!/usr/bin/env python3
"""
STEP 1: Extract Latent Features from 5-Fold Pre-trained Models & Rank by Importance
====================================================================================

This script operates on your EXISTING 5-fold ConvAutoencoderWithAttention models.
It does NOT retrain anything.

What it does:
  1. Loads 5 pre-trained models (best_model_fold_1.pth ... best_model_fold_5.pth)
  2. For ALL 39 samples (ALS + CTRL), extracts:
       a) Per-timepoint latent vectors z_t (shape: [N, 3, latent_dim])
       b) Attention-aggregated z_agg (shape: [N, latent_dim])
       c) Attention weights per timepoint
  3. Ranks latent dimensions by MULTIPLE importance criteria:
       - Classifier head weight magnitude (from the trained classification head)
       - Univariate AUC (how well each dim alone separates ALS vs CTRL)
       - Absolute Welch t-statistic (group difference in each dim)
       - Spearman correlation with binary label
  4. Produces a consensus ranking across all 5 models
  5. Saves:
       - Per-model latent features (z_agg and per-timepoint z_t)
       - Per-model importance tables
       - Ensemble-averaged importance ranking
       - Ensemble z_agg (sign-aligned and averaged)

USAGE:
------
python step1_extract_and_rank_latent_features.py \
    --arch_file conv_autoencoder_detailed.py \
    --model_dir ./5_fold_models_original \
    --tp_dirs ./out_0h,./out_6h,./out_24h \
    --labels_csv sample_labels.csv \
    --output_dir ./step1_latent_features \
    --latent_dim 512
"""

import os
import sys
import argparse
import importlib.util
from pathlib import Path
from typing import List, Dict, Tuple, Optional
from collections import OrderedDict

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import torch
import torch.nn as nn

from sklearn.metrics import roc_auc_score
from scipy.stats import ttest_ind, spearmanr, rankdata

import warnings
warnings.filterwarnings("ignore")

SEED = 42
np.random.seed(SEED)
torch.manual_seed(SEED)
DEVICE = torch.device("cuda" if torch.cuda.is_available() else "cpu")


# =============================================================================
# Architecture import
# =============================================================================

def import_arch_module(arch_file: str):
    """Dynamically import the model architecture module."""
    spec = importlib.util.spec_from_file_location("arch_module", arch_file)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Could not import architecture file: {arch_file}")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# =============================================================================
# Data loading
# =============================================================================

def load_labels(csv_path: str):
    """Load sample labels CSV. Returns (codes, y_binary, group_strings)."""
    df = pd.read_csv(csv_path)
    codes = df["code"].tolist()
    groups = df["group"].values
    y = (groups == "ALS").astype(int)
    return codes, y, groups


def list_sample_codes(tp_dir: str) -> List[str]:
    """List all sample codes (without _Xh.xlsx suffix) in a timepoint directory."""
    codes = []
    for f in Path(tp_dir).glob("*.xlsx"):
        if f.name.startswith("~$"):
            continue
        # Remove the timepoint suffix to get the code
        # e.g., "1_3d__570_opj_with_emission_0h.xlsx" -> extract code
        name = f.stem  # e.g., "1_3d__570_opj_with_emission_0h"
        # Remove the _0h / _6h / _24h suffix
        for suffix in ["_0h", "_6h", "_24h"]:
            if name.endswith(suffix):
                name = name[:-len(suffix)]
                break
        codes.append(name)
    return codes


def find_file_for_code(tp_dir: str, code: str) -> Optional[str]:
    """Find the Excel file for a given sample code in a timepoint directory."""
    tp_path = Path(tp_dir)
    
    # 1. Try exact match (e.g. "code.xlsx")
    # This handles your case where files are just named by the code
    exact_path = tp_path / f"{code}.xlsx"
    if exact_path.exists():
        return str(exact_path)

    # 2. Try flexible suffix matching (e.g. "code_0h.xlsx")
    for f in tp_path.glob("*.xlsx"):
        if f.name.startswith("~$"):
            continue
        stem = f.stem
        
        # Check against suffixes
        for suffix in ["_0h", "_6h", "_24h"]:
            if stem.endswith(suffix):
                # Check if the stem *without* suffix matches the code
                file_code = stem[:-len(suffix)]
                if file_code == code:
                    return str(f)
                
                # Double check for potential "._" vs "_" issues if needed
                if file_code.replace('.', '_') == code:
                    return str(f)

    return None

def load_excel_as_array(filepath: str) -> np.ndarray:
    """
    Read EEM matrix from Excel. Skip first row (headers) and first column (emission labels).
    Min-max normalize to [0, 1].
    """
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    data = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx == 0:
            continue
        row_data = []
        for col_idx, cell in enumerate(row):
            if col_idx == 0:
                continue
            if isinstance(cell, (int, float)) and cell is not None:
                row_data.append(float(cell))
            else:
                row_data.append(0.0)
        if row_data:
            data.append(row_data)
    arr = np.array(data, dtype=np.float32)
    mn, mx = arr.min(), arr.max()
    if mx > mn:
        arr = (arr - mn) / (mx - mn)
    else:
        arr = np.zeros_like(arr)
    return arr


def load_all_samples(codes: List[str], tp_dirs: List[str]) -> np.ndarray:
    """Load all samples across all timepoints. Returns [N, 3, H, W]."""
    arrays = []
    for i, code in enumerate(codes):
        tp_images = []
        for tp_dir in tp_dirs:
            fpath = find_file_for_code(tp_dir, code)
            if fpath is None:
                raise FileNotFoundError(f"No file found for code '{code}' in {tp_dir}")
            arr = load_excel_as_array(fpath)
            tp_images.append(arr)
        # Verify consistent shapes
        shapes = [a.shape for a in tp_images]
        if len(set(shapes)) > 1:
            print(f"[WARN] Inconsistent shapes for {code}: {shapes}. Using min dims.")
            min_h = min(s[0] for s in shapes)
            min_w = min(s[1] for s in shapes)
            tp_images = [a[:min_h, :min_w] for a in tp_images]
        arrays.append(np.stack(tp_images, axis=0))  # [3, H, W]

    X = np.stack(arrays, axis=0)  # [N, 3, H, W]
    print(f"  Loaded {X.shape[0]} samples, shape per sample: {X.shape[1:]}")
    return X


# =============================================================================
# Model loading
# =============================================================================

def _load_state_dict_flexible(model: nn.Module, path: str):
    """Load state dict with flexible key matching (handles prefix mismatches)."""
    sd = torch.load(path, map_location=DEVICE, weights_only=False)

    # Try direct load first
    try:
        model.load_state_dict(sd, strict=True)
        return
    except RuntimeError:
        pass

    # Try stripping 'module.' prefix
    new_sd = OrderedDict()
    for k, v in sd.items():
        new_key = k.replace("module.", "")
        new_sd[new_key] = v
    try:
        model.load_state_dict(new_sd, strict=True)
        return
    except RuntimeError:
        pass

    # Flexible: match by shape
    model_sd = model.state_dict()
    matched = OrderedDict()
    for mk, mv in model_sd.items():
        for sk, sv in sd.items():
            if sv.shape == mv.shape and mk.split(".")[-1] == sk.split(".")[-1]:
                matched[mk] = sv
                break
        else:
            matched[mk] = mv  # keep init
    model.load_state_dict(matched, strict=False)
    n_loaded = sum(1 for mk in model_sd if mk in matched and not torch.equal(matched[mk], model_sd[mk]))
    print(f"    [WARN] Flexible load: matched {n_loaded}/{len(model_sd)} params from {Path(path).name}")


# =============================================================================
# Feature extraction
# =============================================================================

def extract_latents_per_timepoint(model: nn.Module, X: torch.Tensor) -> np.ndarray:
    """
    Extract per-timepoint latent vectors.
    X: [N, 3, H, W]
    Returns: [N, 3, latent_dim]
    """
    model.eval()
    N, T, H, W = X.shape
    X_5d = X.unsqueeze(2)  # [N, 3, 1, H, W]
    z_list = []
    with torch.no_grad():
        for t in range(T):
            z_t = model.encoder(X_5d[:, t])  # [N, latent_dim]
            z_list.append(z_t.cpu().numpy())
    return np.stack(z_list, axis=1)  # [N, 3, latent_dim]


def compute_z_agg(model: nn.Module, X: torch.Tensor) -> Tuple[np.ndarray, np.ndarray]:
    """
    Compute attention-aggregated latent z_agg and attention weights.
    X: [N, 3, H, W]
    Returns: (z_agg [N, latent_dim], attn_weights [N, 3])
    """
    model.eval()
    N, T, H, W = X.shape
    X_5d = X.unsqueeze(2)
    z_list = []
    with torch.no_grad():
        for t in range(T):
            z_t = model.encoder(X_5d[:, t])
            z_list.append(z_t)
        z_stack = torch.stack(z_list, dim=1)  # [N, 3, latent_dim]
        z_agg, attn = model.attention(z_stack)
    return z_agg.detach().cpu().numpy(), attn.detach().cpu().numpy()


# =============================================================================
# Importance ranking methods
# =============================================================================

def importance_classifier_head(model: nn.Module, latent_dim: int) -> np.ndarray:
    """
    Extract importance from classifier head weights.
    Uses L2 norm of weight vectors mapping each latent dim to output logits.
    """
    # Find the first Linear layer in the classifier that takes latent_dim input
    lin = None
    for m in model.classifier.modules():
        if isinstance(m, nn.Linear) and m.in_features == latent_dim:
            lin = m
            break
    if lin is None:
        # Fallback: search all modules
        for m in model.modules():
            if isinstance(m, nn.Linear) and m.in_features == latent_dim:
                lin = m
                break
    if lin is None:
        raise RuntimeError(f"No Linear(in_features={latent_dim}) found in model.")

    w = lin.weight.detach().cpu().numpy()  # [out_features, latent_dim]
    # L2 norm across output dimension = importance of each latent dim
    imp = np.sqrt(np.sum(w ** 2, axis=0))  # [latent_dim]
    return imp


def importance_univariate_auc(z_agg: np.ndarray, y: np.ndarray) -> np.ndarray:
    """
    For each latent dimension, compute AUC for separating ALS (1) vs CTRL (0).
    Returns AUC values (0.5 = chance, 1.0 = perfect, values <0.5 inverted to >0.5).
    """
    D = z_agg.shape[1]
    aucs = np.full(D, 0.5)
    for d in range(D):
        vals = z_agg[:, d]
        if np.std(vals) < 1e-12:
            continue
        try:
            auc = roc_auc_score(y, vals)
            aucs[d] = max(auc, 1 - auc)  # Ensure ≥ 0.5
        except Exception:
            pass
    return aucs


def importance_welch_t(z_agg: np.ndarray, y: np.ndarray) -> np.ndarray:
    """Absolute Welch t-statistic per latent dimension."""
    D = z_agg.shape[1]
    t_stats = np.zeros(D)
    idx0 = y == 0
    idx1 = y == 1
    for d in range(D):
        try:
            t, _ = ttest_ind(z_agg[idx0, d], z_agg[idx1, d], equal_var=False)
            t_stats[d] = abs(t)
        except Exception:
            pass
    return t_stats


def importance_spearman(z_agg: np.ndarray, y: np.ndarray) -> np.ndarray:
    """Absolute Spearman rho with binary label per latent dimension."""
    D = z_agg.shape[1]
    rhos = np.zeros(D)
    for d in range(D):
        try:
            rho, _ = spearmanr(z_agg[:, d], y)
            rhos[d] = abs(rho)
        except Exception:
            pass
    return rhos


def compute_all_importances(model: nn.Module, z_agg: np.ndarray, y: np.ndarray,
                            latent_dim: int) -> pd.DataFrame:
    """Compute all importance metrics for one model. Returns DataFrame indexed by dim."""
    imp_head = importance_classifier_head(model, latent_dim)
    imp_auc = importance_univariate_auc(z_agg, y)
    imp_t = importance_welch_t(z_agg, y)
    imp_rho = importance_spearman(z_agg, y)

    df = pd.DataFrame({
        "dim": np.arange(latent_dim),
        "head_weight_l2": imp_head,
        "univariate_auc": imp_auc,
        "welch_t": imp_t,
        "spearman_abs_rho": imp_rho,
    })

    # Rank each metric (higher = more important, so ascending=False → rank 1 = best)
    for col in ["head_weight_l2", "univariate_auc", "welch_t", "spearman_abs_rho"]:
        df[f"{col}_rank"] = rankdata(-df[col].values, method="min")

    # Mean rank across all criteria
    rank_cols = [c for c in df.columns if c.endswith("_rank")]
    df["mean_rank"] = df[rank_cols].mean(axis=1)
    df = df.sort_values("mean_rank")

    return df


# =============================================================================
# Ensemble aggregation
# =============================================================================

def ensemble_sign_align_and_average(Zs: List[np.ndarray]) -> np.ndarray:
    """
    Sign-align all fold z_agg arrays to fold 0 (reference), then average.
    Zs: list of [N, latent_dim] arrays.
    """
    Z_ref = Zs[0]
    aligned = [Z_ref.copy()]
    for z in Zs[1:]:
        z2 = z.copy()
        for d in range(z.shape[1]):
            a = Z_ref[:, d]
            b = z2[:, d]
            if np.std(a) > 1e-12 and np.std(b) > 1e-12:
                corr = np.corrcoef(a, b)[0, 1]
                if corr < 0:
                    z2[:, d] = -z2[:, d]
        aligned.append(z2)
    return np.mean(np.stack(aligned), axis=0)


def consensus_importance(importance_dfs: List[pd.DataFrame], latent_dim: int) -> pd.DataFrame:
    """
    Average importance metrics across models. Recompute consensus ranking.
    """
    metric_cols = ["head_weight_l2", "univariate_auc", "welch_t", "spearman_abs_rho"]
    rank_cols = [f"{c}_rank" for c in metric_cols]

    # Collect per-model values
    all_metrics = {col: [] for col in metric_cols}
    all_ranks = {col: [] for col in rank_cols}

    for df in importance_dfs:
        df_sorted = df.sort_values("dim")
        for col in metric_cols:
            all_metrics[col].append(df_sorted[col].values)
        for col in rank_cols:
            all_ranks[col].append(df_sorted[col].values)

    # Average
    result = pd.DataFrame({"dim": np.arange(latent_dim)})
    for col in metric_cols:
        result[f"{col}_mean"] = np.mean(all_metrics[col], axis=0)
        result[f"{col}_std"] = np.std(all_metrics[col], axis=0)
    for col in rank_cols:
        result[f"{col}_mean"] = np.mean(all_ranks[col], axis=0)

    # Consensus mean rank
    avg_rank_cols = [f"{c}_rank_mean" for c in metric_cols]
    result["consensus_mean_rank"] = result[avg_rank_cols].mean(axis=1)

    # Final consensus rank
    result["consensus_rank"] = rankdata(result["consensus_mean_rank"].values, method="min")
    result = result.sort_values("consensus_rank")

    return result


# =============================================================================
# Main
# =============================================================================

def main():
    parser = argparse.ArgumentParser(
        description="Step 1: Extract latent features from 5-fold models and rank by importance"
    )
    parser.add_argument("--arch_file", type=str, required=True,
                        help="Path to conv_autoencoder_detailed.py (architecture definition)")
    parser.add_argument("--model_dir", type=str, required=True,
                        help="Directory containing best_model_fold_*.pth files")
    parser.add_argument("--tp_dirs", type=str, required=True,
                        help="Comma-separated: dir_0h,dir_6h,dir_24h")
    parser.add_argument("--labels_csv", type=str, required=True,
                        help="CSV with columns: code, group")
    parser.add_argument("--output_dir", type=str, default="./step1_latent_features",
                        help="Output directory")
    parser.add_argument("--latent_dim", type=int, default=512,
                        help="Latent dimension of the autoencoder")
    parser.add_argument("--top_k", type=int, default=20,
                        help="Number of top dimensions to highlight in summary")
    args = parser.parse_args()

    outdir = Path(args.output_dir)
    outdir.mkdir(parents=True, exist_ok=True)

    tp_dirs = [d.strip() for d in args.tp_dirs.split(",")]
    assert len(tp_dirs) == 3, f"Expected 3 timepoint dirs, got {len(tp_dirs)}"
    TP_NAMES = ["0h", "6h", "24h"]

    print("=" * 70)
    print("STEP 1: Extract Latent Features & Rank by Importance")
    print("=" * 70)
    print(f"Architecture: {args.arch_file}")
    print(f"Model dir:    {args.model_dir}")
    print(f"Timepoints:   {tp_dirs}")
    print(f"Labels:       {args.labels_csv}")
    print(f"Output:       {outdir}")
    print(f"Latent dim:   {args.latent_dim}")
    print(f"Top K:        {args.top_k}")
    print(f"Device:       {DEVICE}")

    # -------------------------------------------------------------------------
    # Load labels
    # -------------------------------------------------------------------------
    codes_csv, y, groups = load_labels(args.labels_csv)
    N = len(codes_csv)
    n_als = (y == 1).sum()
    n_ctrl = (y == 0).sum()
    print(f"\nLabels: {N} samples ({n_als} ALS, {n_ctrl} CTRL)")

    # -------------------------------------------------------------------------
    # Find common codes across all timepoint directories
    # -------------------------------------------------------------------------
    codes_per_tp = [set(list_sample_codes(d)) for d in tp_dirs]
    common_codes = codes_per_tp[0] & codes_per_tp[1] & codes_per_tp[2]
    codes = [c for c in codes_csv if c in common_codes]

    if len(codes) < N:
        missing = [c for c in codes_csv if c not in common_codes]
        print(f"[WARN] {len(missing)} codes missing from timepoint dirs: {missing}")

    # Reindex y and groups to match available codes
    keep = np.array([c in common_codes for c in codes_csv])
    y = y[keep]
    groups = groups[keep]
    N = len(codes)
    print(f"Using {N} samples with data across all timepoints")

    # -------------------------------------------------------------------------
    # Load EEM data
    # -------------------------------------------------------------------------
    print("\nLoading EEM spectral data...")
    X_np = load_all_samples(codes, tp_dirs)
    _, T, H, W = X_np.shape
    print(f"Data shape: [{N}, {T}, {H}, {W}]")

    X = torch.tensor(X_np, dtype=torch.float32, device=DEVICE)

    # -------------------------------------------------------------------------
    # Import architecture and load models
    # -------------------------------------------------------------------------
    print("\nImporting architecture...")
    arch_mod = import_arch_module(args.arch_file)

    fold_paths = sorted(Path(args.model_dir).glob("best_model_fold*.pth"))
    if len(fold_paths) == 0:
        # Try alternate naming
        fold_paths = sorted(Path(args.model_dir).glob("best_model_fold_*.pth"))
    if len(fold_paths) == 0:
        raise FileNotFoundError(f"No best_model_fold*.pth found in {args.model_dir}")

    print(f"\nFound {len(fold_paths)} fold models:")
    for fp in fold_paths:
        print(f"  {fp.name}")

    models = []
    for fp in fold_paths:
        print(f"\n  Loading {fp.name}...")
        model = arch_mod.ConvAutoencoderWithAttention(
            in_channels=1, latent_dim=args.latent_dim, num_classes=2
        ).to(DEVICE)
        # Initialize dynamic layers
        with torch.no_grad():
            dummy = torch.zeros((1, 3, 1, H, W), device=DEVICE)
            _ = model(dummy)
        _load_state_dict_flexible(model, str(fp))
        model.eval()
        models.append(model)

    M = len(models)
    print(f"\nLoaded {M} models successfully")

    # -------------------------------------------------------------------------
    # Extract features and compute importance for each model
    # -------------------------------------------------------------------------
    print("\n" + "=" * 70)
    print("Extracting latent features and computing importance...")
    print("=" * 70)

    all_z_agg = []          # [M] list of [N, latent_dim]
    all_z_per_tp = []       # [M] list of [N, 3, latent_dim]
    all_attn_weights = []   # [M] list of [N, 3]
    all_importance_dfs = [] # [M] list of DataFrames

    for m_idx, model in enumerate(models):
        fold_name = fold_paths[m_idx].stem
        print(f"\n--- Model {m_idx + 1}/{M}: {fold_name} ---")

        # Extract per-timepoint latents
        z_per_tp = extract_latents_per_timepoint(model, X)  # [N, 3, latent_dim]
        all_z_per_tp.append(z_per_tp)

        # Extract z_agg and attention
        z_agg, attn = compute_z_agg(model, X)  # [N, D], [N, 3]
        all_z_agg.append(z_agg)
        all_attn_weights.append(attn)

        # Attention summary
        mean_attn = attn.mean(axis=0)
        print(f"  Attention weights (mean): 0h={mean_attn[0]:.3f}, 6h={mean_attn[1]:.3f}, 24h={mean_attn[2]:.3f}")

        # Latent statistics
        z_std = z_agg.std(axis=0)
        n_active = (z_std > 0.01).sum()
        print(f"  Active dims (std > 0.01): {n_active}/{args.latent_dim}")

        # Compute importance
        imp_df = compute_all_importances(model, z_agg, y, args.latent_dim)
        all_importance_dfs.append(imp_df)

        # Top K for this model
        top_dims = imp_df.head(args.top_k)["dim"].values
        top_aucs = imp_df.head(args.top_k)["univariate_auc"].values
        top_heads = imp_df.head(args.top_k)["head_weight_l2"].values
        print(f"  Top {args.top_k} dims: {top_dims.tolist()}")
        print(f"  Their AUCs:  {[f'{a:.3f}' for a in top_aucs]}")

        # Save per-model outputs
        model_dir = outdir / fold_name
        model_dir.mkdir(exist_ok=True)

        # z_agg
        z_agg_df = pd.DataFrame(z_agg, columns=[f"dim{d}" for d in range(args.latent_dim)])
        z_agg_df.insert(0, "code", codes)
        z_agg_df.insert(1, "group", groups)
        z_agg_df.to_csv(model_dir / "z_agg.csv", index=False)

        # Per-timepoint latents
        for t_idx, tp_name in enumerate(TP_NAMES):
            z_tp_df = pd.DataFrame(z_per_tp[:, t_idx, :],
                                   columns=[f"dim{d}" for d in range(args.latent_dim)])
            z_tp_df.insert(0, "code", codes)
            z_tp_df.insert(1, "group", groups)
            z_tp_df.to_csv(model_dir / f"z_{tp_name}.csv", index=False)

        # Attention weights
        attn_df = pd.DataFrame(attn, columns=TP_NAMES)
        attn_df.insert(0, "code", codes)
        attn_df.to_csv(model_dir / "attention_weights.csv", index=False)

        # Importance table
        imp_df.to_csv(model_dir / "dim_importance.csv", index=False)

    # -------------------------------------------------------------------------
    # Ensemble: sign-align and average
    # -------------------------------------------------------------------------
    print("\n" + "=" * 70)
    print("Computing ensemble (sign-aligned) latent features...")
    print("=" * 70)

    Z_ensemble = ensemble_sign_align_and_average(all_z_agg)

    # Ensemble importance
    ensemble_imp = consensus_importance(all_importance_dfs, args.latent_dim)

    # Compute additional ensemble-level statistics
    ensemble_auc = importance_univariate_auc(Z_ensemble, y)
    ensemble_t = importance_welch_t(Z_ensemble, y)
    ensemble_rho = importance_spearman(Z_ensemble, y)

    ensemble_imp_extra = pd.DataFrame({
        "dim": np.arange(args.latent_dim),
        "ensemble_auc": ensemble_auc,
        "ensemble_welch_t": ensemble_t,
        "ensemble_spearman_rho": ensemble_rho,
    })
    ensemble_imp = ensemble_imp.merge(ensemble_imp_extra, on="dim")
    ensemble_imp = ensemble_imp.sort_values("consensus_rank")

    # Save ensemble outputs
    z_ens_df = pd.DataFrame(Z_ensemble, columns=[f"dim{d}" for d in range(args.latent_dim)])
    z_ens_df.insert(0, "code", codes)
    z_ens_df.insert(1, "group", groups)
    z_ens_df.to_csv(outdir / "ensemble_z_agg.csv", index=False)

    ensemble_imp.to_csv(outdir / "ensemble_dim_importance.csv", index=False)

    # -------------------------------------------------------------------------
    # Summary: per-timepoint latent features (ensemble averaged)
    # -------------------------------------------------------------------------
    print("\nComputing ensemble per-timepoint latents (sign-aligned)...")
    Z_ref_per_tp = all_z_per_tp[0]  # [N, 3, D]
    aligned_per_tp = [Z_ref_per_tp.copy()]
    for z_tp in all_z_per_tp[1:]:
        z2 = z_tp.copy()
        for d in range(args.latent_dim):
            for t in range(3):
                a = Z_ref_per_tp[:, t, d]
                b = z2[:, t, d]
                if np.std(a) > 1e-12 and np.std(b) > 1e-12:
                    if np.corrcoef(a, b)[0, 1] < 0:
                        z2[:, t, d] *= -1
        aligned_per_tp.append(z2)
    Z_ens_per_tp = np.mean(np.stack(aligned_per_tp), axis=0)  # [N, 3, D]

    for t_idx, tp_name in enumerate(TP_NAMES):
        z_tp_df = pd.DataFrame(
            Z_ens_per_tp[:, t_idx, :],
            columns=[f"dim{d}" for d in range(args.latent_dim)]
        )
        z_tp_df.insert(0, "code", codes)
        z_tp_df.insert(1, "group", groups)
        z_tp_df.to_csv(outdir / f"ensemble_z_{tp_name}.csv", index=False)

    # -------------------------------------------------------------------------
    # Print summary
    # -------------------------------------------------------------------------
    print("\n" + "=" * 70)
    print(f"ENSEMBLE CONSENSUS: Top {args.top_k} Latent Dimensions")
    print("=" * 70)

    top = ensemble_imp.head(args.top_k)
    print(f"\n{'Rank':<6}{'Dim':<8}{'Head_L2':<12}{'AUC':<10}{'Welch_t':<12}{'Spearman':<12}{'Ens_AUC':<10}")
    print("-" * 70)
    for _, row in top.iterrows():
        print(f"{int(row['consensus_rank']):<6}"
              f"{int(row['dim']):<8}"
              f"{row['head_weight_l2_mean']:.4f}    "
              f"{row['univariate_auc_mean']:.3f}   "
              f"{row['welch_t_mean']:.3f}      "
              f"{row['spearman_abs_rho_mean']:.3f}      "
              f"{row['ensemble_auc']:.3f}")

    # Attention weights summary across models
    print(f"\nAttention Weights (mean across all models and samples):")
    all_attn_stacked = np.stack(all_attn_weights)  # [M, N, 3]
    grand_mean_attn = all_attn_stacked.mean(axis=(0, 1))
    grand_std_attn = all_attn_stacked.mean(axis=1).std(axis=0)
    for t_idx, tp_name in enumerate(TP_NAMES):
        print(f"  {tp_name}: {grand_mean_attn[t_idx]:.3f} ± {grand_std_attn[t_idx]:.3f}")

    # Cross-model agreement
    print(f"\nCross-model agreement (top {args.top_k} overlap):")
    per_model_top = [set(df.head(args.top_k)["dim"].values) for df in all_importance_dfs]
    for i in range(M):
        for j in range(i + 1, M):
            overlap = len(per_model_top[i] & per_model_top[j])
            print(f"  Fold {i+1} vs Fold {j+1}: {overlap}/{args.top_k} shared")

    union_all = set()
    for s in per_model_top:
        union_all |= s
    intersect_all = per_model_top[0]
    for s in per_model_top[1:]:
        intersect_all &= s
    print(f"  Union of all top-{args.top_k}: {len(union_all)} unique dims")
    print(f"  Intersection of all top-{args.top_k}: {len(intersect_all)} dims: {sorted(intersect_all)}")

    # -------------------------------------------------------------------------
    # Save manifest
    # -------------------------------------------------------------------------
    manifest = {
        "n_samples": N,
        "n_als": int(n_als),
        "n_ctrl": int(n_ctrl),
        "n_models": M,
        "latent_dim": args.latent_dim,
        "top_k": args.top_k,
        "data_shape": [int(x) for x in X_np.shape],
        "top_dims_consensus": ensemble_imp.head(args.top_k)["dim"].astype(int).tolist(),
        "top_dims_ensemble_auc": [
            float(ensemble_imp.head(args.top_k)["ensemble_auc"].values[i])
            for i in range(args.top_k)
        ],
        "attention_weights_mean": {
            tp: float(grand_mean_attn[t]) for t, tp in enumerate(TP_NAMES)
        },
        "cross_model_intersection": sorted([int(d) for d in intersect_all]),
        "output_files": [
            "ensemble_z_agg.csv",
            "ensemble_z_0h.csv",
            "ensemble_z_6h.csv",
            "ensemble_z_24h.csv",
            "ensemble_dim_importance.csv",
        ] + [f"{fp.stem}/z_agg.csv" for fp in fold_paths]
          + [f"{fp.stem}/dim_importance.csv" for fp in fold_paths],
    }

    import json
    with open(outdir / "step1_manifest.json", "w") as f:
        json.dump(manifest, f, indent=2)

    print(f"\n{'=' * 70}")
    print("All outputs saved to:", outdir)
    print(f"{'=' * 70}")
    print("\nFiles:")
    print("  ensemble_z_agg.csv           - Ensemble-averaged aggregated latents [N x D]")
    print("  ensemble_z_0h/6h/24h.csv     - Ensemble per-timepoint latents [N x D]")
    print("  ensemble_dim_importance.csv   - Consensus importance ranking across all models")
    print("  step1_manifest.json           - Summary metadata")
    for fp in fold_paths:
        print(f"  {fp.stem}/                    - Per-model latents, importance, attention")

    print("\nStep 1 complete. Use the top_dims from ensemble_dim_importance.csv for Step 2.")


if __name__ == "__main__":
    main()
