#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
apply_proxy_latents_to_new_dataset.py

Compute proxy latents for a NEW dataset using exported parameters from the
"shallow conv + ridge" proxy distillation.

Inputs:
- deterministic_feature_spec.json  (contains per-dimension ROI, input scaler, ridge coef/intercept, etc.)
- optional proxy_latent_scalers.json (optional extra output standardization)
- a dataset folder with timepoint subfolders (e.g., out_0h, out_6h, out_24h)
  each containing per-subject Excel files: <subject>.xlsx

Outputs:
- proxy_latents.csv with columns: subject, (optional) group, proxy_dim_<d1> ... proxy_dim_<dk>

This script does NOT train anything.
"""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
from typing import Dict, List, Tuple, Optional

import numpy as np
import pandas as pd
from openpyxl import load_workbook


# -----------------------------
# Fixed kernels (must match training script)
# -----------------------------
def _kernels() -> Dict[str, np.ndarray]:
    sobel_x = np.array([[ 1, 0,-1],
                        [ 2, 0,-2],
                        [ 1, 0,-1]], dtype=float)
    sobel_y = np.array([[ 1, 2, 1],
                        [ 0, 0, 0],
                        [-1,-2,-1]], dtype=float)
    laplace_4 = np.array([[ 0, 1, 0],
                          [ 1,-4, 1],
                          [ 0, 1, 0]], dtype=float)
    diag_45 = np.array([[ 2, 1, 0],
                        [ 1, 0,-1],
                        [ 0,-1,-2]], dtype=float)
    diag_135 = np.array([[ 0, 1, 2],
                         [-1, 0, 1],
                         [-2,-1, 0]], dtype=float)
    blur = np.array([[1,2,1],
                     [2,4,2],
                     [1,2,1]], dtype=float) / 16.0
    return {
        "sobel_x": sobel_x,
        "sobel_y": sobel_y,
        "laplace4": laplace_4,
        "diag_45": diag_45,
        "diag_135": diag_135,
        "blur": blur,
    }


STAT_NAMES = ("mean", "std", "min", "max", "p05", "p50", "p95", "mean_abs", "energy")


def _conv2_same_symm(img: np.ndarray, k: np.ndarray) -> np.ndarray:
    """
    2D convolution, output same size, symmetric boundary.
    Uses scipy if available; otherwise a safe numpy fallback.
    """
    try:
        from scipy.signal import convolve2d  # type: ignore
        return convolve2d(img, k, mode="same", boundary="symm")
    except Exception:
        kh, kw = k.shape
        ph, pw = kh // 2, kw // 2
        padded = np.pad(img, ((ph, ph), (pw, pw)), mode="symmetric")
        out = np.zeros_like(img, dtype=float)
        kf = np.flipud(np.fliplr(k))
        for i in range(out.shape[0]):
            for j in range(out.shape[1]):
                patch = padded[i:i+kh, j:j+kw]
                out[i, j] = float(np.sum(patch * kf))
        return out


def _summ_stats(x: np.ndarray) -> Dict[str, float]:
    x = np.asarray(x, dtype=float)
    # robust percentiles even for small arrays
    return {
        "mean": float(np.mean(x)),
        "std": float(np.std(x)),
        "min": float(np.min(x)),
        "max": float(np.max(x)),
        "p05": float(np.percentile(x, 5)),
        "p50": float(np.percentile(x, 50)),
        "p95": float(np.percentile(x, 95)),
        "mean_abs": float(np.mean(np.abs(x))),
        "energy": float(np.mean(x * x)),
    }


def _parse_float(v, allow_comma_decimal: bool) -> float:
    if v is None:
        return 0.0
    if isinstance(v, (int, float, np.integer, np.floating)):
        if np.isnan(v):
            return 0.0
        return float(v)
    # strings
    s = str(v).strip()
    if s == "":
        return 0.0
    if allow_comma_decimal and "," in s and "." not in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except Exception:
        # treat non-numeric as zero
        return 0.0


def read_eem_excel(filepath: Path, allow_comma_decimal: bool = False) -> np.ndarray:
    """
    Reads an EEM Excel file in the format used in your projects:
    - first row: excitation wavelengths (header)
    - first col: emission identifiers (header)
    - numeric grid starts at (row=2, col=2) in Excel 1-indexed terms.
    Returns a 2D numpy array (H x W) of floats.
    """
    wb = load_workbook(str(filepath), data_only=True, read_only=True)
    ws = wb.active

    max_row = ws.max_row
    max_col = ws.max_column
    if max_row < 2 or max_col < 2:
        raise ValueError(f"File {filepath} is too small: rows={max_row} cols={max_col}")

    # Extract numeric grid excluding header row/col
    H = max_row - 1
    W = max_col - 1
    mat = np.zeros((H, W), dtype=float)

    # rows 2..max_row, cols 2..max_col
    for i, r in enumerate(range(2, max_row + 1)):
        row_cells = ws[r]
        # ws[r] returns all columns; slice 1: (skip first header col)
        for j, cell in enumerate(row_cells[1:max_col]):
            mat[i, j] = _parse_float(cell.value, allow_comma_decimal=allow_comma_decimal)

    return mat


def _load_labels(labels_csv: Optional[str]) -> Tuple[Optional[pd.DataFrame], Optional[str], Optional[str]]:
    """
    Returns (df, subject_col, group_col)
    subject col auto-detected among common names.
    group col auto-detected among common names.
    """
    if not labels_csv:
        return None, None, None

    df = pd.read_csv(labels_csv)
    cols = [c.lower() for c in df.columns]

    subj_candidates = ["subject", "sample", "id", "code", "filename"]
    grp_candidates = ["group", "label", "class", "condition"]

    subject_col = None
    group_col = None

    for cand in subj_candidates:
        if cand in cols:
            subject_col = df.columns[cols.index(cand)]
            break
    if subject_col is None:
        # fallback: first column
        subject_col = df.columns[0]

    for cand in grp_candidates:
        if cand in cols:
            group_col = df.columns[cols.index(cand)]
            break

    return df, subject_col, group_col


def _infer_subjects_from_dir(tp_dir: Path) -> List[str]:
    subjects = []
    for p in sorted(tp_dir.glob("*.xlsx")):
        subjects.append(p.stem)
    return subjects


def _build_mean_over_time(subject: str, data_dir: Path, tp_dirs: List[str], allow_comma_decimal: bool) -> np.ndarray:
    mats = []
    for tp in tp_dirs:
        f = data_dir / tp / f"{subject}.xlsx"
        if not f.exists():
            raise FileNotFoundError(f"Missing file for subject={subject}: {f}")
        mats.append(read_eem_excel(f, allow_comma_decimal=allow_comma_decimal))
    # ensure same shape
    shapes = {m.shape for m in mats}
    if len(shapes) != 1:
        raise ValueError(f"Shape mismatch across timepoints for {subject}: {shapes}")
    return np.mean(np.stack(mats, axis=0), axis=0)


def _extract_features_for_dim(mean_mat: np.ndarray, dim_spec: dict, kernels: Dict[str, np.ndarray]) -> np.ndarray:
    """
    Uses dim_spec["roi"] and dim_spec["features"] to build the feature vector in the exact order.
    Feature names follow the pattern:
      <kernel_name>:<stat>
    where kernel_name is one of: sobel_x, sobel_y, laplace4, diag_45, diag_135, blur, raw.
    stats include: mean,std,min,max,p05,p50,p95,mean_abs,energy
    """
    roi = dim_spec["roi"]
    # ROI is inclusive bounds according to spec
    y0, y1 = int(roi["y0"]), int(roi["y1"])
    x0, x1 = int(roi["x0"]), int(roi["x1"])

    patch = mean_mat[y0:y1+1, x0:x1+1].astype(float, copy=False)
    if patch.size == 0:
        raise ValueError(f"Empty ROI after slicing: roi={roi}, mat_shape={mean_mat.shape}")

    # compute all kernel responses once
    responses: Dict[str, np.ndarray] = {}
    responses["raw"] = patch
    for kname, k in kernels.items():
        responses[kname] = _conv2_same_symm(patch, k)

    # precompute stats for each response
    resp_stats: Dict[Tuple[str, str], float] = {}
    for kname, arr in responses.items():
        s = _summ_stats(arr.ravel())
        for stat_name, val in s.items():
            resp_stats[(kname, stat_name)] = val

    feats = []
    for fname in dim_spec["features"]:
        kname, stat = fname.split(":")
        if (kname, stat) not in resp_stats:
            raise KeyError(f"Missing feature {fname}. Available kernels={list(responses.keys())}, stats={list(STAT_NAMES)}")
        feats.append(resp_stats[(kname, stat)])
    return np.asarray(feats, dtype=float)


def _standardize(x: np.ndarray, mean: np.ndarray, scale: np.ndarray) -> np.ndarray:
    mean = np.asarray(mean, dtype=float)
    scale = np.asarray(scale, dtype=float)
    # avoid divide by zero
    scale_safe = np.where(scale == 0.0, 1.0, scale)
    return (x - mean) / scale_safe


def main():
    ap = argparse.ArgumentParser(description="Apply exported proxy (shallow conv + ridge) to a new dataset.")
    ap.add_argument("--feature_spec_json", required=True, help="Path to deterministic_feature_spec.json")
    ap.add_argument("--proxy_scalers_json", default=None, help="Optional path to proxy_latent_scalers.json (output standardization).")
    ap.add_argument("--data_dir", required=True, help="Root directory containing timepoint subfolders.")
    ap.add_argument("--tp_dirs", default=None, help="Comma-separated timepoint directory names. If omitted, uses tp_dirs from spec.")
    ap.add_argument("--labels_csv", default=None, help="Optional labels CSV to define subject list and include group column.")
    ap.add_argument("--out_csv", required=True, help="Where to write proxy_latents.csv")
    ap.add_argument("--allow_comma_decimal", action="store_true", help="Treat comma as decimal separator when parsing strings.")
    ap.add_argument("--apply_output_scaler", action="store_true", help="Apply per-dimension output scaling (recommended).")
    ap.add_argument("--strict", action="store_true", help="Fail on any missing subject file; otherwise skip with warning.")
    args = ap.parse_args()

    spec = json.load(open(args.feature_spec_json, "r", encoding="utf-8"))
    data_dir = Path(args.data_dir)

    tp_dirs = spec.get("tp_dirs", [])
    if args.tp_dirs:
        tp_dirs = [t.strip() for t in args.tp_dirs.split(",") if t.strip()]

    if not tp_dirs:
        raise ValueError("No tp_dirs found (provide --tp_dirs or ensure spec has tp_dirs).")

    for tp in tp_dirs:
        d = data_dir / tp
        if not d.exists():
            raise FileNotFoundError(f"Timepoint directory not found: {d}")

    # output scalers (optional)
    out_scalers = None
    if args.proxy_scalers_json:
        out_scalers = json.load(open(args.proxy_scalers_json, "r", encoding="utf-8"))

    # subject list
    labels_df, subj_col, grp_col = _load_labels(args.labels_csv)
    if labels_df is not None:
        subjects = [str(s) for s in labels_df[subj_col].tolist()]
        group_map = {}
        if grp_col is not None:
            for _, row in labels_df.iterrows():
                group_map[str(row[subj_col])] = row[grp_col]
    else:
        subjects = _infer_subjects_from_dir(data_dir / tp_dirs[0])
        group_map = {}

    kernels = _kernels()

    # dims to compute
    dims_spec = spec["dims"]
    dims_list = [int(d["dim"]) for d in dims_spec]
    out_cols = [f"proxy_dim_{d}" for d in dims_list]

    rows = []
    skipped = 0

    for subj in subjects:
        try:
            mean_mat = _build_mean_over_time(subj, data_dir, tp_dirs, allow_comma_decimal=args.allow_comma_decimal)
            out_row = {"subject": subj}
            if subj in group_map:
                out_row["group"] = group_map[subj]

            for dspec in dims_spec:
                dim = int(dspec["dim"])
                # feature extraction
                x = _extract_features_for_dim(mean_mat, dspec, kernels)
                # input standardization (per-dim)
                x_std = _standardize(x, np.array(dspec["input_scaler_mean"]), np.array(dspec["input_scaler_scale"]))
                # ridge prediction
                coef = np.asarray(dspec["ridge_coef"], dtype=float)
                intercept = float(dspec["ridge_intercept"])
                yhat = float(np.dot(coef, x_std) + intercept)

                # output scaling
                if args.apply_output_scaler:
                    # Prefer scaler embedded in spec for that dim (most faithful)
                    if isinstance(dspec.get("output_scaler", None), dict) and "mean" in dspec["output_scaler"] and "scale" in dspec["output_scaler"]:
                        mu = float(dspec["output_scaler"]["mean"])
                        sd = float(dspec["output_scaler"]["scale"]) if float(dspec["output_scaler"]["scale"]) != 0.0 else 1.0
                        yhat = (yhat - mu) / sd
                    elif out_scalers is not None:
                        key = f"proxy_dim_{dim}"
                        if key in out_scalers:
                            mu = float(out_scalers[key].get("mean", 0.0))
                            sd = float(out_scalers[key].get("std", 1.0)) if float(out_scalers[key].get("std", 1.0)) != 0.0 else 1.0
                            yhat = (yhat - mu) / sd

                out_row[f"proxy_dim_{dim}"] = yhat

            rows.append(out_row)
        except Exception as e:
            skipped += 1
            msg = f"[WARN] skipping subject={subj}: {e}"
            print(msg)
            if args.strict:
                raise

    df_out = pd.DataFrame(rows)
    # ensure consistent column order
    base_cols = ["subject"] + (["group"] if "group" in df_out.columns else [])
    df_out = df_out[base_cols + out_cols]

    out_path = Path(args.out_csv)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    df_out.to_csv(out_path, index=False)

    print(f"Wrote: {out_path}")
    if skipped:
        print(f"Skipped {skipped} subject(s). Use --strict to fail instead.")


if __name__ == "__main__":
    main()
