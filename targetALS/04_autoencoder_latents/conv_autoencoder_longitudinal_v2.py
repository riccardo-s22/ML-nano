"""
Convolutional Autoencoder with Longitudinal LSTM (Trajectory Analysis)
Version: v3 (Fixes 'Zero Data' bug)

Key Changes:
1. Robust Data Loading: Forces string-to-float conversion for Excel data.
2. Longitudinal LSTM: Models the 0h->6h->24h evolution trajectory.
3. Global Normalization: Preserves intensity differences (quenching).
"""

import torch
import torch.nn as nn
import torch.optim as optim
from torch.utils.data import Dataset, DataLoader
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from sklearn.model_selection import StratifiedKFold
from sklearn.preprocessing import LabelEncoder
import os
from typing import Tuple, List, Dict

# Set random seeds for reproducibility
torch.manual_seed(42)
np.random.seed(42)

class ExcelImageDataset(Dataset):
    """
    Optimized dataset that pre-loads all Excel files into RAM.
    Includes ROBUST parsing to handle numbers stored as text.
    """
    
    def __init__(self, codes: List[str], labels: List[int], 
                 timepoint_dirs: List[str], global_min=None, global_max=None):
        self.codes = codes
        self.labels = labels
        self.timepoint_dirs = timepoint_dirs
        
        # Pre-load data into memory immediately
        print(f"Pre-loading {len(codes)} subjects x {len(timepoint_dirs)} timepoints...")
        self.data_cache = self._preload_all_data()
        
        # Normalization parameters
        self.min_val = global_min
        self.max_val = global_max

    def _preload_all_data(self) -> Dict[str, List[np.ndarray]]:
        """Load all Excel files once and store in a dictionary."""
        cache = {}
        for code in self.codes:
            subject_timepoints = []
            for tp_dir in self.timepoint_dirs:
                filepath = os.path.join(tp_dir, f"{code}.xlsx")
                img = self._load_excel_file(filepath)
                subject_timepoints.append(img)
            cache[code] = subject_timepoints
        return cache
    
    def _load_excel_file(self, filepath: str) -> np.ndarray:
        """
        Parses Excel file with FORCED float conversion.
        Fixes the bug where string-formatted numbers were read as 0.0.
        """
        try:
            wb = load_workbook(filepath, data_only=True)
            ws = wb.active
            data = []
            
            # Counter to verify we are actually reading data
            valid_cells = 0
            
            for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
                if row_idx == 0: continue # Skip header row
                
                row_data = []
                for col_idx, cell in enumerate(row):
                    if col_idx == 0: continue # Skip first column (wavelengths)
                    
                    if cell is not None:
                        try:
                            # FORCE conversion to float. 
                            # Handles "123.45" (string) and 123.45 (float)
                            val = float(cell)
                            row_data.append(val)
                            valid_cells += 1
                        except ValueError:
                            # Cell contained text that isn't a number
                            row_data.append(0.0)
                    else:
                        row_data.append(0.0)
                
                if len(row_data) > 0:
                    data.append(row_data)
            
            # Warn if file seemed empty (helps debugging)
            if valid_cells == 0:
                print(f"WARNING: No numeric data found in {os.path.basename(filepath)}")
                return np.zeros((10, 10), dtype=np.float32)
                
            return np.array(data, dtype=np.float32)
            
        except Exception as e:
            print(f"Error loading {filepath}: {e}")
            return np.zeros((10, 10), dtype=np.float32)

    def compute_stats(self):
        """Calculate min/max across this specific dataset."""
        all_values = []
        for code in self.codes:
            for img in self.data_cache[code]:
                all_values.append(img)
        
        full_stack = np.array(all_values)
        if full_stack.size == 0: return 0.0, 1.0
        return float(full_stack.min()), float(full_stack.max())

    def set_normalization(self, min_val, max_val):
        """Apply external normalization stats."""
        self.min_val = min_val
        self.max_val = max_val

    def __len__(self):
        return len(self.codes)
    
    def __getitem__(self, idx):
        code = self.codes[idx]
        label = self.labels[idx]
        
        raw_images = self.data_cache[code] # List of 3 numpy arrays
        processed_tensors = []
        
        for img in raw_images:
            # Global Normalization
            if self.min_val is not None and self.max_val is not None:
                denom = self.max_val - self.min_val
                if denom == 0: denom = 1.0
                norm_img = (img - self.min_val) / denom
            else:
                norm_img = img # Fallback (should not happen)

            # Add channel dim: (H, W) -> (1, H, W)
            if len(norm_img.shape) == 2:
                norm_img = norm_img[np.newaxis, :, :]
                
            processed_tensors.append(torch.FloatTensor(norm_img))
        
        # Stack timepoints: (3, 1, H, W)
        images = torch.stack(processed_tensors)
        return images, torch.tensor(label, dtype=torch.long)


class ConvEncoder(nn.Module):
    """Shared Convolutional Encoder."""
    def __init__(self, in_channels: int, latent_dim: int):
        super(ConvEncoder, self).__init__()
        self.conv1 = nn.Sequential(
            nn.Conv2d(in_channels, 16, 3, padding=1), nn.BatchNorm2d(16), nn.ReLU(),
            nn.MaxPool2d(2, 2)
        )
        self.conv2 = nn.Sequential(
            nn.Conv2d(16, 32, 3, padding=1), nn.BatchNorm2d(32), nn.ReLU(),
            nn.MaxPool2d(2, 2)
        )
        self.conv3 = nn.Sequential(
            nn.Conv2d(32, 64, 3, padding=1), nn.BatchNorm2d(64), nn.ReLU(),
            nn.MaxPool2d(2, 2)
        )
        self.fc = None
        self.latent_dim = latent_dim
        self.encoded_shape = None
        self.input_shape = None
        
    def forward(self, x):
        if self.input_shape is None: self.input_shape = (x.shape[2], x.shape[3])
        x = self.conv1(x)
        x = self.conv2(x)
        x = self.conv3(x)
        if self.encoded_shape is None: self.encoded_shape = x.shape[1:]
        if self.fc is None:
            flatten_size = x.shape[1] * x.shape[2] * x.shape[3]
            self.fc = nn.Linear(flatten_size, self.latent_dim).to(x.device)
        x = x.view(x.size(0), -1)
        return self.fc(x)

class ConvDecoder(nn.Module):
    """Mirror Decoder."""
    def __init__(self, latent_dim, out_channels, encoded_shape, target_shape):
        super(ConvDecoder, self).__init__()
        self.encoded_c, self.encoded_h, self.encoded_w = encoded_shape
        self.target_h, self.target_w = target_shape
        self.fc = nn.Linear(latent_dim, self.encoded_c * self.encoded_h * self.encoded_w)
        self.deconv3 = nn.Sequential(
            nn.ConvTranspose2d(64, 32, 3, stride=2, padding=1, output_padding=1),
            nn.BatchNorm2d(32), nn.ReLU()
        )
        self.deconv2 = nn.Sequential(
            nn.ConvTranspose2d(32, 16, 3, stride=2, padding=1, output_padding=1),
            nn.BatchNorm2d(16), nn.ReLU()
        )
        self.deconv1 = nn.Sequential(
            nn.ConvTranspose2d(16, 16, 3, stride=2, padding=1, output_padding=1),
            nn.BatchNorm2d(16), nn.ReLU()
        )
        self.final_conv = nn.Conv2d(16, out_channels, 3, padding=1)
        self.sigmoid = nn.Sigmoid()

    def forward(self, z):
        x = self.fc(z)
        x = x.view(x.size(0), self.encoded_c, self.encoded_h, self.encoded_w)
        x = self.deconv3(x)
        x = self.deconv2(x)
        x = self.deconv1(x)
        x = self.final_conv(x)
        if x.shape[2] != self.target_h or x.shape[3] != self.target_w:
            x = torch.nn.functional.interpolate(x, size=(self.target_h, self.target_w), mode='bilinear')
        return self.sigmoid(x)

class TemporalEvolutionEncoder(nn.Module):
    """
    LSTM that traces the trajectory (0h -> 6h -> 24h).
    """
    def __init__(self, input_dim, hidden_dim=32):
        super(TemporalEvolutionEncoder, self).__init__()
        # Batch_first=True expects input: (batch, seq_len, features)
        self.lstm = nn.LSTM(input_size=input_dim, hidden_size=hidden_dim, num_layers=1, batch_first=True)
    
    def forward(self, x_seq):
        # x_seq: (batch, 3, latent_dim)
        # We only need the final hidden state (hn) which summarizes the path
        _, (hn, _) = self.lstm(x_seq)
        return hn.squeeze(0) # (batch, hidden_dim)

class Classifier(nn.Module):
    def __init__(self, input_dim, num_classes):
        super(Classifier, self).__init__()
        self.net = nn.Sequential(
            nn.Linear(input_dim, 32),
            nn.ReLU(),
            nn.Dropout(0.4), 
            nn.Linear(32, num_classes)
        )
    def forward(self, x):
        return self.net(x)

class AutoencoderWithTrajectory(nn.Module):
    def __init__(self, in_channels, latent_dim, num_classes):
        super(AutoencoderWithTrajectory, self).__init__()
        self.in_channels = in_channels
        self.latent_dim = latent_dim
        
        self.encoder = ConvEncoder(in_channels, latent_dim)
        self.temporal = TemporalEvolutionEncoder(input_dim=latent_dim, hidden_dim=32)
        self.classifier = Classifier(input_dim=32, num_classes=num_classes)
        self.decoder = None

    def forward(self, images):
        batch_size, num_timepoints = images.shape[0], images.shape[1]
        
        # 1. Encode sequence
        latent_list = []
        for t in range(num_timepoints):
            z = self.encoder(images[:, t])
            latent_list.append(z)
            
        if self.decoder is None:
            self.decoder = ConvDecoder(
                self.latent_dim, self.in_channels, 
                self.encoder.encoded_shape, self.encoder.input_shape
            ).to(images.device)
            
        # 2. Trace Trajectory
        latent_seq = torch.stack(latent_list, dim=1) # (batch, 3, latent_dim)
        trajectory = self.temporal(latent_seq)
        
        # 3. Classify Trajectory
        logits = self.classifier(trajectory)
        
        # 4. Reconstruct (Regularization)
        reconstructions = []
        for t in range(num_timepoints):
            rec = self.decoder(latent_list[t])
            reconstructions.append(rec)
        reconstructions = torch.stack(reconstructions, dim=1)
        
        return reconstructions, logits

def train_epoch(model, loader, optimizer, alpha, beta, device):
    model.train()
    total_loss, correct, total = 0, 0, 0
    mse = nn.MSELoss()
    ce = nn.CrossEntropyLoss()
    
    for imgs, lbls in loader:
        imgs, lbls = imgs.to(device), lbls.to(device)
        optimizer.zero_grad()
        recons, logits = model(imgs)
        
        l_rec = mse(recons, imgs)
        l_cls = ce(logits, lbls)
        loss = alpha * l_rec + beta * l_cls
        
        loss.backward()
        optimizer.step()
        
        total_loss += loss.item()
        _, preds = torch.max(logits, 1)
        total += lbls.size(0)
        correct += (preds == lbls).sum().item()
        
    return total_loss/len(loader), 100*correct/total

def validate(model, loader, device, alpha, beta):
    model.eval()
    total_loss, correct, total = 0, 0, 0
    mse = nn.MSELoss()
    ce = nn.CrossEntropyLoss()
    
    with torch.no_grad():
        for imgs, lbls in loader:
            imgs, lbls = imgs.to(device), lbls.to(device)
            recons, logits = model(imgs)
            
            l_rec = mse(recons, imgs)
            l_cls = ce(logits, lbls)
            loss = alpha * l_rec + beta * l_cls
            total_loss += loss.item()
            
            _, preds = torch.max(logits, 1)
            total += lbls.size(0)
            correct += (preds == lbls).sum().item()
            
    if len(loader) == 0: return 0, 0
    return total_loss/len(loader), 100*correct/total

def main():
    # --- CONFIG ---
    LABELS_CSV = 'sample_labels.csv'
    TIMEPOINT_DIRS = ['out_0h', 'out_6h', 'out_24h']
    LATENT_DIM = 64 
    BATCH_SIZE = 4
    EPOCHS = 100
    LR = 0.0005
    N_SPLITS = 5
    ALPHA, BETA = 1.0, 1.5
    
    device = torch.device('cuda' if torch.cuda.is_available() else 'cpu')
    print(f"Using device: {device}")
    
    # --- DATA SETUP ---
    df = pd.read_csv(LABELS_CSV)
    le = LabelEncoder()
    df['enc_group'] = le.fit_transform(df['group'])
    codes, labels = df['code'].astype(str).values, df['enc_group'].values
    
    skf = StratifiedKFold(n_splits=N_SPLITS, shuffle=True, random_state=42)
    fold_results = []
    
    for fold, (train_idx, val_idx) in enumerate(skf.split(codes, labels)):
        print(f"\n{'='*30} FOLD {fold+1}/{N_SPLITS} {'='*30}")
        
        # 1. Init Datasets
        train_ds = ExcelImageDataset(codes[train_idx], labels[train_idx], TIMEPOINT_DIRS)
        g_min, g_max = train_ds.compute_stats()
        train_ds.set_normalization(g_min, g_max)
        
        print(f"  Data Range: {g_min:.2f} to {g_max:.2f}")
        if g_max == 0:
            print("  ERROR: Data is all zeros! Check Excel formatting.")
            continue

        val_ds = ExcelImageDataset(codes[val_idx], labels[val_idx], TIMEPOINT_DIRS)
        val_ds.set_normalization(g_min, g_max)
        
        train_loader = DataLoader(train_ds, batch_size=BATCH_SIZE, shuffle=True)
        val_loader = DataLoader(val_ds, batch_size=BATCH_SIZE, shuffle=False)
        
        # 2. Init Model
        sample_img, _ = train_ds[0]
        _, channels, h, w = sample_img.shape
        model = AutoencoderWithTrajectory(channels, LATENT_DIM, len(le.classes_)).to(device)
        optimizer = optim.Adam(model.parameters(), lr=LR)
        
        # 3. Train
        best_acc = 0.0
        for epoch in range(EPOCHS):
            t_loss, t_acc = train_epoch(model, train_loader, optimizer, ALPHA, BETA, device)
            v_loss, v_acc = validate(model, val_loader, device, ALPHA, BETA)
            
            if v_acc > best_acc:
                best_acc = v_acc
                torch.save(model.state_dict(), f'best_model_fold{fold+1}.pth')
                
            if (epoch+1) % 20 == 0:
                print(f"  Ep {epoch+1}: Train Loss={t_loss:.4f}, Acc={t_acc:.0f}% | Val Loss={v_loss:.4f}, Acc={v_acc:.0f}%")
        
        print(f"  >> Best Val Accuracy: {best_acc:.2f}%")
        fold_results.append(best_acc)
        
    print(f"\nOverall CV Accuracy: {np.mean(fold_results):.2f}% (+/- {np.std(fold_results):.2f}%)")

if __name__ == "__main__":
    main()