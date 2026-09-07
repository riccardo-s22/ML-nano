#!/usr/bin/env python3
"""
STEP 2: Compute Per-Pixel Reconstruction MSE (averaged across 5 models)
========================================================================

For each of the 5 fold models and each timepoint (0h, 6h, 24h):
  1. Forward pass: reconstruct each sample's EEM matrix
  2. Compute per-pixel squared error: (reconstruction - original)^2
  3. Average squared error across all N samples → mean MSE map [H × W]
  4. Average the 5 per-model MSE maps → ensemble MSE map per timepoint

Outputs:
  - Per-model MSE maps: mse_model{k}_{tp}.npy   [H × W]
  - Ensemble MSE maps:  mse_ensemble_{tp}.npy    [H × W]
  - Per-sample MSE:     mse_per_sample.csv       (mean MSE per sample per tp per model)
  - Visualizations:     mse_ensemble_heatmaps.png

USAGE:
------
python step2_compute_mse_maps.py \
    --arch_file conv_autoencoder_detailed.py \
    --model_dir ./5_fold_models_original \
    --tp_dirs ./out_0h,./out_6h,./out_24h \
    --labels_csv sample_labels.csv \
    --output_dir ./step2_mse_maps \
    --latent_dim 512
"""

import os
import sys
import argparse
import importlib.util
from pathlib import Path
from typing import List, Optional
from collections import OrderedDict

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import torch
import torch.nn as nn

import warnings
warnings.filterwarnings("ignore")

SEED = 42
np.random.seed(SEED)
torch.manual_seed(SEED)
DEVICE = torch.device("cuda" if torch.cuda.is_available() else "cpu")

TP_NAMES = ["0h", "6h", "24h"]


# =============================================================================
# Architecture import (same as Step 1)
# =============================================================================

def import_arch_module(arch_file: str):
    spec = importlib.util.spec_from_file_location("arch_module", arch_file)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Could not import architecture file: {arch_file}")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# =============================================================================
# Data loading (same as Step 1)
# =============================================================================

def load_labels(csv_path: str):
    df = pd.read_csv(csv_path)
    codes = df["code"].tolist()
    groups = df["group"].values
    y = (groups == "ALS").astype(int)
    return codes, y, groups


def list_sample_codes(tp_dir: str) -> List[str]:
    codes = []
    for f in Path(tp_dir).glob("*.xlsx"):
        if f.name.startswith("~$"):
            continue
        name = f.stem
        for suffix in ["_0h", "_6h", "_24h"]:
            if name.endswith(suffix):
                name = name[:-len(suffix)]
                break
        codes.append(name)
    return codes


def find_file_for_code(tp_dir: str, code: str) -> Optional[str]:
    """Find the Excel file for a given sample code in a timepoint directory."""
    tp_path = Path(tp_dir)
    
    # 1. Try exact match (e.g. "code.xlsx")
    # This handles your specific case where filenames match codes exactly
    exact_path = tp_path / f"{code}.xlsx"
    if exact_path.exists():
        return str(exact_path)

    # 2. Try flexible suffix matching (e.g. "code_0h.xlsx")
    for f in tp_path.glob("*.xlsx"):
        if f.name.startswith("~$"):
            continue
        stem = f.stem
        
        # Check against suffixes
        for suffix in ["_0h", "_6h", "_24h"]:
            if stem.endswith(suffix):
                # Check if the stem *without* suffix matches the code
                file_code = stem[:-len(suffix)]
                
                # Direct match
                if file_code == code:
                    return str(f)
                
                # Match with dot/underscore replacement
                if file_code.replace('.', '_') == code:
                    return str(f)

    return None

def load_excel_as_array(filepath: str) -> np.ndarray:
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    data = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx == 0:
            continue
        row_data = []
        for col_idx, cell in enumerate(row):
            if col_idx == 0:
                continue
            if isinstance(cell, (int, float)) and cell is not None:
                row_data.append(float(cell))
            else:
                row_data.append(0.0)
        if row_data:
            data.append(row_data)
    arr = np.array(data, dtype=np.float32)
    mn, mx = arr.min(), arr.max()
    if mx > mn:
        arr = (arr - mn) / (mx - mn)
    else:
        arr = np.zeros_like(arr)
    return arr


def load_all_samples(codes: List[str], tp_dirs: List[str]) -> np.ndarray:
    arrays = []
    for code in codes:
        tp_images = []
        for tp_dir in tp_dirs:
            fpath = find_file_for_code(tp_dir, code)
            if fpath is None:
                raise FileNotFoundError(f"No file found for code '{code}' in {tp_dir}")
            arr = load_excel_as_array(fpath)
            tp_images.append(arr)
        shapes = [a.shape for a in tp_images]
        if len(set(shapes)) > 1:
            min_h = min(s[0] for s in shapes)
            min_w = min(s[1] for s in shapes)
            tp_images = [a[:min_h, :min_w] for a in tp_images]
        arrays.append(np.stack(tp_images, axis=0))
    X = np.stack(arrays, axis=0)
    return X


# =============================================================================
# Model loading (same as Step 1)
# =============================================================================

def _load_state_dict_flexible(model: nn.Module, path: str):
    sd = torch.load(path, map_location=DEVICE, weights_only=False)
    if isinstance(sd, dict) and "state_dict" in sd:
        sd = sd["state_dict"]
    try:
        model.load_state_dict(sd, strict=True)
        return
    except RuntimeError:
        pass
    cleaned = {}
    for k, v in sd.items():
        kk = k
        for pref in ["module.", "model."]:
            if kk.startswith(pref):
                kk = kk[len(pref):]
        cleaned[kk] = v
    missing, unexpected = model.load_state_dict(cleaned, strict=False)
    if missing:
        print(f"    [WARN] Missing keys: {missing[:5]}{'...' if len(missing)>5 else ''}")


# =============================================================================
# MSE computation
# =============================================================================

@torch.no_grad()
def compute_sqerr_per_sample_tp(model: nn.Module, X: torch.Tensor) -> np.ndarray:
    """
    Compute per-pixel squared reconstruction error.

    Args:
        model: trained ConvAutoencoderWithAttention
        X: input tensor [N, 3, H, W]

    Returns:
        sqerr: [N, 3, H, W] per-pixel squared error
    """
    model.eval()
    X_5d = X.unsqueeze(2)            # [N, 3, 1, H, W]
    recon_5d, _, _ = model(X_5d)     # [N, 3, 1, H, W]
    sq = (recon_5d[:, :, 0] - X) ** 2  # [N, 3, H, W]
    return sq.detach().cpu().numpy().astype(np.float32)


# =============================================================================
# Main
# =============================================================================

def main():
    parser = argparse.ArgumentParser(
        description="Step 2: Compute per-pixel MSE maps averaged across 5 fold models"
    )
    parser.add_argument("--arch_file", type=str, required=True)
    parser.add_argument("--model_dir", type=str, required=True)
    parser.add_argument("--tp_dirs", type=str, required=True,
                        help="Comma-separated: dir_0h,dir_6h,dir_24h")
    parser.add_argument("--labels_csv", type=str, required=True)
    parser.add_argument("--output_dir", type=str, default="./step2_mse_maps")
    parser.add_argument("--latent_dim", type=int, default=512)
    args = parser.parse_args()

    outdir = Path(args.output_dir)
    outdir.mkdir(parents=True, exist_ok=True)

    tp_dirs = [d.strip() for d in args.tp_dirs.split(",")]
    assert len(tp_dirs) == 3

    print("=" * 70)
    print("STEP 2: Compute Per-Pixel MSE Maps")
    print("=" * 70)

    # -------------------------------------------------------------------------
    # Load labels & data
    # -------------------------------------------------------------------------
    codes_csv, y, groups = load_labels(args.labels_csv)
    codes_per_tp = [set(list_sample_codes(d)) for d in tp_dirs]
    common = codes_per_tp[0] & codes_per_tp[1] & codes_per_tp[2]
    codes = [c for c in codes_csv if c in common]
    keep = np.array([c in common for c in codes_csv])
    y = y[keep]
    groups = groups[keep]
    N = len(codes)
    print(f"Samples: {N} ({(y==1).sum()} ALS, {(y==0).sum()} CTRL)")

    print("\nLoading EEM data...")
    X_np = load_all_samples(codes, tp_dirs)
    _, T, H, W = X_np.shape
    print(f"Shape: [{N}, {T}, {H}, {W}]")

    X = torch.tensor(X_np, dtype=torch.float32, device=DEVICE)

    # -------------------------------------------------------------------------
    # Load models
    # -------------------------------------------------------------------------
    arch_mod = import_arch_module(args.arch_file)

    fold_paths = sorted(Path(args.model_dir).glob("best_model_fold*.pth"))
    if len(fold_paths) == 0:
        fold_paths = sorted(Path(args.model_dir).glob("best_model_fold_*.pth"))
    if len(fold_paths) == 0:
        raise FileNotFoundError(f"No best_model_fold*.pth in {args.model_dir}")

    M = len(fold_paths)
    print(f"\nLoading {M} fold models...")

    models = []
    for fp in fold_paths:
        model = arch_mod.ConvAutoencoderWithAttention(
            in_channels=1, latent_dim=args.latent_dim, num_classes=2
        ).to(DEVICE)
        with torch.no_grad():
            dummy = torch.zeros((1, 3, 1, H, W), device=DEVICE)
            _ = model(dummy)
        _load_state_dict_flexible(model, str(fp))
        model.eval()
        models.append(model)
        print(f"  Loaded: {fp.name}")

    # -------------------------------------------------------------------------
    # Compute MSE for each model
    # -------------------------------------------------------------------------
    print("\n" + "=" * 70)
    print("Computing per-pixel squared error for each model...")
    print("=" * 70)

    # sqerr_all: [M, N, 3, H, W]
    sqerr_all = []
    for m_idx, model in enumerate(models):
        print(f"  Model {m_idx + 1}/{M}: {fold_paths[m_idx].name}...", end=" ")
        sqerr = compute_sqerr_per_sample_tp(model, X)  # [N, 3, H, W]
        sqerr_all.append(sqerr)

        # Per-timepoint summary
        for t in range(3):
            mean_mse = sqerr[:, t].mean()
            print(f"{TP_NAMES[t]}={mean_mse:.6f}", end="  ")
        print()

    sqerr_all = np.stack(sqerr_all, axis=0)  # [M, N, 3, H, W]

    # -------------------------------------------------------------------------
    # Average MSE across samples → per-model mean MSE map [M, 3, H, W]
    # -------------------------------------------------------------------------
    print("\nAveraging across samples (per model, per timepoint)...")
    mse_per_model = sqerr_all.mean(axis=1)  # [M, 3, H, W]

    # Save per-model MSE maps
    for m_idx in range(M):
        for t in range(3):
            fname = f"mse_model{m_idx+1}_{TP_NAMES[t]}.npy"
            np.save(outdir / fname, mse_per_model[m_idx, t])

    # -------------------------------------------------------------------------
    # Average across 5 models → ensemble MSE map [3, H, W]
    # -------------------------------------------------------------------------
    print("Averaging across 5 models → ensemble MSE maps...")
    mse_ensemble = mse_per_model.mean(axis=0)  # [3, H, W]
    mse_ensemble_std = mse_per_model.std(axis=0)  # [3, H, W]

    for t in range(3):
        np.save(outdir / f"mse_ensemble_{TP_NAMES[t]}.npy", mse_ensemble[t])
        np.save(outdir / f"mse_ensemble_std_{TP_NAMES[t]}.npy", mse_ensemble_std[t])

    # -------------------------------------------------------------------------
    # Also save per-sample mean MSE (diagnostic)
    # -------------------------------------------------------------------------
    print("Computing per-sample summary...")
    rows = []
    for i, code in enumerate(codes):
        for m_idx in range(M):
            for t in range(3):
                rows.append({
                    "code": code,
                    "group": groups[i],
                    "model": m_idx + 1,
                    "timepoint": TP_NAMES[t],
                    "mean_mse": float(sqerr_all[m_idx, i, t].mean()),
                    "max_mse": float(sqerr_all[m_idx, i, t].max()),
                    "median_mse": float(np.median(sqerr_all[m_idx, i, t])),
                })
    df_sample_mse = pd.DataFrame(rows)
    df_sample_mse.to_csv(outdir / "mse_per_sample.csv", index=False)

    # Also save the full per-sample ensemble-averaged MSE (averaged across models)
    sqerr_ens = sqerr_all.mean(axis=0)  # [N, 3, H, W] - ensemble avg per sample
    np.save(outdir / "sqerr_ensemble_per_sample.npy", sqerr_ens)

    # -------------------------------------------------------------------------
    # Summary statistics
    # -------------------------------------------------------------------------
    print("\n" + "=" * 70)
    print("ENSEMBLE MSE SUMMARY (averaged over 5 models, then over samples)")
    print("=" * 70)

    for t in range(3):
        m = mse_ensemble[t]
        print(f"\n  {TP_NAMES[t]}:")
        print(f"    Mean MSE:   {m.mean():.6f}")
        print(f"    Median MSE: {np.median(m):.6f}")
        print(f"    Max MSE:    {m.max():.6f}")
        print(f"    Std MSE:    {m.std():.6f}")
        print(f"    % pixels with MSE < 0.01: {(m < 0.01).mean()*100:.1f}%")
        print(f"    % pixels with MSE < 0.05: {(m < 0.05).mean()*100:.1f}%")
        print(f"    % pixels with MSE < 0.10: {(m < 0.10).mean()*100:.1f}%")

    # ALS vs CTRL reconstruction quality
    als_mask = (y == 1)
    ctrl_mask = (y == 0)
    print(f"\n  Per-group mean MSE (ensemble):")
    for t in range(3):
        als_mse = sqerr_ens[als_mask, t].mean()
        ctrl_mse = sqerr_ens[ctrl_mask, t].mean()
        print(f"    {TP_NAMES[t]}: ALS={als_mse:.6f}  CTRL={ctrl_mse:.6f}  ratio={als_mse/ctrl_mse:.3f}")

    # Cross-model consistency
    print(f"\n  Cross-model consistency (correlation of MSE maps):")
    for t in range(3):
        corrs = []
        for i in range(M):
            for j in range(i + 1, M):
                r = np.corrcoef(mse_per_model[i, t].ravel(), mse_per_model[j, t].ravel())[0, 1]
                corrs.append(r)
        print(f"    {TP_NAMES[t]}: mean r = {np.mean(corrs):.4f} (range {np.min(corrs):.4f}-{np.max(corrs):.4f})")

    # -------------------------------------------------------------------------
    # Visualization
    # -------------------------------------------------------------------------
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        fig, axes = plt.subplots(2, 3, figsize=(18, 10))

        # Top row: ensemble mean MSE
        for t in range(3):
            ax = axes[0, t]
            im = ax.imshow(mse_ensemble[t], aspect="auto", cmap="hot", origin="lower")
            ax.set_title(f"Ensemble Mean MSE — {TP_NAMES[t]}", fontsize=12)
            ax.set_xlabel("Excitation index")
            ax.set_ylabel("Emission index")
            plt.colorbar(im, ax=ax, shrink=0.8)

        # Bottom row: ensemble std across models
        for t in range(3):
            ax = axes[1, t]
            im = ax.imshow(mse_ensemble_std[t], aspect="auto", cmap="viridis", origin="lower")
            ax.set_title(f"Cross-model Std — {TP_NAMES[t]}", fontsize=12)
            ax.set_xlabel("Excitation index")
            ax.set_ylabel("Emission index")
            plt.colorbar(im, ax=ax, shrink=0.8)

        plt.suptitle("Step 2: Per-Pixel Reconstruction MSE (ensemble of 5 models)", fontsize=14, y=1.02)
        plt.tight_layout()
        plt.savefig(outdir / "mse_ensemble_heatmaps.png", dpi=150, bbox_inches="tight")
        plt.close()
        print(f"\nSaved: mse_ensemble_heatmaps.png")
    except Exception as e:
        print(f"\n[WARN] Could not generate visualization: {e}")

    # -------------------------------------------------------------------------
    # Save manifest
    # -------------------------------------------------------------------------
    import json
    manifest = {
        "n_samples": N,
        "n_models": M,
        "data_shape": [int(N), int(T), int(H), int(W)],
        "ensemble_mean_mse_per_tp": {
            TP_NAMES[t]: float(mse_ensemble[t].mean()) for t in range(3)
        },
        "output_files": {
            "ensemble_mse_maps": [f"mse_ensemble_{tp}.npy" for tp in TP_NAMES],
            "ensemble_std_maps": [f"mse_ensemble_std_{tp}.npy" for tp in TP_NAMES],
            "per_model_mse_maps": [f"mse_model{m+1}_{tp}.npy" for m in range(M) for tp in TP_NAMES],
            "per_sample_summary": "mse_per_sample.csv",
            "full_per_sample_sqerr": "sqerr_ensemble_per_sample.npy",
            "visualization": "mse_ensemble_heatmaps.png",
        }
    }
    with open(outdir / "step2_manifest.json", "w") as f:
        json.dump(manifest, f, indent=2)

    print(f"\n{'=' * 70}")
    print(f"All outputs saved to: {outdir}")
    print(f"{'=' * 70}")
    print("\nKey files:")
    print("  mse_ensemble_{0h,6h,24h}.npy  — ensemble MSE map per timepoint [H × W]")
    print("  sqerr_ensemble_per_sample.npy  — per-sample ensemble MSE [N × 3 × H × W]")
    print("  mse_per_sample.csv             — per-sample mean/max/median MSE")
    print("  mse_ensemble_heatmaps.png      — visualization")
    print("\nStep 2 complete. These MSE maps feed into Step 3 (ROI selection).")


if __name__ == "__main__":
    main()
