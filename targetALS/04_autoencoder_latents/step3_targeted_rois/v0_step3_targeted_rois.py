#!/usr/bin/env python3
"""
STEP 3 (TARGETED): Pixel Attribution & ROI Extraction
=====================================================

USAGE:
  python step3_targeted_rois.py \
      --tp_dirs ./out_0h,./out_6h,./out_24h \
      --step1_dir ./step1_latent_features \
      --step2_dir ./step2_mse_maps \
      --model_dir ./5_fold_models_original \
      --output_dir ./step3_targeted_rois \
      --mask_percentile 95 \
      --mse_threshold 0.5 \
      --arch_file conv_autoencoder_detailed.py
"""

import os
import sys
import argparse
import importlib.util
from pathlib import Path
from typing import List, Dict, Tuple, Optional

import numpy as np
import pandas as pd
from openpyxl import load_workbook
import torch
import torch.nn as nn
from sklearn.linear_model import Ridge
from sklearn.preprocessing import StandardScaler
from scipy import ndimage

import warnings
warnings.filterwarnings("ignore")

SEED = 42
torch.manual_seed(SEED)
np.random.seed(SEED)

TP_NAMES = ["0h", "6h", "24h"]

# =============================================================================
# CONFIG: The Top 5 Clinically Validated Features
# =============================================================================
TARGET_FEATURES = {
    "best_model_fold1": [103, 179],
    "best_model_fold3": [477, 376, 370]
}

# =============================================================================
# Architecture Loading
# =============================================================================
def import_arch_module(arch_file: str):
    spec = importlib.util.spec_from_file_location("arch_module", arch_file)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Could not import architecture file: {arch_file}")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod

def _load_state_dict_flexible(model: nn.Module, path: str):
    state = torch.load(path, map_location="cpu")
    if "state_dict" in state:
        state = state["state_dict"]
    
    new_state = {}
    for k, v in state.items():
        if k.startswith("module."): k = k[7:]
        elif k.startswith("model."): k = k[6:]
        new_state[k] = v
        
    try:
        model.load_state_dict(new_state, strict=True)
    except Exception as e:
        print(f"[WARN] strict=True failed: {e}. Trying strict=False.")
        model.load_state_dict(new_state, strict=False)

# =============================================================================
# Data Loading Helpers
# =============================================================================
def list_sample_codes(tp_dir: str) -> List[str]:
    codes = []
    for f in Path(tp_dir).glob("*.xlsx"):
        if f.name.startswith("~$"): continue
        name = f.stem
        for suffix in ["_0h", "_6h", "_24h"]:
            if name.endswith(suffix):
                name = name[:-len(suffix)]
                break
        codes.append(name)
    return codes

def find_file_for_code(tp_dir: str, code: str) -> Optional[str]:
    tp_path = Path(tp_dir)
    exact = tp_path / f"{code}.xlsx"
    if exact.exists(): return str(exact)
    for f in tp_path.glob("*.xlsx"):
        if f.name.startswith("~$"): continue
        stem = f.stem
        for suffix in ["_0h", "_6h", "_24h"]:
            if stem.endswith(suffix):
                if stem[:-len(suffix)] == code: return str(f)
                if stem[:-len(suffix)].replace('.', '_') == code: return str(f)
    return None

def load_excel_as_array(filepath: str) -> np.ndarray:
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    data = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx == 0: continue
        row_data = []
        for col_idx, cell in enumerate(row):
            if col_idx == 0: continue
            if isinstance(cell, (int, float)) and cell is not None:
                row_data.append(float(cell))
            else:
                row_data.append(0.0)
        if row_data: data.append(row_data)
    arr = np.array(data, dtype=np.float32)
    mn, mx = arr.min(), arr.max()
    if mx > mn: arr = (arr - mn) / (mx - mn)
    else: arr = np.zeros_like(arr)
    return arr

def load_all_samples(codes: List[str], tp_dirs: List[str]) -> np.ndarray:
    arrays = []
    for code in codes:
        tp_images = []
        for tp_dir in tp_dirs:
            fpath = find_file_for_code(tp_dir, code)
            if fpath is None: raise FileNotFoundError(f"Missing '{code}' in {tp_dir}")
            arr = load_excel_as_array(fpath)
            tp_images.append(arr)
        
        shapes = [a.shape for a in tp_images]
        if len(set(shapes)) > 1:
            min_h = min(s[0] for s in shapes)
            min_w = min(s[1] for s in shapes)
            tp_images = [a[:min_h, :min_w] for a in tp_images]
        arrays.append(np.stack(tp_images, axis=0))
    return np.stack(arrays, axis=0)

# =============================================================================
# Original Importance Functions
# =============================================================================
def compute_gradient_saliency(model: nn.Module, X: torch.Tensor,
                              dims: List[int]) -> Dict[int, Dict[int, np.ndarray]]:
    N, T, H, W = X.shape
    model.eval()

    grad_maps = {d: {t: np.zeros((H, W), dtype=np.float64) for t in range(T)} for d in dims}

    batch_size = min(8, N)
    n_batches = (N + batch_size - 1) // batch_size

    for b in range(n_batches):
        s = b * batch_size
        e = min(s + batch_size, N)
        B = e - s
        X_batch = X[s:e].clone().requires_grad_(True)
        X_5d = X_batch.unsqueeze(2)

        latents = [model.encoder(X_5d[:, t]) for t in range(T)]

        for d in dims:
            for t in range(T):
                target = latents[t][:, d].sum()
                model.zero_grad()
                if X_batch.grad is not None:
                    X_batch.grad.zero_()
                target.backward(retain_graph=True)

                grads = X_batch.grad.detach().abs()  # [B, 3, H, W]
                # Same gradient indexing as the original
                grad_maps[d][t] += grads[:, t].mean(dim=0).cpu().numpy() * B

        X_batch = X_batch.detach()

    for d in dims:
        for t in range(T):
            grad_maps[d][t] /= N

    return grad_maps

def compute_linear_probing(X: torch.Tensor, z_per_tp: Dict[str, np.ndarray],
                           dims: List[int]) -> Dict[int, Dict[int, np.ndarray]]:
    X_np = X.detach().cpu().numpy()
    N, T, H, W = X_np.shape
    
    linear_maps = {d: {t: np.zeros((H, W), dtype=float) for t in range(T)} for d in dims}
    scaler = StandardScaler()
    
    for t_idx, tp in enumerate(TP_NAMES):
        X_t = X_np[:, t_idx].reshape(N, -1)  # [N, H*W]
        ok_cols = (X_t.var(axis=0) > 1e-12)
        X_t_valid = X_t[:, ok_cols]
        if X_t_valid.shape[1] == 0: continue
        
        Xs = scaler.fit_transform(X_t_valid)
        
        for d in dims:
            y = z_per_tp[tp][:, d]
            ridge = Ridge(alpha=100.0, fit_intercept=True)
            ridge.fit(Xs, y)
            coefs = np.abs(ridge.coef_)
            
            map_1d = np.zeros(H * W, dtype=float)
            map_1d[ok_cols] = coefs
            linear_maps[d][t_idx] = map_1d.reshape(H, W)
            
    return linear_maps

def normalize_map(m: np.ndarray) -> np.ndarray:
    mn, mx = m.min(), m.max()
    if mx - mn < 1e-12: return np.zeros_like(m)
    return (m - mn) / (mx - mn)

def compute_consensus_maps(grad_maps: Dict, linear_maps: Dict,
                           dims: List[int], T: int = 3) -> Dict[int, Dict[int, np.ndarray]]:
    consensus = {d: {} for d in dims}
    for d in dims:
        for t in range(T):
            gm = normalize_map(grad_maps[d][t])
            lm = normalize_map(linear_maps[d][t])
            cm = (gm + lm) / 2.0
            consensus[d][t] = cm
    return consensus

def extract_roi_for_map(importance_map: np.ndarray, mse_map: np.ndarray,
                        importance_percentile: float, mse_threshold: float,
                        min_area: int = 3) -> np.ndarray:
    
    imp_thresh = np.percentile(importance_map, importance_percentile)
    roi_mask = (importance_map >= imp_thresh) & (mse_map < mse_threshold)

    labeled, n_comp = ndimage.label(roi_mask)
    filtered_mask = np.zeros_like(roi_mask)
    
    for cid in range(1, n_comp + 1):
        cmask = labeled == cid
        if cmask.sum() >= min_area:
            filtered_mask |= cmask
            
    return filtered_mask

# =============================================================================
# Main
# =============================================================================
def main():
    parser = argparse.ArgumentParser(description="Targeted ROI Extraction + MSE Filter")
    parser.add_argument("--tp_dirs", type=str, required=True)
    parser.add_argument("--step1_dir", type=str, required=True)
    parser.add_argument("--step2_dir", type=str, required=True)
    parser.add_argument("--model_dir", type=str, required=True)
    parser.add_argument("--output_dir", type=str, default="./step3_targeted_rois")
    parser.add_argument("--mask_percentile", type=float, default=95.0)
    parser.add_argument("--mse_threshold", type=float, default=0.5)
    parser.add_argument("--arch_file", type=str, required=True)
    parser.add_argument("--latent_dim", type=int, default=512)
    args = parser.parse_args()

    outdir = Path(args.output_dir)
    outdir.mkdir(parents=True, exist_ok=True)
    step1_dir = Path(args.step1_dir)
    step2_dir = Path(args.step2_dir)
    model_dir = Path(args.model_dir)
    
    tp_dirs = [d.strip() for d in args.tp_dirs.split(",")]
    
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    print("=" * 70)
    print("STEP 3: TARGETED PIXEL ATTRIBUTION + MSE FILTER")
    print("=" * 70)
    print(f"Device:          {device}")
    print(f"Mask Percentile: {args.mask_percentile}th")
    print(f"MSE Threshold:   {args.mse_threshold}")
    
    # 1. Load MSE Maps
    print("\nLoading MSE maps from Step 2...")
    mse_ensemble = {}
    for t_name in TP_NAMES:
        mse_path = step2_dir / f"mse_ensemble_{t_name}.npy"
        if not mse_path.exists():
            raise FileNotFoundError(f"Missing MSE map: {mse_path}")
        mse_ensemble[t_name] = np.load(mse_path)
        print(f"  {t_name}: max MSE = {mse_ensemble[t_name].max():.4f}")

    # 2. Load EEM data
    codes_per_tp = [set(list_sample_codes(d)) for d in tp_dirs]
    common_codes = sorted(list(codes_per_tp[0] & codes_per_tp[1] & codes_per_tp[2]))
    if not common_codes:
        raise ValueError("No common samples across timepoints!")
    
    print(f"\nLoading {len(common_codes)} samples...")
    X_np = load_all_samples(common_codes, tp_dirs)
    N, T, H, W = X_np.shape
    X_tensor = torch.tensor(X_np, dtype=torch.float32).to(device)

    # 3. Load Architecture Dynamically
    print(f"Loading architecture from {args.arch_file}...")
    arch_mod = import_arch_module(args.arch_file)
    
    arch_class = None
    if hasattr(arch_mod, 'ConvAutoencoderWithAttention'):
        arch_class = arch_mod.ConvAutoencoderWithAttention
    else:
        # Fallback
        import inspect
        for name, obj in inspect.getmembers(arch_mod, inspect.isclass):
            if issubclass(obj, nn.Module) and obj is not nn.Module:
                arch_class = obj
                break
                
    if arch_class is None:
        raise RuntimeError("Could not find an autoencoder class in arch_file")

    all_masks_dict = {}

    for m_name, target_dims in TARGET_FEATURES.items():
        m_step1_dir = step1_dir / m_name
        
        possible_paths = [
            model_dir / m_name / "model.pth",      
            model_dir / f"{m_name}.pth",           
            model_dir / f"model_fold{m_name[-1]}.pth" 
        ]
        
        model_path = None
        for p in possible_paths:
            if p.exists():
                model_path = p
                break
                
        if model_path is None:
            print(f"\n[WARN] Skipping {m_name}, could not find .pth file.")
            continue

        print(f"\nEvaluating Model: {m_name} (using {model_path.name})")
        
        # --- FIXED INSTANTIATION ---
        # Explicitly match the original step 3 parameters
        model = arch_class(
            in_channels=1, 
            latent_dim=args.latent_dim, 
            num_classes=2
        ).to(device)
            
        _load_state_dict_flexible(model, str(model_path))
        model.eval()
        
        # Load Z vectors for Linear Probing
        z_per_tp = {}
        z_df = pd.read_csv(m_step1_dir / "z_agg.csv")
        z_codes = z_df["code"].astype(str).tolist()
        code_idx_map = {c: i for i, c in enumerate(z_codes)}
        reindex = [code_idx_map[c] for c in common_codes]
        
        for tp in TP_NAMES:
            df_tp = pd.read_csv(m_step1_dir / f"z_{tp}.csv")
            dim_cols = [f"dim{d}" for d in range(args.latent_dim)]
            z_per_tp[tp] = df_tp[dim_cols].values[reindex]

        # Compute Importance Maps
        print(f"  -> Computing Gradient Saliency...")
        grad_maps = compute_gradient_saliency(model, X_tensor, target_dims)
        
        print(f"  -> Computing Linear Probing...")
        linear_maps = compute_linear_probing(X_tensor, z_per_tp, target_dims)
        
        print(f"  -> Building Consensus...")
        consensus_maps = compute_consensus_maps(grad_maps, linear_maps, target_dims, T=3)

        m_outdir = outdir / m_name
        m_outdir.mkdir(exist_ok=True)

        # Extract ROIs
        for d in target_dims:
            for t_idx, tp_name in enumerate(TP_NAMES):
                imp_map = consensus_maps[d][t_idx]
                mse_map = mse_ensemble[tp_name]
                
                # Filter Logic
                mask = extract_roi_for_map(imp_map, mse_map, args.mask_percentile, args.mse_threshold)
                
                # Save mask
                npy_path = m_outdir / f"roi_mask_dim{d}_{tp_name}.npy"
                np.save(npy_path, mask)
                
                dict_key = f"{m_name}_dim{d}_{tp_name}"
                all_masks_dict[dict_key] = mask
                
                n_px = mask.sum()
                print(f"      Dim {d} - {tp_name}: Valid Pixels = {n_px}/{H*W}")

    if all_masks_dict:
        npz_path = outdir / "roi_all_masks.npz"
        np.savez_compressed(npz_path, **all_masks_dict)
        print(f"\nSaved all masks to {npz_path}")
    
    print("\nTargeted Step 3 Complete!")

if __name__ == "__main__":
    main()