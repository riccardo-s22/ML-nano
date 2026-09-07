#!/usr/bin/env python3
"""
STEP 4 v2 (FINAL): Targeted Two-Stage Proxy Reconstruction
===========================================================

Changes from v1:
  1. Reads TARGET_FEATURES from Step OPT v2's selected_features_config.json
     instead of hardcoded features.
  2. Falls back to --target_features CLI arg if config not found.
  3. Reports balanced accuracy and AUC alongside accuracy.
  4. Saves per-sample proxy matrix with sample codes for downstream use.
  5. Reports proxy reconstruction quality (R² per feature).

USAGE:
------
python step4_final_proxy_v2.py \
    --tp_dirs ./out_0h,./out_6h,./out_24h \
    --labels_csv sample_labels.csv \
    --step1_dir ./step1_latent_features \
    --step3_dir ./step3_targeted_rois_v2 \
    --opt_config ./opt_selection_results_v2/selected_features_config.json \
    --output_dir ./step4_final_results_v2
"""

import os
import json
import argparse
import numpy as np
import pandas as pd
from pathlib import Path
from typing import List, Dict, Tuple, Optional
from openpyxl import load_workbook
from sklearn.linear_model import Ridge
from sklearn.model_selection import KFold, LeaveOneOut
from sklearn.preprocessing import StandardScaler
from sklearn.svm import SVC
from sklearn.ensemble import RandomForestClassifier
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    accuracy_score, balanced_accuracy_score, roc_auc_score, confusion_matrix
)
from scipy.stats import binom

import warnings
warnings.filterwarnings("ignore")

SEED = 42
np.random.seed(SEED)

TP_NAMES = ["0h", "6h", "24h"]


# =============================================================================
# Config Loading
# =============================================================================

def load_config(config_path: str) -> Tuple[List[str], Dict[str, List[int]]]:
    """Load target features and model_dims from Step OPT v2 config."""
    with open(config_path, 'r') as f:
        config = json.load(f)
    target_features = config["target_features"]
    model_dims = config["model_dims"]
    return target_features, model_dims


def parse_target_features(feature_list: List[str]) -> Dict[str, List[int]]:
    """Parse feature names -> {model_name: [dim_indices]}"""
    needed = {}
    for f in feature_list:
        idx = f.rfind("_dim")
        if idx == -1:
            continue
        mname = f[:idx]
        dim_str = f[idx + 4:]
        if mname not in needed:
            needed[mname] = []
        needed[mname].append(int(dim_str))
    return needed


# =============================================================================
# Data Loading Helpers
# =============================================================================

def load_labels(csv_path: str):
    df = pd.read_csv(csv_path)
    codes = df["code"].tolist()
    groups = df["group"].values
    y = (groups == "ALS").astype(int)
    return codes, y, groups


def list_sample_codes(tp_dir: str) -> List[str]:
    codes = []
    for f in Path(tp_dir).glob("*.xlsx"):
        if f.name.startswith("~$"):
            continue
        name = f.stem
        for suffix in ["_0h", "_6h", "_24h"]:
            if name.endswith(suffix):
                name = name[:-len(suffix)]
                break
        codes.append(name)
    return codes


def find_file_for_code(tp_dir: str, code: str) -> Optional[str]:
    tp_path = Path(tp_dir)
    exact_path = tp_path / f"{code}.xlsx"
    if exact_path.exists():
        return str(exact_path)
    for f in tp_path.glob("*.xlsx"):
        if f.name.startswith("~$"):
            continue
        stem = f.stem
        for suffix in ["_0h", "_6h", "_24h"]:
            if stem.endswith(suffix):
                file_code = stem[:-len(suffix)]
                if file_code == code:
                    return str(f)
                if file_code.replace('.', '_') == code:
                    return str(f)
    return None


def load_excel_as_array(filepath: str) -> np.ndarray:
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    data = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx == 0:
            continue
        row_data = []
        for col_idx, cell in enumerate(row):
            if col_idx == 0:
                continue
            if isinstance(cell, (int, float)) and cell is not None:
                row_data.append(float(cell))
            else:
                row_data.append(0.0)
        if row_data:
            data.append(row_data)
    arr = np.array(data, dtype=np.float32)
    mn, mx = arr.min(), arr.max()
    if mx > mn:
        arr = (arr - mn) / (mx - mn)
    else:
        arr = np.zeros_like(arr)
    return arr


def load_all_samples(codes: List[str], tp_dirs: List[str]) -> np.ndarray:
    arrays = []
    for code in codes:
        tp_images = []
        for tp_dir in tp_dirs:
            fpath = find_file_for_code(tp_dir, code)
            if fpath is None:
                raise FileNotFoundError(f"Missing {code} in {tp_dir}")
            arr = load_excel_as_array(fpath)
            tp_images.append(arr)
        shapes = [a.shape for a in tp_images]
        if len(set(shapes)) > 1:
            min_h = min(s[0] for s in shapes)
            min_w = min(s[1] for s in shapes)
            tp_images = [a[:min_h, :min_w] for a in tp_images]
        arrays.append(np.stack(tp_images, axis=0))
    return np.stack(arrays, axis=0)


# =============================================================================
# Ridge Regression Helpers
# =============================================================================

def fit_predict_ridge(X_tr, y_tr, X_te, alpha):
    mu = X_tr.mean(axis=0, keepdims=True)
    sd = X_tr.std(axis=0, keepdims=True)
    sd[sd < 1e-12] = 1.0
    Xs_tr = (X_tr - mu) / sd
    Xs_te = (X_te - mu) / sd
    reg = Ridge(alpha=float(alpha), fit_intercept=True, random_state=SEED)
    reg.fit(Xs_tr, y_tr)
    return reg.predict(Xs_tr), reg.predict(Xs_te)


def select_alpha_inner_cv(X_tr, y_tr, alpha_grid, inner_folds=3):
    n = X_tr.shape[0]
    if n < 6:
        return float(alpha_grid[0])
    k = min(int(inner_folds), n - 1)
    kf = KFold(n_splits=k, shuffle=True, random_state=SEED)
    best_alpha = float(alpha_grid[0])
    best_score = -np.inf
    for a in alpha_grid:
        scores = []
        for tr_i, va_i in kf.split(np.arange(n)):
            _, pred_va = fit_predict_ridge(X_tr[tr_i], y_tr[tr_i], X_tr[va_i], alpha=a)
            ss_res = np.sum((y_tr[va_i] - pred_va) ** 2)
            ss_tot = np.sum((y_tr[va_i] - np.mean(y_tr[va_i])) ** 2)
            r2 = 1.0 - ss_res / ss_tot if ss_tot > 1e-12 else 0.0
            scores.append(r2)
        if np.mean(scores) > best_score:
            best_score = np.mean(scores)
            best_alpha = float(a)
    return best_alpha


# =============================================================================
# Power calculation (same as Step OPT v2)
# =============================================================================

def calculate_binomial_power(accuracy: float, n_samples: int, alpha: float = 0.05) -> float:
    """Exact binomial power for classification accuracy vs chance (0.5)."""
    if accuracy <= 0.5:
        return 0.0
    n = int(n_samples)
    p0 = 0.5
    p1 = float(accuracy)
    k_crit = int(binom.ppf(1.0 - alpha, n, p0)) + 1
    if k_crit > n:
        return 0.0
    power = 1.0 - binom.cdf(k_crit - 1, n, p1)
    return float(power)


# =============================================================================
# Model Bundle Loading
# =============================================================================

class ModelBundle:
    def __init__(self, model_name, top_dims, z_agg, z_per_tp, roi_masks):
        self.model_name = model_name
        self.top_dims = top_dims
        self.z_agg = z_agg
        self.z_per_tp = z_per_tp
        self.roi_masks = roi_masks


def load_targeted_bundles(step1_dir: Path, step3_dir: Path,
                          needed_dims: Dict[str, List[int]],
                          codes: List[str], H: int, W: int,
                          latent_dim: int) -> List[ModelBundle]:

    # Pre-load ensemble fallbacks
    ens_z_tp = {}
    for tp in TP_NAMES:
        p = step1_dir / f"ensemble_z_{tp}.csv"
        if p.exists():
            ens_z_tp[tp] = pd.read_csv(p)
    ens_z_agg_df = None
    if (step1_dir / "ensemble_z_agg.csv").exists():
        ens_z_agg_df = pd.read_csv(step1_dir / "ensemble_z_agg.csv")

    # ROI npz
    npz_path = step3_dir / "roi_all_masks.npz"
    npz = np.load(npz_path) if npz_path.exists() else None

    bundles = []
    dim_cols = [f"dim{d}" for d in range(latent_dim)]

    for mname, dims in needed_dims.items():
        s1_mdir = step1_dir / mname
        s3_mdir = step3_dir / mname

        # Load Z vectors
        if (s1_mdir / "z_agg.csv").exists():
            df = pd.read_csv(s1_mdir / "z_agg.csv")
        elif ens_z_agg_df is not None:
            df = ens_z_agg_df
        else:
            raise FileNotFoundError(f"No z_agg for {mname}")

        z_codes = df["code"].astype(str).tolist()
        code_map = {c: i for i, c in enumerate(z_codes)}
        reindex = [code_map[c] for c in codes]
        z_agg = df[dim_cols].values[reindex]

        z_tp_full = {}
        for tp in TP_NAMES:
            if (s1_mdir / f"z_{tp}.csv").exists():
                df_tp = pd.read_csv(s1_mdir / f"z_{tp}.csv")
            elif tp in ens_z_tp:
                df_tp = ens_z_tp[tp]
            else:
                raise FileNotFoundError(f"Missing z_{tp} for {mname}")
            z_tp_full[tp] = df_tp[dim_cols].values[reindex]

        # Load ROI masks
        roi_masks = {}
        for d in dims:
            roi_masks[d] = {}
            for tp in TP_NAMES:
                mask = None

                # Try individual .npy first
                fpath = s3_mdir / f"roi_mask_dim{d}_{tp}.npy"
                if fpath.exists():
                    mask = np.load(fpath).astype(bool)

                # Try NPZ
                if mask is None and npz is not None:
                    key = f"{mname}_dim{d}_{tp}"
                    if key in npz:
                        mask = npz[key].astype(bool)

                # Fallback: full mask
                if mask is None:
                    print(f"  [WARN] No ROI mask for {mname} dim{d} {tp}, using full matrix")
                    mask = np.ones((H, W), dtype=bool)
                elif mask.shape != (H, W):
                    print(f"  [WARN] ROI shape mismatch for {mname} dim{d} {tp}, using full matrix")
                    mask = np.ones((H, W), dtype=bool)

                roi_masks[d][tp] = mask

        bundles.append(ModelBundle(mname, dims, z_agg, z_tp_full, roi_masks))

    if npz:
        npz.close()
    return bundles


# =============================================================================
# Proxy Builder
# =============================================================================

def build_proxy_single_fold(X_np, bundle, train_idx, test_idx, alpha_grid):
    """Build proxy features for one LOOCV fold."""
    n_train = len(train_idx)
    proxies_tr = np.zeros((n_train, len(bundle.top_dims)))
    proxies_te = np.zeros((1, len(bundle.top_dims)))

    for i, d in enumerate(bundle.top_dims):
        # Stage 1: per-timepoint Ridge (pixel -> z_t[d])
        s1_tr = np.zeros((n_train, 3))
        s1_te = np.zeros((1, 3))

        for t_i, tp in enumerate(TP_NAMES):
            y = bundle.z_per_tp[tp][:, d]
            y_tr = y[train_idx]
            mask = bundle.roi_masks[d][tp]

            pixels = X_np[:, t_i].reshape(X_np.shape[0], -1)
            mask_flat = mask.ravel()
            if not mask_flat.any():
                s1_tr[:, t_i] = np.mean(y_tr)
                s1_te[:, t_i] = np.mean(y_tr)
                continue

            X_des = pixels[:, mask_flat]
            X_tr, X_te = X_des[train_idx], X_des[test_idx:test_idx + 1]

            a = select_alpha_inner_cv(X_tr, y_tr, alpha_grid)
            ptr, pte = fit_predict_ridge(X_tr, y_tr, X_te, a)
            s1_tr[:, t_i] = ptr
            s1_te[:, t_i] = pte[0]

        # Stage 2: aggregate (per-tp predictions -> z_agg[d])
        y_agg = bundle.z_agg[:, d]
        y_agg_tr = y_agg[train_idx]

        if s1_tr.std() < 1e-12:
            proxies_tr[:, i] = np.mean(y_agg_tr)
            proxies_te[:, i] = np.mean(y_agg_tr)
        else:
            a2 = select_alpha_inner_cv(s1_tr, y_agg_tr, alpha_grid)
            ptr, pte = fit_predict_ridge(s1_tr, y_agg_tr, s1_te, a2)
            proxies_tr[:, i] = ptr
            proxies_te[:, i] = pte[0]

    return proxies_tr, proxies_te


# =============================================================================
# Main
# =============================================================================

def main():
    parser = argparse.ArgumentParser(description="Step 4 v2: Proxy reconstruction & classification")
    parser.add_argument("--tp_dirs", required=True)
    parser.add_argument("--labels_csv", required=True)
    parser.add_argument("--step1_dir", required=True)
    parser.add_argument("--step3_dir", required=True)
    parser.add_argument("--output_dir", default="./step4_final_results_v2")
    parser.add_argument("--latent_dim", type=int, default=512)
    # Config input
    parser.add_argument("--opt_config", type=str, default=None,
                        help="Path to selected_features_config.json from Step OPT v2")
    parser.add_argument("--target_features", type=str, default=None,
                        help="Comma-separated feature names (fallback)")
    args = parser.parse_args()

    outdir = Path(args.output_dir)
    outdir.mkdir(parents=True, exist_ok=True)

    # --- Load target features ---
    if args.opt_config and Path(args.opt_config).exists():
        TARGET_FEATURES, needed_dims = load_config(args.opt_config)
        print(f"Loaded {len(TARGET_FEATURES)} features from config: {args.opt_config}")
    elif args.target_features:
        TARGET_FEATURES = [f.strip() for f in args.target_features.split(",")]
        needed_dims = parse_target_features(TARGET_FEATURES)
    else:
        raise ValueError("Must provide either --opt_config or --target_features")

    print("=" * 70)
    print("STEP 4 v2: PROXY RECONSTRUCTION & CLASSIFICATION")
    print("=" * 70)
    print(f"Target features: {TARGET_FEATURES}")
    print(f"Models/dims: {needed_dims}")

    # 1. Load Data
    codes, y, groups = load_labels(args.labels_csv)
    tp_dirs = args.tp_dirs.split(",")

    found_codes = set(list_sample_codes(tp_dirs[0]))
    for d in tp_dirs[1:]:
        found_codes &= set(list_sample_codes(d))

    keep_mask = [c in found_codes for c in codes]
    codes = [c for c, k in zip(codes, keep_mask) if k]
    y = y[np.array(keep_mask)]
    groups = groups[np.array(keep_mask)]
    N = len(codes)
    n_als = int(y.sum())
    n_ctrl = N - n_als
    print(f"Samples: {N} (ALS={n_als}, CTRL={n_ctrl})")

    X_np = load_all_samples(codes, tp_dirs)
    _, T, H, W = X_np.shape

    # 2. Load Bundles
    bundles = load_targeted_bundles(
        Path(args.step1_dir), Path(args.step3_dir),
        needed_dims, codes, H, W, args.latent_dim
    )

    # 3. LOOCV
    print(f"\nRunning LOOCV Proxy Reconstruction ({len(TARGET_FEATURES)} features)...")

    proxy_oof = np.zeros((N, len(TARGET_FEATURES)))
    classifiers = {
        "SVM": SVC(kernel="linear", probability=True, class_weight="balanced", random_state=SEED),
        "RF": RandomForestClassifier(n_estimators=500, class_weight="balanced_subsample", random_state=SEED),
        "LR": LogisticRegression(class_weight="balanced", solver="liblinear", random_state=SEED)
    }
    y_preds = {k: np.zeros(N, dtype=int) for k in classifiers}
    y_probs = {k: np.zeros(N) for k in classifiers}

    loo = LeaveOneOut()
    alpha_grid = [0.1, 1.0, 10.0, 100.0]

    for fold, (train_idx, test_idx) in enumerate(loo.split(np.arange(N))):
        test_i = test_idx[0]
        if fold % 5 == 0:
            print(f"  Fold {fold + 1}/{N}...")

        # Build proxy matrix for this fold
        bundle_res = {}
        for b in bundles:
            tr, te = build_proxy_single_fold(X_np, b, train_idx, test_i, alpha_grid)
            bundle_res[b.model_name] = (tr, te)

        fold_X_tr = np.zeros((len(train_idx), len(TARGET_FEATURES)))
        fold_X_te = np.zeros((1, len(TARGET_FEATURES)))

        for i, fname in enumerate(TARGET_FEATURES):
            idx = fname.rfind("_dim")
            mname = fname[:idx]
            dim = int(fname[idx + 4:])

            b_tr, b_te = bundle_res[mname]
            col_idx = needed_dims[mname].index(dim)

            fold_X_tr[:, i] = b_tr[:, col_idx]
            fold_X_te[:, i] = b_te[:, col_idx]

        proxy_oof[test_i] = fold_X_te[0]

        scaler = StandardScaler()
        X_tr_s = scaler.fit_transform(fold_X_tr)
        X_te_s = scaler.transform(fold_X_te)

        for name, clf in classifiers.items():
            clf.fit(X_tr_s, y[train_idx])
            if hasattr(clf, "predict_proba"):
                prob = clf.predict_proba(X_te_s)[0, 1]
            else:
                prob = clf.decision_function(X_te_s)[0]
            y_probs[name][test_i] = prob
            y_preds[name][test_i] = int(prob >= 0.5)

    # 4. Results
    print("\n" + "=" * 70)
    print(f"FINAL RESULTS ({len(TARGET_FEATURES)} Proxy Features)")
    print("=" * 70)

    res_rows = []
    for name in classifiers:
        acc = accuracy_score(y, y_preds[name])
        bal_acc = balanced_accuracy_score(y, y_preds[name])
        auc = roc_auc_score(y, y_probs[name])
        cm = confusion_matrix(y, y_preds[name])
        power = calculate_binomial_power(acc, N)

        print(f"\n{name}:")
        print(f"  Accuracy:          {acc:.3f}")
        print(f"  Balanced Accuracy: {bal_acc:.3f}")
        print(f"  AUC:               {auc:.3f}")
        print(f"  Binomial Power:    {power:.3f}")
        print(f"  Confusion Matrix:  {cm.tolist()}")

        tn, fp, fn, tp_ = cm.ravel()
        sens = tp_ / (tp_ + fn) if (tp_ + fn) > 0 else 0.0
        spec = tn / (tn + fp) if (tn + fp) > 0 else 0.0
        print(f"  Sensitivity:       {sens:.3f}")
        print(f"  Specificity:       {spec:.3f}")

        res_rows.append({
            "classifier": name,
            "accuracy": acc,
            "balanced_accuracy": bal_acc,
            "auc": auc,
            "sensitivity": sens,
            "specificity": spec,
            "binomial_power": power,
        })

    # 5. Proxy reconstruction quality
    print("\n--- Proxy Reconstruction Quality (R² vs true latent) ---")
    r2_rows = []
    for i, fname in enumerate(TARGET_FEATURES):
        idx = fname.rfind("_dim")
        mname = fname[:idx]
        dim = int(fname[idx + 4:])

        # Get the true z_agg for this dim
        for b in bundles:
            if b.model_name == mname:
                true_z = b.z_agg[:, dim]
                break

        proxy_z = proxy_oof[:, i]
        ss_res = np.sum((true_z - proxy_z) ** 2)
        ss_tot = np.sum((true_z - true_z.mean()) ** 2)
        r2 = 1.0 - ss_res / ss_tot if ss_tot > 1e-12 else 0.0
        corr = np.corrcoef(true_z, proxy_z)[0, 1] if np.std(proxy_z) > 1e-12 else 0.0

        r2_rows.append({"feature": fname, "r2": r2, "pearson_r": corr})
        print(f"  {fname}: R²={r2:.3f}, Pearson r={corr:.3f}")

    # 6. Save outputs
    proxy_df = pd.DataFrame(proxy_oof, columns=TARGET_FEATURES)
    proxy_df.insert(0, "code", codes)
    proxy_df.insert(1, "group", groups)
    proxy_df.to_csv(outdir / "final_proxy_features.csv", index=False)

    pd.DataFrame(res_rows).to_csv(outdir / "final_classification_metrics.csv", index=False)
    pd.DataFrame(r2_rows).to_csv(outdir / "proxy_reconstruction_quality.csv", index=False)

    # Save per-sample predictions
    pred_df = pd.DataFrame({"code": codes, "group": groups, "true_label": y})
    for name in classifiers:
        pred_df[f"{name}_pred"] = y_preds[name]
        pred_df[f"{name}_prob"] = y_probs[name]
    pred_df.to_csv(outdir / "per_sample_predictions.csv", index=False)

    # Save step4 config for traceability
    step4_config = {
        "step4_version": "v2",
        "target_features": TARGET_FEATURES,
        "n_samples": N,
        "n_als": n_als,
        "n_ctrl": n_ctrl,
        "results": res_rows,
        "proxy_quality": r2_rows,
    }
    with open(outdir / "step4_config.json", "w") as f:
        json.dump(step4_config, f, indent=2, default=str)

    print(f"\nAll outputs saved to {outdir}")


if __name__ == "__main__":
    main()
