"""
Cache the EEM tensor used by the conv-autoencoder into a single npz so the
PLS-DA baseline can reuse *exactly* the same inputs (per-file min-max to [0,1],
3 timepoints stacked) and the same canonical 5-fold OOF split.
"""
import os
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = "/mnt/c/Users/riccardo-s/Documents/CNT/targetALS/exp5"
TP_DIRS = [os.path.join(ROOT, d) for d in ["out_0h", "out_6h", "out_24h"]]
LABELS_CSV = os.path.join(ROOT, "sample_labels.csv")
OOF_CSV = os.path.join(ROOT, "model_metrics_eval", "existing_models_oof_predictions.csv")
OUT = os.path.join(ROOT, "plsda_baseline", "eem_cache.npz")


def load_excel(fp):
    """Identical to eval_5fold_metrics.load_excel: strip header row + first col,
    per-file min-max normalize to [0,1]."""
    wb = load_workbook(fp, data_only=True)
    ws = wb.active
    data = []
    for ri, row in enumerate(ws.iter_rows(values_only=True)):
        if ri == 0:
            continue
        rd = []
        for ci, c in enumerate(row):
            if ci == 0:
                continue
            rd.append(float(c) if isinstance(c, (int, float)) and c is not None else 0.0)
        if rd:
            data.append(rd)
    arr = np.array(data, dtype=np.float32)
    mn, mx = arr.min(), arr.max()
    arr = (arr - mn) / (mx - mn) if mx > mn else np.zeros_like(arr)
    return arr


def main():
    lab = pd.read_csv(LABELS_CSV)
    codes = lab["code"].astype(str).tolist()
    y = np.array([1 if str(g).strip().lower() == "als" else 0 for g in lab["group"]], dtype=int)

    oof = pd.read_csv(OOF_CSV)
    fold_map = dict(zip(oof["code"].astype(str), oof["held_out_fold"].astype(int)))
    fold = np.array([fold_map[c] for c in codes], dtype=int)

    X = []
    for c in codes:
        tps = [load_excel(os.path.join(d, f"{c}.xlsx")) for d in TP_DIRS]
        X.append(np.stack(tps, axis=0))  # [3, H, W]
    X = np.stack(X, axis=0).astype(np.float32)  # [N, 3, H, W]
    print("X shape:", X.shape, "y:", y.sum(), "ALS /", len(y))

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    np.savez_compressed(OUT, X=X, y=y, fold=fold, codes=np.array(codes))
    print("wrote", OUT)


if __name__ == "__main__":
    main()
