"""
Evaluate the saved 5-fold autoencoder+classifier models in 5_fold_models_original/
using the architecture defined in conv_autoencoder_detailed.py.

Step 1 (this run): load data, run each fold model on ALL 39 subjects, cache
per-model predictions, and run a self-consistency check that the assumed
per-fold validation split (recovered from attention_reports CSVs) matches the
held-out subjects of each saved model.
"""
import os, sys, importlib.util
import numpy as np
import pandas as pd
import torch
from openpyxl import load_workbook

ROOT = "/mnt/c/Users/riccardo-s/Documents/CNT/targetALS/exp5"
MODEL_DIR = os.path.join(ROOT, "5_fold_models_original")
TP_DIRS = [os.path.join(ROOT, d) for d in ["out_0h", "out_6h", "out_24h"]]
LABELS_CSV = os.path.join(ROOT, "sample_labels.csv")
OUT_DIR = os.path.join(ROOT, "model_metrics_eval")
os.makedirs(OUT_DIR, exist_ok=True)

DEVICE = torch.device("cpu")
LATENT_DIM = 512
torch.manual_seed(42); np.random.seed(42)

# ---- import model classes from conv_autoencoder_detailed.py ----
spec = importlib.util.spec_from_file_location(
    "cae", os.path.join(ROOT, "conv_autoencoder_detailed.py"))
cae = importlib.util.module_from_spec(spec)
spec.loader.exec_module(cae)
Model = cae.ConvAutoencoderWithAttention


def load_excel(fp):
    wb = load_workbook(fp, data_only=True); ws = wb.active
    data = []
    for ri, row in enumerate(ws.iter_rows(values_only=True)):
        if ri == 0:
            continue
        rd = []
        for ci, c in enumerate(row):
            if ci == 0:
                continue
            rd.append(float(c) if isinstance(c, (int, float)) and c is not None else 0.0)
        if rd:
            data.append(rd)
    arr = np.array(data, dtype=np.float32)
    mn, mx = arr.min(), arr.max()
    arr = (arr - mn) / (mx - mn) if mx > mn else np.zeros_like(arr)
    return arr


def main():
    lab = pd.read_csv(LABELS_CSV)
    codes = lab["code"].astype(str).tolist()
    y = np.array([1 if str(g).strip().lower() == "als" else 0 for g in lab["group"]], dtype=int)
    code2y = dict(zip(codes, y))

    # assumed per-fold val split from attention_reports CSVs (deterministic seed)
    val_by_fold = {}
    for k in range(1, 6):
        df = pd.read_csv(os.path.join(ROOT, "attention_reports", f"attention_weights_fold{k}.csv"))
        val_by_fold[k] = df[df["split"] == "val"]["code"].astype(str).tolist()

    # load all samples -> X [N,3,H,W]
    print("Loading data...")
    X = []
    for c in codes:
        tps = [load_excel(os.path.join(d, f"{c}.xlsx")) for d in TP_DIRS]
        X.append(np.stack(tps, axis=0))
    X = np.stack(X, axis=0).astype(np.float32)
    print("X shape:", X.shape)
    X_t = torch.from_numpy(X).unsqueeze(2)  # [N,3,1,H,W]

    # predictions matrix: prob_class1 for each (model_fold, sample)
    prob = np.full((5, len(codes)), np.nan)
    for k in range(1, 6):
        mpath = os.path.join(MODEL_DIR, f"best_model_fold{k}.pth")
        print(f"\nLoading {mpath}")
        model = Model(in_channels=1, latent_dim=LATENT_DIM, num_classes=2).to(DEVICE)
        # trigger lazy init of encoder.fc and decoder via dummy forward
        with torch.no_grad():
            _ = model(X_t[:1])
        state = torch.load(mpath, map_location=DEVICE)
        model.load_state_dict(state, strict=True)
        model.eval()
        with torch.no_grad():
            _, logits, _ = model(X_t)
            p1 = torch.softmax(logits, dim=1)[:, 1].cpu().numpy()
        prob[k - 1] = p1

    # cache
    np.savez(os.path.join(OUT_DIR, "pred_cache.npz"),
             prob=prob, codes=np.array(codes), y=y)

    # self-consistency: per fold model, train vs val accuracy under assumed split
    print("\n==== SELF-CONSISTENCY CHECK (assumed split) ====")
    code_idx = {c: i for i, c in enumerate(codes)}
    for k in range(1, 6):
        val_codes = set(val_by_fold[k])
        val_i = [code_idx[c] for c in codes if c in val_codes]
        tr_i = [code_idx[c] for c in codes if c not in val_codes]
        pk = prob[k - 1]
        pred = (pk >= 0.5).astype(int)
        tr_acc = (pred[tr_i] == y[tr_i]).mean()
        va_acc = (pred[val_i] == y[val_i]).mean()
        print(f"fold{k}: train_acc={tr_acc:.3f} (n={len(tr_i)})  "
              f"val_acc={va_acc:.3f} (n={len(val_i)})  gap={tr_acc-va_acc:+.3f}")
    print("\nIf train_acc is consistently ~1.0 and > val_acc, the assumed split "
          "matches the saved models' held-out sets.")


if __name__ == "__main__":
    main()
