"""
NO-C9 variant: identical faithful re-run of the ORIGINAL training procedure,
but EXCLUDING the two control subjects carrying a C9orf72 mutation:
    NEUTK792LKM -> 10.3d__570.opj_with_emission  (CTRL)
    NEUPL452UJJ -> 30.3d__570.opj_with_emission  (CTRL)
Cohort drops from 39 (20 ALS / 19 CTRL) to 37 (20 ALS / 17 CTRL). Everything
else (seeds, StratifiedKFold(5, shuffle, rs=42), 100 epochs, batch 4, best-val-
acc checkpoint, frozen lazy fc/decoder) is byte-for-byte the same as
rerun_faithful.py. Outputs go to model_metrics_rerun_noC9/ (original untouched).

Ported from:
  C:\\Users\\riccardo-s\\Documents\\CNT\\paper\\CNN\\input_importances\\conv_autoencoder_detailed.py

Key fidelity points (these are what my first re-run got wrong):
  * BATCH_SIZE = 4
  * NUM_EPOCHS = 100, NO early stopping (always full 100 epochs)
  * best checkpoint selected by best VALIDATION ACCURACY (strict >), not val_loss
  * optimizer built BEFORE the first forward pass  ->  encoder.fc and decoder
    are lazily created afterwards and are therefore NEVER trained (they keep
    PyTorch's default init). Verified against the saved models.
  * seeds set once globally (torch.manual_seed(42), np.random.seed(42)); RNG
    flows across folds (no per-fold reseed)
  * StratifiedKFold(5, shuffle=True, random_state=42) on codes/labels
  * data: out_0h, out_6h, out_24h ; labels: sample_labels.csv (LabelEncoder ->
    ALS=0, CTRL=1)

NOTE: exact weights won't reproduce (original likely ran on GPU / different
torch build; CPU conv math differs), but the PROCEDURE matches.
"""
import os, time, json, importlib.util
import numpy as np
import pandas as pd
import torch
import torch.nn as nn
import torch.optim as optim
from torch.utils.data import DataLoader, Dataset
from openpyxl import load_workbook
from sklearn.model_selection import StratifiedKFold
from sklearn.preprocessing import LabelEncoder
from sklearn.metrics import (roc_auc_score, accuracy_score, precision_score,
                             recall_score, f1_score)

# global seeds, exactly like the original module top
torch.manual_seed(42)
np.random.seed(42)

# portable root so this runs under the Windows `als-prs` torch env (WSL has no torch)
ROOT = os.path.dirname(os.path.abspath(__file__))
LABELS_CSV = os.path.join(ROOT, "sample_labels.csv")
TIMEPOINT_DIRS = [os.path.join(ROOT, d) for d in ["out_0h", "out_6h", "out_24h"]]
OUT = os.path.join(ROOT, "model_metrics_rerun_noC9")
MDL = os.path.join(OUT, "models")
os.makedirs(MDL, exist_ok=True)

# C9orf72-mutation controls to exclude (subject IDs NEUTK792LKM / NEUPL452UJJ)
EXCLUDE_CODES = {"10.3d__570.opj_with_emission", "30.3d__570.opj_with_emission"}

DEVICE = torch.device("cpu")
LATENT_DIM = 512
BATCH_SIZE = 4
NUM_EPOCHS = 100
LEARNING_RATE = 0.001
ALPHA = 1.0
BETA = 1.0
N_SPLITS = 5
torch.set_num_threads(max(1, os.cpu_count() or 1))

# import the model classes + ExcelImageDataset from the exp5 copy (identical arch)
spec = importlib.util.spec_from_file_location(
    "cae", os.path.join(ROOT, "conv_autoencoder_detailed.py"))
cae = importlib.util.module_from_spec(spec); spec.loader.exec_module(cae)
ConvAutoencoderWithAttention = cae.ConvAutoencoderWithAttention


def load_excel_image(fp):
    """Identical to ExcelImageDataset.load_excel_as_image: skip row0/col0,
    non-numeric->0.0, per-file min-max normalize to [0,1]; returns (1,H,W)."""
    wb = load_workbook(fp, data_only=True); ws = wb.active
    data = []
    for ri, row in enumerate(ws.iter_rows(values_only=True)):
        if ri == 0:
            continue
        rd = [float(c) if isinstance(c, (int, float)) and c is not None else 0.0
              for ci, c in enumerate(row) if ci != 0]
        if rd:
            data.append(rd)
    arr = np.array(data, dtype=np.float32)
    mn, mx = arr.min(), arr.max()
    arr = (arr - mn) / (mx - mn) if mx > mn else np.zeros_like(arr)
    return arr[np.newaxis, :, :]


class CachedDataset(Dataset):
    """Returns preloaded (3,1,H,W) tensors + long label -- byte-identical to
    what ExcelImageDataset would yield, but with no per-epoch disk I/O."""
    def __init__(self, X, y):
        self.X = X; self.y = y
    def __len__(self):
        return len(self.y)
    def __getitem__(self, i):
        return self.X[i], torch.tensor(int(self.y[i]), dtype=torch.long)


def train_epoch(model, loader, optimizer, device, alpha, beta):
    model.train()
    rc, cc = nn.MSELoss(), nn.CrossEntropyLoss()
    tot = tr = tc = 0.0; correct = n = 0
    for images, labels in loader:
        images, labels = images.to(device), labels.to(device)
        optimizer.zero_grad()
        recon, logits, _ = model(images)
        recon_loss = rc(recon, images); class_loss = cc(logits, labels)
        loss = alpha * recon_loss + beta * class_loss
        loss.backward(); optimizer.step()
        tot += loss.item(); tr += recon_loss.item(); tc += class_loss.item()
        correct += (logits.argmax(1) == labels).sum().item(); n += labels.size(0)
    nb = len(loader)
    return tot / nb, tr / nb, tc / nb, 100.0 * correct / n


@torch.no_grad()
def validate(model, loader, device, alpha, beta):
    model.eval()
    rc, cc = nn.MSELoss(), nn.CrossEntropyLoss()
    tot = tr = tc = 0.0; correct = n = 0
    for images, labels in loader:
        images, labels = images.to(device), labels.to(device)
        recon, logits, _ = model(images)
        recon_loss = rc(recon, images); class_loss = cc(logits, labels)
        tot += (alpha * recon_loss + beta * class_loss).item()
        tr += recon_loss.item(); tc += class_loss.item()
        correct += (logits.argmax(1) == labels).sum().item(); n += labels.size(0)
    nb = len(loader)
    return tot / nb, tr / nb, tc / nb, 100.0 * correct / n


@torch.no_grad()
def predict_probs(model, loader, device):
    """P(class1) and labels for a loader (class1 == CTRL under LabelEncoder)."""
    model.eval(); ps = []; ys = []
    for images, labels in loader:
        _, logits, _ = model(images.to(device))
        ps.append(torch.softmax(logits, 1)[:, 1].cpu().numpy()); ys.append(labels.numpy())
    return np.concatenate(ps), np.concatenate(ys)


def main():
    log = open(os.path.join(OUT, "faithful_log.txt"), "w")
    def p(*a):
        m = " ".join(str(x) for x in a); print(m); log.write(m + "\n"); log.flush()

    labels_df = pd.read_csv(LABELS_CSV)
    n_before = len(labels_df)
    dropped = labels_df[labels_df["code"].astype(str).isin(EXCLUDE_CODES)]
    labels_df = labels_df[~labels_df["code"].astype(str).isin(EXCLUDE_CODES)].reset_index(drop=True)
    p(f"Excluded {len(dropped)} C9 control(s): {list(dropped['code'])}")
    p(f"Cohort: {n_before} -> {len(labels_df)} subjects "
      f"({(labels_df['group']=='ALS').sum()} ALS / {(labels_df['group']=='CTRL').sum()} CTRL)")
    assert len(dropped) == len(EXCLUDE_CODES), "not all excluded codes were found in labels!"
    le = LabelEncoder()
    labels_df["enc"] = le.fit_transform(labels_df["group"])  # ALS=0, CTRL=1
    p("LabelEncoder classes (index=encoded value):", list(le.classes_))
    codes = labels_df["code"].astype(str).values
    labels = labels_df["enc"].values
    num_classes = len(le.classes_)

    # preload all samples once (byte-identical to ExcelImageDataset output)
    p("Preloading Excel data into memory ...")
    t_pre = time.time()
    cache = {}
    for c in codes:
        tps = [load_excel_image(os.path.join(d, f"{c}.xlsx")) for d in TIMEPOINT_DIRS]
        cache[c] = torch.from_numpy(np.stack(tps, axis=0))  # (3,1,H,W)
    p(f"  cached {len(cache)} subjects in {time.time()-t_pre:.0f}s; "
      f"sample shape {tuple(cache[codes[0]].shape)}")

    skf = StratifiedKFold(n_splits=N_SPLITS, shuffle=True, random_state=42)

    curves = {}
    oof_p1 = np.full(len(codes), np.nan)   # P(CTRL)
    oof_fold = np.full(len(codes), -1, dtype=int)
    oof_code = list(codes)
    rows = []

    for fold, (train_idx, val_idx) in enumerate(skf.split(codes, labels)):
        p(f"\n===== FOLD {fold+1}/{N_SPLITS} | train={len(train_idx)} val={len(val_idx)} =====")
        tr_codes, tr_lab = codes[train_idx], labels[train_idx]
        va_codes, va_lab = codes[val_idx], labels[val_idx]

        train_ds = CachedDataset([cache[c] for c in tr_codes], tr_lab)
        val_ds = CachedDataset([cache[c] for c in va_codes], va_lab)
        train_loader = DataLoader(train_ds, batch_size=BATCH_SIZE, shuffle=True)
        val_loader = DataLoader(val_ds, batch_size=BATCH_SIZE, shuffle=False)

        model = ConvAutoencoderWithAttention(
            in_channels=1, latent_dim=LATENT_DIM, num_classes=num_classes).to(DEVICE)
        # FAITHFUL: optimizer built BEFORE the lazy encoder.fc & decoder exist,
        # so they are excluded from training. We additionally create them via a
        # dummy eval-mode forward (no BatchNorm update) and freeze requires_grad:
        # gradients still flow THROUGH their fixed init weights to the conv
        # layers, so trained params evolve identically -- this is just faster.
        optimizer = optim.Adam(model.parameters(), lr=LEARNING_RATE)
        with torch.no_grad():
            model.eval(); _ = model(train_ds[0][0].unsqueeze(0).to(DEVICE))
        for prm in model.encoder.fc.parameters():
            prm.requires_grad_(False)
        for prm in model.decoder.parameters():
            prm.requires_grad_(False)
        model.train()

        hist = {k: [] for k in ["tr_loss", "tr_recon", "tr_class", "tr_acc",
                                 "va_loss", "va_recon", "va_class", "va_acc"]}
        best_val_acc = -1.0; ckpt = os.path.join(MDL, f"best_model_fold{fold+1}.pth")
        t0 = time.time()
        for epoch in range(NUM_EPOCHS):
            trl, trr, trc, tra = train_epoch(model, train_loader, optimizer, DEVICE, ALPHA, BETA)
            val, var, vac, vaa = validate(model, val_loader, DEVICE, ALPHA, BETA)
            for k, v in zip(hist, [trl, trr, trc, tra, val, var, vac, vaa]):
                hist[k].append(v)
            if vaa > best_val_acc:
                best_val_acc = vaa
                torch.save(model.state_dict(), ckpt)
            if (epoch + 1) % 10 == 0:
                p(f"  ep{epoch+1:3d}: tr_acc={tra:5.1f} va_acc={vaa:5.1f} "
                  f"tr_loss={trl:.3f} va_loss={val:.3f} best_va_acc={best_val_acc:.1f} "
                  f"({time.time()-t0:.0f}s)")
        curves[f"fold{fold+1}"] = hist
        p(f"  FOLD {fold+1} best_val_acc={best_val_acc:.2f}%")

        # reload best-val-acc checkpoint, score val fold
        m2 = ConvAutoencoderWithAttention(in_channels=1, latent_dim=LATENT_DIM,
                                          num_classes=num_classes).to(DEVICE)
        with torch.no_grad():
            m2.eval(); _ = m2(next(iter(val_loader))[0].to(DEVICE))  # init lazy layers
        m2.load_state_dict(torch.load(ckpt, map_location=DEVICE), strict=True)
        p1, yv = predict_probs(m2, val_loader, DEVICE)  # p1 = P(CTRL); yv: 1=CTRL
        # store in clinical orientation: P(ALS) = 1 - P(CTRL), y_als = 1-yv
        for j, gi in enumerate(val_idx):
            oof_p1[gi] = p1[j]; oof_fold[gi] = fold + 1
        y_als = 1 - yv; pal = 1.0 - p1; pred_als = (pal >= 0.5).astype(int)
        try:
            auc = roc_auc_score(y_als, pal)
        except ValueError:
            auc = float("nan")
        rows.append(dict(fold=fold + 1, n=len(val_idx), best_val_acc=best_val_acc,
                         accuracy=accuracy_score(y_als, pred_als), auc=auc,
                         precision=precision_score(y_als, pred_als, zero_division=0),
                         recall=recall_score(y_als, pred_als, zero_division=0),
                         f1=f1_score(y_als, pred_als, zero_division=0)))
        p(f"  [fold{fold+1}] acc={rows[-1]['accuracy']:.3f} auc={auc:.3f}")

    # ---- assemble OOF in clinical orientation (positive = ALS) ----
    y_als_full = 1 - labels  # labels: ALS=0,CTRL=1 -> y_als: ALS=1
    oof_pal = 1.0 - oof_p1
    oof_pred = (oof_pal >= 0.5).astype(int)

    np.savez(os.path.join(OUT, "rerun_cache.npz"),
             oof_p=oof_pal, oof_fold=oof_fold, y=y_als_full, codes=np.array(oof_code))
    json.dump(curves, open(os.path.join(OUT, "curves.json"), "w"))

    fold_df = pd.DataFrame(rows)
    summ = pd.concat([fold_df, pd.DataFrame([
        dict(fold="mean", **{m: fold_df[m].mean() for m in ["accuracy", "auc", "precision", "recall", "f1"]}),
        dict(fold="std", **{m: fold_df[m].std() for m in ["accuracy", "auc", "precision", "recall", "f1"]}),
        dict(fold="OOF", n=len(y_als_full), accuracy=accuracy_score(y_als_full, oof_pred),
             auc=roc_auc_score(y_als_full, oof_pal),
             precision=precision_score(y_als_full, oof_pred, zero_division=0),
             recall=recall_score(y_als_full, oof_pred, zero_division=0),
             f1=f1_score(y_als_full, oof_pred, zero_division=0))])], ignore_index=True)
    summ.to_csv(os.path.join(OUT, "rerun_metrics.csv"), index=False)
    pd.DataFrame({"code": oof_code,
                  "true_label": ["ALS" if v else "CTRL" for v in y_als_full],
                  "held_out_fold": oof_fold, "P_ALS": np.round(oof_pal, 4),
                  "pred": ["ALS" if v else "CTRL" for v in oof_pred],
                  "correct": (oof_pred == y_als_full)}).to_csv(
        os.path.join(OUT, "rerun_oof_predictions.csv"), index=False)
    p("\n" + summ.round(3).to_string(index=False))
    p("\nDone. Outputs in", OUT)
    log.close()


if __name__ == "__main__":
    main()
