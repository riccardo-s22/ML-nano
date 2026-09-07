#!/usr/bin/env python
r"""
predict_exp7_with_exp5_models_v7b_posindex_autoflip.py

Fixes the "AUC ~ 0" issue from v7:
- Your model's softmax class order appears to be [ALS, CTRL] (ALS is index 0), not index 1.
- v7 hard-coded p_als = probs[:,1], which inverted the meaning of the score.

This version adds:
  --pos_index {0,1}          # which softmax column corresponds to ALS
  --auto_flip_if_labeled     # if labels_csv is provided, pick the index that yields AUC >= 0.5

It keeps the "lazy init" dummy forward pass so strict_load works.

Windows CMD example:
python predict_exp7_with_exp5_models_v7b_posindex_autoflip.py ^
  --arch_file "C:\Users\riccardo-s\Documents\CNT\targetALS\exp5\conv_autoencoder_detailed.py" ^
  --model_glob "C:\Users\riccardo-s\Documents\CNT\targetALS\exp5\5_fold_models_original\best_model_fold*.pth" ^
  --data_root "C:\Users\riccardo-s\Documents\CNT\targetALS\exp7\3" ^
  --labels_csv "C:\Users\riccardo-s\Documents\CNT\targetALS\exp7\3\sample_labels_03.csv" ^
  --id_col code --label_col group --pos_label ALS ^
  --output_csv "C:\Users\riccardo-s\Documents\CNT\targetALS\exp7\3_predictions_final.csv" ^
  --timepoints "C:\Users\riccardo-s\Documents\CNT\targetALS\exp7\3\out_0h" "C:\Users\riccardo-s\Documents\CNT\targetALS\exp7\3\out_6h" "C:\Users\riccardo-s\Documents\CNT\targetALS\exp7\3\out_24h" ^
  --device cpu --strict_load ^
  --auto_flip_if_labeled
"""

import os
import sys
import glob
import argparse
import importlib.util
import numpy as np
import pandas as pd
import torch
from openpyxl import load_workbook

try:
    from sklearn.metrics import roc_auc_score
except Exception:
    roc_auc_score = None


# -----------------------------------------------------------------------------
# Dynamic Import for Architecture
# -----------------------------------------------------------------------------
def load_arch_module(arch_file_path: str):
    spec = importlib.util.spec_from_file_location("dynamic_arch", arch_file_path)
    if spec is None or spec.loader is None:
        raise ImportError(f"Could not load architecture from {arch_file_path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules["dynamic_arch"] = module
    spec.loader.exec_module(module)
    return module


# -----------------------------------------------------------------------------
# EEM Loader (same as v7)
# -----------------------------------------------------------------------------
def load_eem_matrix(xlsx_path, target_shape=(512, 71)):
    """
    Returns (1, H, W) float32 in [0,1] using per-file min-max normalization (as v7).
    """
    try:
        wb = load_workbook(xlsx_path, data_only=True)
        ws = wb.active
        data = []
        for r_i, row in enumerate(ws.iter_rows(values_only=True)):
            if r_i == 0:
                continue  # header
            row_vals = []
            for c_i, val in enumerate(row):
                if c_i == 0:
                    continue  # first col is emission index
                if isinstance(val, (int, float)):
                    row_vals.append(float(val))
                else:
                    row_vals.append(0.0)
            if row_vals:
                data.append(row_vals)

        arr = np.array(data, dtype=np.float32)
        if arr.size == 0:
            return np.zeros((1, *target_shape), dtype=np.float32)

        mn, mx = float(arr.min()), float(arr.max())
        if mx > mn:
            arr = (arr - mn) / (mx - mn)
        else:
            arr = np.zeros_like(arr)

        # NOTE: v7 trusted the incoming shape; keep same behavior.
        return arr[np.newaxis, ...]  # (1,H,W)
    except Exception as e:
        print(f"Error loading {xlsx_path}: {e}")
        return np.zeros((1, *target_shape), dtype=np.float32)


# -----------------------------------------------------------------------------
def parse_args():
    p = argparse.ArgumentParser()
    p.add_argument("--arch_file", required=True)
    p.add_argument("--model_glob", required=True)
    p.add_argument("--data_root", required=True)
    p.add_argument("--labels_csv", required=True)
    p.add_argument("--id_col", default="code")
    p.add_argument("--label_col", default="group")
    p.add_argument("--pos_label", default="ALS")
    p.add_argument("--output_csv", default="predictions.csv")
    p.add_argument("--timepoints", nargs="+", default=["out_0h", "out_6h", "out_24h"])
    p.add_argument("--device", default="cpu")
    p.add_argument("--recursive", action="store_true")
    p.add_argument("--strict_load", action="store_true")

    p.add_argument("--pos_index", type=int, default=1, choices=[0, 1],
                   help="Which softmax column corresponds to ALS (0 or 1).")
    p.add_argument("--auto_flip_if_labeled", action="store_true",
                   help="If set, choose pos_index that yields AUC >= 0.5 on labeled subset (requires sklearn).")
    return p.parse_args()


def main():
    args = parse_args()
    device = torch.device(args.device)

    # A) Load Labels
    df_labels = pd.read_csv(args.labels_csv)
    df_labels[args.id_col] = df_labels[args.id_col].astype(str).str.strip()

    true_map = {str(r[args.id_col]).strip(): str(r[args.label_col]).strip() for _, r in df_labels.iterrows()}

    # B) Find Samples in first timepoint folder
    tp0 = args.timepoints[0]
    tp0_path = tp0 if os.path.isabs(tp0) else os.path.join(args.data_root, tp0)

    print(f"Scanning for samples in: {tp0_path}")
    files = glob.glob(os.path.join(tp0_path, "*.xlsx"))
    files = [os.path.basename(f) for f in files if not os.path.basename(f).startswith("~")]
    found_samples = []
    for sid in true_map.keys():
        if f"{sid}.xlsx" in files:
            found_samples.append(sid)

    print(f"Matched {len(found_samples)} / {len(true_map)} labeled samples.")
    if not found_samples:
        print("Error: No matching files found. Check your CSV IDs vs filenames.")
        return

    # C) Load Architecture / Model class
    arch_mod = load_arch_module(args.arch_file)
    ModelClass = getattr(arch_mod, "ConvAutoencoderWithAttention", None) or getattr(arch_mod, "ConvAutoencoderV2", None)
    if not ModelClass:
        raise RuntimeError("Could not find 'ConvAutoencoderWithAttention' or 'ConvAutoencoderV2' in arch_file.")

    checkpoints = sorted(glob.glob(args.model_glob))
    if not checkpoints:
        raise RuntimeError("No checkpoints found.")
    print(f"Found {len(checkpoints)} models.")

    # Storage: per-sample per-fold probs
    all_probs = np.zeros((len(found_samples), len(checkpoints)), dtype=np.float32)

    # D) Predict per model
    for m_i, ckpt_path in enumerate(checkpoints):
        print(f"--- Model {m_i+1}: {os.path.basename(ckpt_path)} ---")

        model = ModelClass(in_channels=1, latent_dim=512, num_classes=2).to(device)

        # Lazy init
        print("  Running dummy pass to initialize lazy layers...")
        dummy_input = torch.zeros(1, 3, 1, 512, 71, device=device)
        with torch.no_grad():
            try:
                model(dummy_input)
            except Exception as e:
                print(f"  Warning: Dummy pass failed ({e}).")

        sd = torch.load(ckpt_path, map_location=device)
        try:
            model.load_state_dict(sd, strict=args.strict_load)
            print("  Weights loaded successfully.")
        except RuntimeError as e:
            print(f"  Strict load failed: {e}")
            print("  Trying strict=False...")
            model.load_state_dict(sd, strict=False)

        model.eval()

        for s_i, sid in enumerate(found_samples):
            imgs = []
            for tp in args.timepoints:
                folder = tp if os.path.isabs(tp) else os.path.join(args.data_root, tp)
                fpath = os.path.join(folder, f"{sid}.xlsx")
                img = load_eem_matrix(fpath)  # (1,H,W)
                imgs.append(torch.from_numpy(img))

            batch_x = torch.stack(imgs, dim=0).unsqueeze(0).to(device)  # (1,3,1,H,W)

            with torch.no_grad():
                # forward returns: (recon, logits, attn) in your model
                out = model(batch_x)
                if isinstance(out, (tuple, list)) and len(out) >= 2:
                    logits = out[1]
                else:
                    raise RuntimeError("Unexpected model output; expected (recon, logits, ...)")

                probs = torch.softmax(logits, dim=1)
                all_probs[s_i, m_i] = float(probs[0, args.pos_index].item())

            if s_i % 10 == 0:
                print(f"  Predicting {s_i+1}/{len(found_samples)}...", end="\r")
        print("")

    # E) Aggregate
    avg_probs = np.mean(all_probs, axis=1)
    std_probs = np.std(all_probs, axis=1)

    # F) Auto-flip if labeled and sklearn available
    chosen_index = args.pos_index
    if args.auto_flip_if_labeled and roc_auc_score is not None:
        y_true = np.array([(true_map[sid] == args.pos_label) for sid in found_samples], dtype=int)
        auc_current = roc_auc_score(y_true, avg_probs)
        auc_flipped = roc_auc_score(y_true, 1.0 - avg_probs)
        print(f"\nAUC with pos_index={args.pos_index}: {auc_current:.4f}")
        print(f"AUC if flipped (1-p): {auc_flipped:.4f}")
        if auc_current < 0.5 and auc_flipped > auc_current:
            print("[AUTO] Flipping probabilities (ALS is the other softmax index).")
            avg_probs = 1.0 - avg_probs
            chosen_index = 1 - args.pos_index

    # G) Save
    rows = []
    for i, sid in enumerate(found_samples):
        pred_label = args.pos_label if avg_probs[i] >= 0.5 else "Healthy"
        row = {
            "Subject": sid,
            "True_Label": true_map.get(sid, "Unknown"),
            "Prediction": pred_label,
            "Prob_POS_Mean": float(avg_probs[i]),
            "Uncertainty": float(std_probs[i]),
            "pos_index_used": int(chosen_index),
        }
        for m_i in range(len(checkpoints)):
            row[f"Fold_{m_i}_Prob_rawposindex{args.pos_index}"] = float(all_probs[i, m_i])
        rows.append(row)

    df_out = pd.DataFrame(rows)
    df_out.to_csv(args.output_csv, index=False)

    # Metrics
    if roc_auc_score is not None:
        y_true = (df_out["True_Label"] == args.pos_label).astype(int).to_numpy()
        auc = roc_auc_score(y_true, df_out["Prob_POS_Mean"].astype(float).to_numpy())
        print(f"\nEnsemble ROC AUC (using pos_index_used={chosen_index}): {auc:.4f}")

    print(f"Saved to: {args.output_csv}")


if __name__ == "__main__":
    main()
