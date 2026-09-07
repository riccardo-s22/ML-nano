#!/usr/bin/env python3
"""
ALSFRS / ALSFRS-slope Chirality Regression Pipeline (No PyTorch Required)
==========================================================================

Extracts features from DNA-SWCNT chirality positions in EEM spectral matrices,
then regresses against continuous ALSFRS score and ALSFRS-slope.

APPROACH: Chirality Direct Features
  - 12 chirality positions × 3 patch stats × 3 timepoints = 108 base features
  - + temporal deltas (6h-0h, 24h-0h, 24h-6h) and ratios → ~288 raw features
  - In-fold f-regression + collinearity pruning per LOOCV fold
  - Regressors: Ridge, SVR, ElasticNet, RF, GBR

LOOCV on 20 ALS subjects. Metrics: Pearson r, Spearman ρ, R², MAE, RMSE.
"""

import os, sys, json, re
from pathlib import Path
from typing import Dict, List, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy import stats as scipy_stats

from sklearn.linear_model import Ridge, ElasticNet
from sklearn.svm import SVR
from sklearn.ensemble import RandomForestRegressor, GradientBoostingRegressor
from sklearn.feature_selection import f_regression
from sklearn.model_selection import LeaveOneOut, KFold
from sklearn.preprocessing import StandardScaler
from sklearn.pipeline import Pipeline
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score

import warnings
warnings.filterwarnings("ignore")

SEED = 42
np.random.seed(SEED)

# =============================================================================
# DNA-SWCNT Chirality Positions
# =============================================================================
CHIRALITY_POSITIONS = [
    {"name": "(8,3)",  "em": 973.98,  "exc": 673.94, "height": 0.480},
    {"name": "(6,5)",  "em": 987.82,  "exc": 577.12, "height": 0.762},
    {"name": "(7,5)",  "em": 1047.81, "exc": 653.32, "height": 0.648},
    {"name": "(10,2)", "em": 1080.60, "exc": 745.92, "height": 0.389},
    {"name": "(9,4)",  "em": 1131.96, "exc": 731.39, "height": 0.553},
    {"name": "(8,4)",  "em": 1130.34, "exc": 599.78, "height": 0.591},
    {"name": "(7,6)",  "em": 1138.19, "exc": 659.79, "height": 0.725},
    {"name": "(8,6)",  "em": 1200.03, "exc": 727.40, "height": 0.318},
    {"name": "(8,7)",  "em": 1288.27, "exc": 740.87, "height": 0.134},
    {"name": "(9,5)",  "em": 1262.98, "exc": 685.15, "height": 0.190},
    {"name": "(10,3)", "em": 1267.70, "exc": 648.97, "height": 0.190},
    {"name": "(10,5)", "em": 1282.97, "exc": 801.23, "height": 0.100},
]

# =============================================================================
# Data Loading
# =============================================================================
def load_excel_as_array(filepath: str) -> np.ndarray:
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    data = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx == 0:
            continue
        row_vals = []
        for col_idx, cell in enumerate(row):
            if col_idx == 0:
                continue
            if isinstance(cell, (int, float)) and cell is not None:
                row_vals.append(float(cell))
            else:
                row_vals.append(np.nan)
        data.append(row_vals)
    arr = np.array(data, dtype=np.float32)
    if np.isnan(arr).any():
        col_med = np.nanmedian(arr, axis=0)
        inds = np.where(np.isnan(arr))
        arr[inds] = np.take(col_med, inds[1])
    mn, mx = float(np.min(arr)), float(np.max(arr))
    if mx > mn:
        arr = (arr - mn) / (mx - mn)
    return arr


def load_spectral_axes(filepath: str) -> Tuple[np.ndarray, np.ndarray]:
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    ex_re = re.compile(r"Excitation_\s*0*([0-9]+(?:\.[0-9]+)?)", re.IGNORECASE)
    excitation = []
    for col in header[1:]:
        if col is None:
            continue
        m = ex_re.search(str(col))
        if m:
            excitation.append(float(m.group(1)))
    emission = []
    for r in rows[1:]:
        if r[0] is not None:
            try:
                emission.append(float(r[0]))
            except:
                pass
    return np.array(emission), np.array(excitation)


def load_all_samples(codes: List[str], tp_dirs: List[str]) -> np.ndarray:
    X_list = []
    for code in codes:
        tp_mats = []
        for tp in tp_dirs:
            fp = os.path.join(tp, f"{code}.xlsx")
            tp_mats.append(load_excel_as_array(fp))
        X_list.append(np.stack(tp_mats, axis=0))
    return np.stack(X_list, axis=0)


# =============================================================================
# Clinical Data Loading
# =============================================================================
def load_clinical_data(alsfrs_path: str, slope_path: str) -> pd.DataFrame:
    """Load ALSFRS scores and slopes, return merged DataFrame with 'code' column."""
    wb = load_workbook(alsfrs_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c else f"col{i}" for i, c in enumerate(rows[0])]
    alsfrs_df = pd.DataFrame(rows[1:], columns=header)

    # Find score column
    score_col = None
    for c in alsfrs_df.columns:
        if 'score' in c.lower() or 'total' in c.lower() or 'alsfrs' in c.lower():
            score_col = c
            break
    if score_col is None:
        # fallback: use last numeric column
        for c in reversed(alsfrs_df.columns):
            if pd.to_numeric(alsfrs_df[c], errors='coerce').notna().sum() > 5:
                score_col = c
                break
    
    # Find code column
    code_col = None
    for c in alsfrs_df.columns:
        if 'code' in c.lower() or 'sample' in c.lower() or 'id' in c.lower():
            code_col = c
            break
    if code_col is None:
        code_col = alsfrs_df.columns[0]

    alsfrs_df['code'] = alsfrs_df[code_col].astype(str).str.strip()
    alsfrs_df['ALSFRS'] = pd.to_numeric(alsfrs_df[score_col], errors='coerce')
    alsfrs_df = alsfrs_df[['code', 'ALSFRS']].dropna()

    slope_df = pd.read_csv(slope_path)
    slope_df.columns = [c.strip() for c in slope_df.columns]
    # Find slope column
    slope_col = None
    for c in slope_df.columns:
        if 'slope' in c.lower():
            slope_col = c
            break
    if slope_col is None:
        slope_col = slope_df.columns[-1]
    
    code_col2 = None
    for c in slope_df.columns:
        if 'code' in c.lower() or 'sample' in c.lower() or 'id' in c.lower():
            code_col2 = c
            break
    if code_col2 is None:
        code_col2 = slope_df.columns[0]

    slope_df['code'] = slope_df[code_col2].astype(str).str.strip()
    slope_df['ALSFRS_slope'] = pd.to_numeric(slope_df[slope_col], errors='coerce')
    slope_df = slope_df[['code', 'ALSFRS_slope']].dropna()

    return alsfrs_df.merge(slope_df, on='code', how='outer')


# =============================================================================
# Chirality Feature Extraction
# =============================================================================
def find_nearest_pixel(target_nm: float, axis_nm: np.ndarray) -> int:
    return int(np.argmin(np.abs(axis_nm - target_nm)))


def extract_chirality_features(
    X_np: np.ndarray,
    emission_axis: np.ndarray,
    excitation_axis: np.ndarray,
    radius: int = 3,
) -> Tuple[pd.DataFrame, List[str]]:
    """Extract patch statistics at each chirality position."""
    N, T, H, W = X_np.shape
    tp_names = ['0h', '6h', '24h']
    stat_funcs = {'mean': np.mean, 'max': np.max, 'std': np.std}
    features = {}

    for ch in CHIRALITY_POSITIONS:
        em_idx = find_nearest_pixel(ch['em'], emission_axis)
        exc_idx = find_nearest_pixel(ch['exc'], excitation_axis)
        ch_name = ch['name'].replace('(', '').replace(')', '').replace(',', '_')

        r0, r1 = max(0, em_idx - radius), min(H, em_idx + radius + 1)
        c0, c1 = max(0, exc_idx - radius), min(W, exc_idx + radius + 1)

        patches = {}
        for t in range(T):
            patch = X_np[:, t, r0:r1, c0:c1]
            patches[tp_names[t]] = patch
            for stat_name, stat_fn in stat_funcs.items():
                features[f"ch{ch_name}_{tp_names[t]}_{stat_name}"] = stat_fn(patch, axis=(1, 2))

        # Temporal deltas
        for stat_name, stat_fn in stat_funcs.items():
            v0 = stat_fn(patches['0h'], axis=(1, 2))
            v6 = stat_fn(patches['6h'], axis=(1, 2))
            v24 = stat_fn(patches['24h'], axis=(1, 2))

            features[f"ch{ch_name}_d6h0h_{stat_name}"] = v6 - v0
            features[f"ch{ch_name}_d24h0h_{stat_name}"] = v24 - v0
            features[f"ch{ch_name}_d24h6h_{stat_name}"] = v24 - v6

            # Ratios (safe division)
            eps = 1e-10
            features[f"ch{ch_name}_r6h0h_{stat_name}"] = v6 / (np.abs(v0) + eps)
            features[f"ch{ch_name}_r24h0h_{stat_name}"] = v24 / (np.abs(v0) + eps)

    feat_df = pd.DataFrame(features)
    return feat_df, list(feat_df.columns)


# =============================================================================
# Feature Selection (in-fold)
# =============================================================================
def select_features_regression(
    X_train: np.ndarray,
    y_train: np.ndarray,
    fanova_k: int = 25,
    collinear_thresh: float = 0.85,
    min_keep: int = 3,
) -> List[int]:
    """In-fold feature selection: top-k by f_regression, then collinearity pruning."""
    if X_train.shape[1] == 0:
        return []
    
    # Remove constant features
    var = np.var(X_train, axis=0)
    valid = var > 1e-12
    if not valid.any():
        return []
    
    F = np.zeros(X_train.shape[1])
    p = np.ones(X_train.shape[1])
    F_ok, p_ok = f_regression(X_train[:, valid], y_train)
    F[valid] = F_ok
    p[valid] = p_ok

    order = np.argsort(-F)
    order = order[:min(fanova_k, len(order))]

    # Greedy collinearity pruning
    keep = []
    for idx in order:
        idx = int(idx)
        if len(keep) == 0:
            keep.append(idx)
            continue
        x = X_train[:, idx]
        ok = True
        for j in keep:
            r = np.corrcoef(x, X_train[:, j])[0, 1]
            if np.isnan(r):
                continue
            if abs(r) > collinear_thresh:
                ok = False
                break
        if ok:
            keep.append(idx)

    if len(keep) < min_keep:
        keep = [int(i) for i in order[:min(min_keep, len(order))]]
    return sorted(set(keep))


# =============================================================================
# Regressors
# =============================================================================
def get_regressors():
    return {
        "Ridge": Pipeline([
            ("scaler", StandardScaler()),
            ("reg", Ridge(alpha=10.0, random_state=SEED)),
        ]),
        "SVR": Pipeline([
            ("scaler", StandardScaler()),
            ("reg", SVR(kernel='rbf', C=1.0, epsilon=0.1)),
        ]),
        "ElasticNet": Pipeline([
            ("scaler", StandardScaler()),
            ("reg", ElasticNet(alpha=0.5, l1_ratio=0.5, max_iter=5000, random_state=SEED)),
        ]),
        "RF": RandomForestRegressor(
            n_estimators=200, max_depth=4, random_state=SEED, n_jobs=-1
        ),
        "GBR": GradientBoostingRegressor(
            n_estimators=100, max_depth=2, learning_rate=0.05, random_state=SEED
        ),
    }


# =============================================================================
# Evaluation
# =============================================================================
def evaluate_regression(y_true, y_pred):
    """Compute regression metrics."""
    n = len(y_true)
    if n < 3 or np.std(y_pred) < 1e-12:
        return {"pearson_r": 0.0, "spearman_rho": 0.0, "r2": -999, 
                "mae": 999, "rmse": 999}
    r, p_r = scipy_stats.pearsonr(y_true, y_pred)
    rho, p_rho = scipy_stats.spearmanr(y_true, y_pred)
    r2 = r2_score(y_true, y_pred)
    mae = mean_absolute_error(y_true, y_pred)
    rmse = np.sqrt(mean_squared_error(y_true, y_pred))
    return {
        "pearson_r": float(r), "pearson_p": float(p_r),
        "spearman_rho": float(rho), "spearman_p": float(p_rho),
        "r2": float(r2), "mae": float(mae), "rmse": float(rmse),
    }


def run_loocv_regression(
    X_feat: np.ndarray,
    y: np.ndarray,
    feat_names: List[str],
    fanova_k: int = 25,
    collinear_thresh: float = 0.85,
) -> pd.DataFrame:
    """Run LOOCV regression with in-fold feature selection."""
    N = len(y)
    regressors = get_regressors()
    loo = LeaveOneOut()

    predictions = {name: np.full(N, np.nan) for name in regressors}
    fs_counts = {fn: 0 for fn in feat_names}

    for train_idx, test_idx_arr in loo.split(np.arange(N)):
        test_idx = int(test_idx_arr[0])
        X_tr, X_te = X_feat[train_idx], X_feat[test_idx:test_idx+1]
        y_tr = y[train_idx]

        # In-fold feature selection
        keep = select_features_regression(X_tr, y_tr, fanova_k=fanova_k, 
                                          collinear_thresh=collinear_thresh)
        if len(keep) == 0:
            keep = list(range(min(5, X_feat.shape[1])))

        for i in keep:
            if i < len(feat_names):
                fs_counts[feat_names[i]] += 1

        X_tr_sel = X_tr[:, keep]
        X_te_sel = X_te[:, keep]

        for name, reg in regressors.items():
            try:
                reg.fit(X_tr_sel, y_tr)
                predictions[name][test_idx] = float(reg.predict(X_te_sel)[0])
            except Exception:
                predictions[name][test_idx] = float(np.mean(y_tr))

    # Evaluate
    results = []
    for name in regressors:
        metrics = evaluate_regression(y, predictions[name])
        metrics['model'] = name
        results.append(metrics)

    return pd.DataFrame(results), predictions, fs_counts


# =============================================================================
# Feature sets: different combinations of base/delta/ratio features
# =============================================================================
def split_feature_sets(feat_df: pd.DataFrame) -> Dict[str, List[str]]:
    """Split features into meaningful subsets."""
    all_cols = list(feat_df.columns)
    
    base_cols = [c for c in all_cols if not ('_d6h' in c or '_d24h' in c or '_r6h' in c or '_r24h' in c)]
    delta_cols = [c for c in all_cols if '_d6h' in c or '_d24h' in c]
    ratio_cols = [c for c in all_cols if '_r6h' in c or '_r24h' in c]
    
    # Per-timepoint
    tp0_cols = [c for c in base_cols if '_0h_' in c]
    tp6_cols = [c for c in base_cols if '_6h_' in c]
    tp24_cols = [c for c in base_cols if '_24h_' in c]
    
    return {
        "all": all_cols,
        "base_only": base_cols,
        "deltas_only": delta_cols,
        "base+deltas": base_cols + delta_cols,
        "tp0h_only": tp0_cols,
        "tp6h_only": tp6_cols,
        "tp24h_only": tp24_cols,
        "tp6h+deltas": tp6_cols + delta_cols,
    }


# =============================================================================
# Main
# =============================================================================
def main():
    # Paths
    tp_dirs = ['/home/claude/tp_0h', '/home/claude/tp_6h', '/home/claude/tp_24h']
    labels_csv = '/home/claude/sample_labels_filecodes.csv'
    alsfrs_csv = '/home/claude/ALSFRS_fixed.xlsx'
    slope_csv = '/home/claude/ALSFRS_slope_fixed.csv'
    output_dir = Path('/home/claude/chirality_regression_results')
    output_dir.mkdir(exist_ok=True)

    print("=" * 70)
    print("CHIRALITY REGRESSION PIPELINE")
    print("=" * 70)

    # Load labels
    labels_df = pd.read_csv(labels_csv)
    all_codes = labels_df['code'].astype(str).tolist()
    all_groups = labels_df['group'].tolist()
    print(f"Total samples: {len(all_codes)}")

    # Load clinical data
    clinical_df = load_clinical_data(alsfrs_csv, slope_csv)
    print(f"Clinical data loaded: {len(clinical_df)} rows")
    print(clinical_df.head())

    # Map clinical codes to file codes
    # Clinical codes use dots (e.g., "1.3d__570.opj_with_emission")
    # File codes use underscores (e.g., "1_3d__570_opj_with_emission")
    clinical_df['file_code'] = clinical_df['code'].str.replace('.', '_', regex=False)
    
    # Find ALS subjects with clinical data
    als_codes = [c for c, g in zip(all_codes, all_groups) if g.upper() == 'ALS']
    print(f"ALS codes in labels: {len(als_codes)}")
    
    # Match
    als_with_clinical = []
    y_alsfrs = []
    y_slope = []
    for code in als_codes:
        match = clinical_df[clinical_df['file_code'] == code]
        if len(match) == 0:
            # Try partial match
            match = clinical_df[clinical_df['file_code'].str.contains(code[:5], na=False)]
        if len(match) > 0:
            row = match.iloc[0]
            if pd.notna(row.get('ALSFRS')) or pd.notna(row.get('ALSFRS_slope')):
                als_with_clinical.append(code)
                y_alsfrs.append(row.get('ALSFRS', np.nan))
                y_slope.append(row.get('ALSFRS_slope', np.nan))

    y_alsfrs = np.array(y_alsfrs, dtype=float)
    y_slope = np.array(y_slope, dtype=float)
    print(f"\nALS subjects with clinical data: {len(als_with_clinical)}")
    
    alsfrs_valid = ~np.isnan(y_alsfrs)
    slope_valid = ~np.isnan(y_slope)
    print(f"  ALSFRS valid: {alsfrs_valid.sum()} (range: {np.nanmin(y_alsfrs):.1f}-{np.nanmax(y_alsfrs):.1f})")
    print(f"  Slope valid:  {slope_valid.sum()} (range: {np.nanmin(y_slope):.1f}-{np.nanmax(y_slope):.1f})")

    if len(als_with_clinical) < 5:
        print("ERROR: Too few ALS subjects with clinical data.")
        sys.exit(1)

    # Load spectral data for ALS subjects
    print("\nLoading spectral data...")
    X_als = load_all_samples(als_with_clinical, tp_dirs)
    N, T, H, W = X_als.shape
    print(f"  Shape: {X_als.shape}")

    # Get spectral axes
    first_file = os.path.join(tp_dirs[0], f"{als_with_clinical[0]}.xlsx")
    emission_axis, excitation_axis = load_spectral_axes(first_file)
    print(f"  Emission: {len(emission_axis)} pts ({emission_axis[0]:.1f}-{emission_axis[-1]:.1f} nm)")
    print(f"  Excitation: {len(excitation_axis)} pts ({excitation_axis[0]:.1f}-{excitation_axis[-1]:.1f} nm)")

    # Verify chirality positions are within grid
    print("\nChirality position mapping:")
    for ch in CHIRALITY_POSITIONS:
        em_idx = find_nearest_pixel(ch['em'], emission_axis)
        exc_idx = find_nearest_pixel(ch['exc'], excitation_axis)
        em_actual = emission_axis[em_idx] if em_idx < len(emission_axis) else -1
        exc_actual = excitation_axis[exc_idx] if exc_idx < len(excitation_axis) else -1
        in_bounds = (0 <= em_idx < H) and (0 <= exc_idx < W)
        print(f"  {ch['name']:>8s}: em={ch['em']:.1f}→idx={em_idx}({em_actual:.1f}nm), "
              f"exc={ch['exc']:.1f}→idx={exc_idx}({exc_actual:.1f}nm) {'✓' if in_bounds else '✗ OUT OF BOUNDS'}")

    # Extract chirality features
    print("\nExtracting chirality features (radius=3)...")
    feat_df, feat_names = extract_chirality_features(X_als, emission_axis, excitation_axis, radius=3)
    print(f"  Total features: {len(feat_names)}")
    
    # Remove features with zero variance
    var = feat_df.var()
    good = var[var > 1e-12].index.tolist()
    feat_df = feat_df[good]
    feat_names = good
    print(f"  After removing zero-variance: {len(feat_names)}")

    # Feature sets
    feature_sets = split_feature_sets(feat_df)
    
    # Save feature matrix
    feat_save = feat_df.copy()
    feat_save.insert(0, 'code', als_with_clinical)
    feat_save.to_csv(output_dir / 'chirality_features_als.csv', index=False)

    # =============================================================================
    # Run regressions for each target × feature set
    # =============================================================================
    targets = {}
    if alsfrs_valid.sum() >= 10:
        targets['ALSFRS'] = (y_alsfrs, alsfrs_valid)
    if slope_valid.sum() >= 10:
        targets['ALSFRS_slope'] = (y_slope, slope_valid)

    all_results = []

    for target_name, (y_full, valid_mask) in targets.items():
        print(f"\n{'='*70}")
        print(f"TARGET: {target_name}")
        print(f"{'='*70}")
        
        idx_valid = np.where(valid_mask)[0]
        y_target = y_full[idx_valid]
        
        for fs_name, fs_cols in feature_sets.items():
            # Filter to available columns
            fs_cols_avail = [c for c in fs_cols if c in feat_df.columns]
            if len(fs_cols_avail) < 3:
                continue
            
            X_fs = feat_df[fs_cols_avail].values[idx_valid]
            
            print(f"\n  Feature set: {fs_name} ({len(fs_cols_avail)} features, {len(y_target)} samples)")
            
            results_df, predictions, fs_counts = run_loocv_regression(
                X_fs, y_target, fs_cols_avail,
                fanova_k=25, collinear_thresh=0.85,
            )
            
            results_df['target'] = target_name
            results_df['feature_set'] = fs_name
            results_df['n_features'] = len(fs_cols_avail)
            results_df['n_samples'] = len(y_target)
            
            all_results.append(results_df)
            
            # Print best
            best = results_df.loc[results_df['pearson_r'].abs().idxmax()]
            print(f"    Best: {best['model']} | r={best['pearson_r']:.3f} (p={best.get('pearson_p', 0):.3f}) | "
                  f"ρ={best['spearman_rho']:.3f} | R²={best['r2']:.3f} | MAE={best['mae']:.2f}")
            
            # Save predictions for best model
            best_name = best['model']
            pred_df = pd.DataFrame({
                'code': [als_with_clinical[i] for i in idx_valid],
                'y_true': y_target,
                'y_pred': predictions[best_name][~np.isnan(predictions[best_name])] if not np.isnan(predictions[best_name]).all() else predictions[best_name],
            })
            safe_name = fs_name.replace('+', '_')
            pred_df.to_csv(output_dir / f'predictions_{target_name}_{safe_name}_{best_name}.csv', index=False)
            
            # Save feature selection counts for this combo
            top_fs = sorted(fs_counts.items(), key=lambda x: -x[1])[:20]
            if top_fs:
                fs_df = pd.DataFrame(top_fs, columns=['feature', 'count'])
                fs_df['fraction'] = fs_df['count'] / len(y_target)
                fs_df.to_csv(output_dir / f'fs_counts_{target_name}_{safe_name}.csv', index=False)

    # Combine all results
    if all_results:
        combined = pd.concat(all_results, ignore_index=True)
        combined = combined.sort_values('pearson_r', ascending=False, key=abs)
        combined.to_csv(output_dir / 'all_results.csv', index=False)
        
        print(f"\n{'='*70}")
        print("TOP 15 RESULTS (by |Pearson r|)")
        print(f"{'='*70}")
        top15 = combined.head(15)
        for _, row in top15.iterrows():
            print(f"  {row['target']:>12s} | {row['feature_set']:>15s} | {row['model']:>10s} | "
                  f"r={row['pearson_r']:+.3f} (p={row.get('pearson_p', 0):.3f}) | "
                  f"ρ={row['spearman_rho']:+.3f} | R²={row['r2']:.3f} | MAE={row['mae']:.2f}")

        # Per-target summary
        for target_name in targets:
            print(f"\n  Best for {target_name}:")
            sub = combined[combined['target'] == target_name].head(5)
            for _, row in sub.iterrows():
                print(f"    {row['feature_set']:>15s} | {row['model']:>10s} | r={row['pearson_r']:+.3f} | ρ={row['spearman_rho']:+.3f}")

    print(f"\nResults saved to: {output_dir}")


if __name__ == "__main__":
    main()
