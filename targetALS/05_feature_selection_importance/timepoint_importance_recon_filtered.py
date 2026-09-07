#!/usr/bin/env python3
"""
TIMEPOINT-SPECIFIC FEATURE IMPORTANCE PIPELINE (Pre-trained Models)
WITH RECONSTRUCTION ERROR FILTERING
=========================================================================

This script loads pre-trained 5-fold CV models and computes feature importance
for each timepoint using gradient saliency and linear probing.

KEY IMPROVEMENT: Filters importance maps using reconstruction error mask.
Only regions where the model achieves low reconstruction error (MSE < threshold)
are considered reliable for interpretation.

WORKFLOW:
---------
1. Load 5 pre-trained models from disk
2. For EACH model:
   a) Compute reconstruction error per timepoint (input vs. reconstructed)
   b) Extract latent features → rank by classification importance → TOP 20
   c) For each top-20 dimension × each timepoint (0h, 6h, 24h):
      - Gradient Saliency: |∂z_d/∂x_t|
      - Linear Probing: Ridge weights
   d) Consensus = average(gradient, linear) per timepoint
3. Cross-model consensus:
   - Average importance across 5 models per timepoint
   - Average reconstruction error across 5 models per timepoint
   - Create mask: MSE < threshold (default 0.05)
   - Apply mask to importance maps
4. Extract ROIs from MASKED importance maps at 90th percentile

USAGE:
------
python timepoint_importance_recon_filtered.py \
    --model_dir C:/Users/.../5_fold_models_original \
    --tp_dirs C:/Users/.../out_0h,C:/Users/.../out_6h,C:/Users/.../out_24h \
    --labels_csv C:/Users/.../sample_labels.csv \
    --output_dir ./importance_results \
    --mse_threshold 0.05
"""

import os
import sys
import re
import argparse
import json
from pathlib import Path
from typing import List, Tuple, Dict, Optional
from collections import defaultdict

import numpy as np
import pandas as pd
import torch
import torch.nn as nn
import torch.nn.functional as F
from torch.utils.data import Dataset, DataLoader
from sklearn.linear_model import Ridge, LogisticRegression
from sklearn.preprocessing import StandardScaler
from sklearn.model_selection import StratifiedKFold
from scipy.stats import rankdata
from scipy import ndimage
from scipy.ndimage import gaussian_filter
import matplotlib.pyplot as plt
import seaborn as sns
from openpyxl import load_workbook
import warnings
warnings.filterwarnings('ignore')

# Reproducibility
SEED = 42
np.random.seed(SEED)
torch.manual_seed(SEED)


# ==============================================================================
# MODEL ARCHITECTURE (must match the saved models exactly)
# ==============================================================================

class ConvEncoder(nn.Module):
    """Convolutional encoder - MUST match original architecture."""
    
    def __init__(self, in_channels: int, latent_dim: int):
        super(ConvEncoder, self).__init__()
        
        self.conv1 = nn.Sequential(
            nn.Conv2d(in_channels, 32, kernel_size=3, padding=1),
            nn.BatchNorm2d(32),
            nn.ReLU(),
            nn.MaxPool2d(2, 2)
        )
        self.conv2 = nn.Sequential(
            nn.Conv2d(32, 64, kernel_size=3, padding=1),
            nn.BatchNorm2d(64),
            nn.ReLU(),
            nn.MaxPool2d(2, 2)
        )
        self.conv3 = nn.Sequential(
            nn.Conv2d(64, 128, kernel_size=3, padding=1),
            nn.BatchNorm2d(128),
            nn.ReLU(),
            nn.MaxPool2d(2, 2)
        )
        self.conv4 = nn.Sequential(
            nn.Conv2d(128, 256, kernel_size=3, padding=1),
            nn.BatchNorm2d(256),
            nn.ReLU()
        )
        
        self.flatten_size = None
        self.fc = None
        self.latent_dim = latent_dim
        self.encoded_shape = None
        
    def forward(self, x):
        x = self.conv1(x)
        x = self.conv2(x)
        x = self.conv3(x)
        x = self.conv4(x)
        
        if self.encoded_shape is None:
            self.encoded_shape = (x.shape[1], x.shape[2], x.shape[3])
        
        if self.fc is None:
            self.flatten_size = x.shape[1] * x.shape[2] * x.shape[3]
            self.fc = nn.Linear(self.flatten_size, self.latent_dim).to(x.device)
        
        x = x.view(x.size(0), -1)
        x = self.fc(x)
        return x


class ConvDecoder(nn.Module):
    """Convolutional decoder - MUST match original architecture."""
    
    def __init__(self, latent_dim: int, out_channels: int, 
                 encoded_shape: Tuple[int, int, int], target_shape: Tuple[int, int]):
        super(ConvDecoder, self).__init__()
        
        self.encoded_channels, self.encoded_h, self.encoded_w = encoded_shape
        self.target_h, self.target_w = target_shape
        
        self.fc = nn.Linear(latent_dim, self.encoded_channels * self.encoded_h * self.encoded_w)
        
        self.deconv1 = nn.Sequential(
            nn.ConvTranspose2d(256, 128, kernel_size=3, stride=2, padding=1, output_padding=1),
            nn.BatchNorm2d(128),
            nn.ReLU()
        )
        self.deconv2 = nn.Sequential(
            nn.ConvTranspose2d(128, 64, kernel_size=3, stride=2, padding=1, output_padding=1),
            nn.BatchNorm2d(64),
            nn.ReLU()
        )
        self.deconv3 = nn.Sequential(
            nn.ConvTranspose2d(64, 32, kernel_size=3, stride=2, padding=1, output_padding=1),
            nn.BatchNorm2d(32),
            nn.ReLU()
        )
        self.final_conv = nn.Conv2d(32, out_channels, kernel_size=3, padding=1)
        self.sigmoid = nn.Sigmoid()
        
    def forward(self, z):
        x = self.fc(z)
        x = x.view(x.size(0), self.encoded_channels, self.encoded_h, self.encoded_w)
        x = self.deconv1(x)
        x = self.deconv2(x)
        x = self.deconv3(x)
        x = self.final_conv(x)
        
        current_h, current_w = x.shape[2], x.shape[3]
        if current_h != self.target_h or current_w != self.target_w:
            x = torch.nn.functional.interpolate(
                x, size=(self.target_h, self.target_w), 
                mode='bilinear', align_corners=False
            )
        
        x = self.sigmoid(x)
        return x


class AttentionAggregation(nn.Module):
    """Attention mechanism - MUST match original architecture."""
    
    def __init__(self, latent_dim: int):
        super(AttentionAggregation, self).__init__()
        self.attention_net = nn.Sequential(
            nn.Linear(latent_dim, 64),
            nn.Tanh(),
            nn.Linear(64, 1)
        )
        
    def forward(self, latent_vectors):
        attention_scores = self.attention_net(latent_vectors)
        attention_scores = attention_scores.squeeze(-1)
        attention_weights = torch.softmax(attention_scores, dim=1)
        attention_weights_expanded = attention_weights.unsqueeze(-1)
        aggregated = torch.sum(latent_vectors * attention_weights_expanded, dim=1)
        return aggregated, attention_weights


class Classifier(nn.Module):
    """Classification head - MUST match original architecture."""
    
    def __init__(self, latent_dim: int, num_classes: int, dropout_rate: float = 0.3):
        super(Classifier, self).__init__()
        self.classifier = nn.Sequential(
            nn.Linear(latent_dim, 128),
            nn.ReLU(),
            nn.Dropout(dropout_rate),
            nn.Linear(128, num_classes)
        )
        
    def forward(self, z_agg):
        return self.classifier(z_agg)


class ConvAutoencoderWithAttention(nn.Module):
    """Complete model - MUST match original architecture."""
    
    def __init__(self, in_channels: int, latent_dim: int, num_classes: int):
        super(ConvAutoencoderWithAttention, self).__init__()
        
        self.encoder = ConvEncoder(in_channels, latent_dim)
        self.attention = AttentionAggregation(latent_dim)
        self.classifier = Classifier(latent_dim, num_classes)
        self.decoder = None
        self.latent_dim = latent_dim
        self.in_channels = in_channels
        
    def init_decoder(self, encoded_shape: tuple, target_shape: tuple):
        self.decoder = ConvDecoder(self.latent_dim, self.in_channels, 
                                   encoded_shape, target_shape)


# ==============================================================================
# DATA LOADING
# ==============================================================================

class ExcelImageDataset(Dataset):
    """Dataset for loading Excel spectral data."""
    
    def __init__(self, codes: List[str], labels: List[int], 
                 timepoint_dirs: List[str], file_suffix: str = ".xlsx"):
        self.codes = codes
        self.labels = labels
        self.timepoint_dirs = timepoint_dirs
        self.file_suffix = file_suffix
        self._cache = {}
        
    def __len__(self):
        return len(self.codes)
    
    def load_excel_as_image(self, filepath: str) -> np.ndarray:
        """Load Excel and convert to normalized numpy array."""
        if filepath in self._cache:
            return self._cache[filepath]
            
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active
        
        data = []
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if row_idx == 0:
                continue
            
            row_data = []
            for col_idx, cell in enumerate(row):
                if col_idx == 0:
                    continue
                if isinstance(cell, (int, float)) and cell is not None:
                    row_data.append(float(cell))
                else:
                    row_data.append(0.0)
            
            if len(row_data) > 0:
                data.append(row_data)
        
        arr = np.array(data, dtype=np.float32)
        
        arr_min, arr_max = arr.min(), arr.max()
        if arr_max > arr_min:
            arr = (arr - arr_min) / (arr_max - arr_min)
        
        self._cache[filepath] = arr
        return arr
    
    def find_file(self, directory: str, code: str) -> str:
        """Find file with given code."""
        for suffix in [self.file_suffix, '.xlsx', '.xls']:
            path = os.path.join(directory, f"{code}{suffix}")
            if os.path.exists(path):
                return path
        
        for f in os.listdir(directory):
            if f.startswith(code) and (f.endswith('.xlsx') or f.endswith('.xls')):
                return os.path.join(directory, f)
        
        raise FileNotFoundError(f"No file found for code {code} in {directory}")
    
    def __getitem__(self, idx):
        code = self.codes[idx]
        label = self.labels[idx]
        
        timepoint_data = []
        for tp_dir in self.timepoint_dirs:
            filepath = self.find_file(tp_dir, code)
            arr = self.load_excel_as_image(filepath)
            timepoint_data.append(arr)
        
        stacked = np.stack(timepoint_data, axis=0)
        tensor = torch.from_numpy(stacked).unsqueeze(1).float()
        
        return tensor, label, code


def load_spectral_axes(filepath: str) -> Tuple[np.ndarray, np.ndarray]:
    """
    Load emission and excitation wavelengths from Excel.
    
    Handles various header formats:
    - Pure numeric: 500, 510, 520...
    - String with number: "Excitation_500", "exc_510", "500nm"...
    - Mixed formats
    """
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    
    excitation_axis = []
    emission_axis = []
    
    def extract_number(value):
        """Extract numeric value from various formats."""
        if value is None:
            return None
        if isinstance(value, (int, float)):
            return float(value)
        # Try to extract number from string (e.g., "Excitation_500" -> 500)
        if isinstance(value, str):
            # First try direct conversion
            try:
                return float(value)
            except ValueError:
                pass
            # Try to find a number in the string
            matches = re.findall(r'[\d.]+', value)
            if matches:
                try:
                    # Take the first number found (usually the wavelength)
                    return float(matches[0])
                except ValueError:
                    pass
        return None
    
    # Get first row for excitation wavelengths
    first_row = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    
    # Debug: print first few header values
    print(f"  DEBUG: First row (first 5 values): {first_row[:5]}")
    
    for col_idx, cell in enumerate(first_row):
        if col_idx == 0:  # Skip first column (usually "Emission" or similar label)
            continue
        num = extract_number(cell)
        if num is not None:
            excitation_axis.append(num)
    
    # Get first column for emission wavelengths
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx == 0:  # Skip header row
            continue
        num = extract_number(row[0])
        if num is not None:
            emission_axis.append(num)
    
    emission_arr = np.array(emission_axis)
    excitation_arr = np.array(excitation_axis)
    
    print(f"  DEBUG: Found {len(emission_arr)} emission values, {len(excitation_arr)} excitation values")
    if len(excitation_arr) > 0:
        print(f"  DEBUG: Excitation range: {excitation_arr.min():.1f} - {excitation_arr.max():.1f} nm")
    if len(emission_arr) > 0:
        print(f"  DEBUG: Emission range: {emission_arr.min():.1f} - {emission_arr.max():.1f} nm")
    
    return emission_arr, excitation_arr


# ==============================================================================
# RECONSTRUCTION ERROR COMPUTATION
# ==============================================================================

def compute_reconstruction_error_per_timepoint(
    model: nn.Module, 
    loader: DataLoader, 
    device: torch.device
) -> Dict[int, np.ndarray]:
    """
    Compute pixel-wise reconstruction error (MSE) for each timepoint.
    
    Returns:
        Dictionary mapping timepoint index (0,1,2) to mean MSE map (em × exc)
    """
    model.eval()
    
    T = 3  # Number of timepoints
    mse_accumulators = {t: [] for t in range(T)}
    
    with torch.no_grad():
        for batch_x, batch_y, batch_codes in loader:
            # batch_x: [B, T, C, H, W] = [batch, 3, 1, emission, excitation]
            batch_x = batch_x.to(device)
            B = batch_x.shape[0]
            
            for t in range(T):
                # Get single timepoint: [B, C, H, W]
                x_t = batch_x[:, t, :, :, :]
                
                # Encode
                z_t = model.encoder(x_t)
                
                # Decode
                x_recon = model.decoder(z_t)
                
                # Compute pixel-wise squared error: [B, C, H, W]
                se = (x_t - x_recon) ** 2
                
                # Mean across batch and channel: [H, W]
                mse_map = se.mean(dim=(0, 1)).cpu().numpy()
                mse_accumulators[t].append(mse_map)
    
    # Average across all batches
    mse_per_timepoint = {}
    for t in range(T):
        mse_per_timepoint[t] = np.mean(np.stack(mse_accumulators[t], axis=0), axis=0)
    
    return mse_per_timepoint


def compute_per_sample_reconstruction_error(
    model: nn.Module,
    loader: DataLoader,
    device: torch.device
) -> Tuple[Dict[int, np.ndarray], Dict[str, Dict[int, float]]]:
    """
    Compute reconstruction error per sample per timepoint.
    
    Returns:
        - mse_per_timepoint: Dictionary mapping timepoint to mean MSE map
        - sample_mse: Dictionary mapping sample_code to {timepoint: scalar_mse}
    """
    model.eval()
    
    T = 3
    mse_maps_per_tp = {t: [] for t in range(T)}
    sample_mse = {}
    
    with torch.no_grad():
        for batch_x, batch_y, batch_codes in loader:
            batch_x = batch_x.to(device)
            B = batch_x.shape[0]
            
            for b in range(B):
                code = batch_codes[b]
                sample_mse[code] = {}
                
                for t in range(T):
                    x_t = batch_x[b:b+1, t, :, :, :]
                    z_t = model.encoder(x_t)
                    x_recon = model.decoder(z_t)
                    
                    se = (x_t - x_recon) ** 2
                    mse_map = se.squeeze().cpu().numpy()  # [H, W]
                    mse_maps_per_tp[t].append(mse_map)
                    
                    sample_mse[code][t] = float(mse_map.mean())
    
    # Average maps across samples
    mse_per_timepoint = {}
    for t in range(T):
        mse_per_timepoint[t] = np.mean(np.stack(mse_maps_per_tp[t], axis=0), axis=0)
    
    return mse_per_timepoint, sample_mse


# ==============================================================================
# LATENT DIMENSION RANKING
# ==============================================================================

def rank_latent_dimensions(
    model: nn.Module, 
    loader: DataLoader, 
    device: torch.device,
    top_k: int = 20
) -> Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """Rank latent dimensions by classification importance."""
    model.eval()
    
    all_z = []
    all_y = []
    
    with torch.no_grad():
        for batch_x, batch_y, batch_codes in loader:
            batch_x = batch_x.to(device)
            B, T = batch_x.shape[0], batch_x.shape[1]
            
            z_timepoints = []
            for t in range(T):
                x_t = batch_x[:, t, :, :, :]
                z_t = model.encoder(x_t)
                z_timepoints.append(z_t)
            
            z_stacked = torch.stack(z_timepoints, dim=1)
            z_agg, _ = model.attention(z_stacked)
            
            all_z.append(z_agg.cpu().numpy())
            all_y.append(batch_y.numpy())
    
    z_all = np.concatenate(all_z, axis=0)
    y_all = np.concatenate(all_y, axis=0)
    
    scaler = StandardScaler()
    z_scaled = scaler.fit_transform(z_all)
    
    clf = LogisticRegression(penalty='l2', C=1.0, max_iter=1000, random_state=SEED)
    clf.fit(z_scaled, y_all)
    
    coeffs = clf.coef_[0]
    abs_coeffs = np.abs(coeffs)
    
    top_indices = np.argsort(abs_coeffs)[::-1][:top_k]
    signs = np.sign(coeffs[top_indices])
    
    return top_indices.astype(int), signs, z_all, y_all


# ==============================================================================
# GRADIENT SALIENCY PER TIMEPOINT
# ==============================================================================

def compute_gradient_saliency_per_timepoint(
    model: nn.Module,
    loader: DataLoader,
    top_dims: np.ndarray,
    signs: np.ndarray,
    device: torch.device
) -> Dict[int, Dict[int, np.ndarray]]:
    """Compute gradient saliency for each latent dimension, per timepoint."""
    model.eval()
    
    gradient_maps = {int(d): {t: [] for t in range(3)} for d in top_dims}
    
    for batch_x, batch_y, batch_codes in loader:
        batch_x = batch_x.to(device)
        B, T = batch_x.shape[0], batch_x.shape[1]
        
        for t in range(T):
            x_t = batch_x[:, t, :, :, :].clone()
            x_t.requires_grad_(True)
            
            z_t = model.encoder(x_t)
            
            for d_idx, d in enumerate(top_dims):
                sign = signs[d_idx]
                target = (sign * z_t[:, int(d)]).sum()
                
                target.backward(retain_graph=True)
                grad = x_t.grad.data.abs()
                
                grad_map = grad.mean(dim=(0, 1)).cpu().numpy()
                gradient_maps[int(d)][t].append(grad_map)
                
                x_t.grad.data.zero_()
    
    for d in top_dims:
        for t in range(3):
            gradient_maps[int(d)][t] = np.mean(
                np.stack(gradient_maps[int(d)][t], axis=0), axis=0
            )
    
    return gradient_maps


# ==============================================================================
# LINEAR PROBING PER TIMEPOINT
# ==============================================================================

def compute_linear_probe_per_timepoint(
    model: nn.Module,
    loader: DataLoader,
    top_dims: np.ndarray,
    z_agg_all: np.ndarray,
    device: torch.device
) -> Dict[int, Dict[int, np.ndarray]]:
    """Compute linear probe importance for each latent dimension, per timepoint."""
    model.eval()
    
    T = 3
    x_flat_per_tp = {t: [] for t in range(T)}
    
    with torch.no_grad():
        for batch_x, batch_y, batch_codes in loader:
            batch_x = batch_x.to(device)
            B = batch_x.shape[0]
            
            for t in range(T):
                x_t = batch_x[:, t, 0, :, :].cpu().numpy()
                x_flat_per_tp[t].append(x_t.reshape(B, -1))
    
    for t in range(T):
        x_flat_per_tp[t] = np.concatenate(x_flat_per_tp[t], axis=0)
    
    input_shape = loader.dataset[0][0].shape[2:]
    
    linear_maps = {int(d): {t: None for t in range(T)} for d in top_dims}
    
    for d in top_dims:
        z_target = z_agg_all[:, int(d)]
        
        for t in range(T):
            X = x_flat_per_tp[t]
            y = z_target
            
            scaler = StandardScaler()
            X_scaled = scaler.fit_transform(X)
            
            ridge = Ridge(alpha=1.0)
            ridge.fit(X_scaled, y)
            
            weights = np.abs(ridge.coef_).reshape(input_shape)
            linear_maps[int(d)][t] = weights
    
    return linear_maps


# ==============================================================================
# CONSENSUS AGGREGATION
# ==============================================================================

def aggregate_importance_per_timepoint(
    gradient_maps: Dict[int, Dict[int, np.ndarray]],
    linear_maps: Dict[int, Dict[int, np.ndarray]],
    top_dims: np.ndarray,
    T: int = 3
) -> Dict[int, np.ndarray]:
    """Create consensus importance maps per timepoint."""
    consensus = {}
    
    for t in range(T):
        g_stack = [gradient_maps[int(d)][t] for d in top_dims]
        l_stack = [linear_maps[int(d)][t] for d in top_dims]
        
        g_avg = np.mean(np.stack(g_stack, axis=0), axis=0)
        l_avg = np.mean(np.stack(l_stack, axis=0), axis=0)
        
        g_norm = (g_avg - g_avg.min()) / (g_avg.max() - g_avg.min() + 1e-10)
        l_norm = (l_avg - l_avg.min()) / (l_avg.max() - l_avg.min() + 1e-10)
        
        consensus[t] = (g_norm + l_norm) / 2.0
    
    return consensus


# ==============================================================================
# ROI EXTRACTION
# ==============================================================================

def extract_rois_at_percentile(
    importance_map: np.ndarray,
    percentile: float = 90,
    emission_axis: np.ndarray = None,
    excitation_axis: np.ndarray = None,
    min_area: int = 4
) -> List[Dict]:
    """Extract connected regions above percentile threshold."""
    threshold = np.percentile(importance_map, percentile)
    binary = importance_map >= threshold
    
    labeled, num_features = ndimage.label(binary)
    
    rois = []
    for comp_id in range(1, num_features + 1):
        comp_mask = labeled == comp_id
        ys, xs = np.where(comp_mask)
        
        if len(xs) < min_area:
            continue
        
        y0, y1 = int(ys.min()), int(ys.max())
        x0, x1 = int(xs.min()), int(xs.max())
        
        roi = {
            'component_id': comp_id,
            'y0': y0, 'y1': y1, 'x0': x0, 'x1': x1,
            'cy': float(ys.mean()), 'cx': float(xs.mean()),
            'area_px': len(xs),
            'mean_importance': float(importance_map[comp_mask].mean()),
            'max_importance': float(importance_map[comp_mask].max()),
            'sum_importance': float(importance_map[comp_mask].sum()),
        }
        
        if emission_axis is not None and len(emission_axis) > y1:
            roi['em_min_nm'] = float(emission_axis[y0])
            roi['em_max_nm'] = float(emission_axis[min(y1, len(emission_axis)-1)])
            roi['em_centroid_nm'] = float(emission_axis[min(int(roi['cy']), len(emission_axis)-1)])
        
        if excitation_axis is not None and len(excitation_axis) > x1:
            roi['exc_min_nm'] = float(excitation_axis[x0])
            roi['exc_max_nm'] = float(excitation_axis[min(x1, len(excitation_axis)-1)])
            roi['exc_centroid_nm'] = float(excitation_axis[min(int(roi['cx']), len(excitation_axis)-1)])
        
        rois.append(roi)
    
    return sorted(rois, key=lambda r: r['mean_importance'], reverse=True)


def extract_rois_from_masked_importance(
    importance_map: np.ndarray,
    mse_mask: np.ndarray,
    percentile: float = 90,
    emission_axis: np.ndarray = None,
    excitation_axis: np.ndarray = None,
    min_area: int = 4
) -> List[Dict]:
    """
    Extract ROIs from importance map after applying MSE mask.
    
    Only regions where mse_mask is True (low reconstruction error) are considered.
    """
    # Apply mask: set importance to 0 where mask is False
    masked_importance = importance_map.copy()
    masked_importance[~mse_mask] = 0
    
    # Get threshold from masked values only
    valid_values = masked_importance[mse_mask]
    if len(valid_values) == 0:
        return []
    
    threshold = np.percentile(valid_values, percentile)
    binary = masked_importance >= threshold
    
    labeled, num_features = ndimage.label(binary)
    
    rois = []
    for comp_id in range(1, num_features + 1):
        comp_mask = labeled == comp_id
        ys, xs = np.where(comp_mask)
        
        if len(xs) < min_area:
            continue
        
        y0, y1 = int(ys.min()), int(ys.max())
        x0, x1 = int(xs.min()), int(xs.max())
        
        roi = {
            'component_id': comp_id,
            'y0': y0, 'y1': y1, 'x0': x0, 'x1': x1,
            'cy': float(ys.mean()), 'cx': float(xs.mean()),
            'area_px': len(xs),
            'mean_importance': float(importance_map[comp_mask].mean()),
            'max_importance': float(importance_map[comp_mask].max()),
            'sum_importance': float(importance_map[comp_mask].sum()),
            'mean_mse_in_roi': float(np.nan),  # Will be filled later
        }
        
        if emission_axis is not None and len(emission_axis) > y1:
            roi['em_min_nm'] = float(emission_axis[y0])
            roi['em_max_nm'] = float(emission_axis[min(y1, len(emission_axis)-1)])
            roi['em_centroid_nm'] = float(emission_axis[min(int(roi['cy']), len(emission_axis)-1)])
        
        if excitation_axis is not None and len(excitation_axis) > x1:
            roi['exc_min_nm'] = float(excitation_axis[x0])
            roi['exc_max_nm'] = float(excitation_axis[min(x1, len(excitation_axis)-1)])
            roi['exc_centroid_nm'] = float(excitation_axis[min(int(roi['cx']), len(excitation_axis)-1)])
        
        rois.append(roi)
    
    return sorted(rois, key=lambda r: r['mean_importance'], reverse=True)


# ==============================================================================
# VISUALIZATION
# ==============================================================================

def plot_importance_maps(maps_dict, output_path, emission_axis=None, excitation_axis=None,
                        title="Importance Maps"):
    """Plot importance maps for all timepoints."""
    tp_labels = ['0h', '6h', '24h']
    
    fig, axes = plt.subplots(1, 3, figsize=(15, 5))
    
    for t in range(3):
        ax = axes[t]
        
        if emission_axis is not None and excitation_axis is not None:
            extent = [excitation_axis[0], excitation_axis[-1],
                     emission_axis[0], emission_axis[-1]]
            im = ax.imshow(maps_dict[t], cmap='hot', aspect='auto',
                          origin='lower', extent=extent)
            ax.set_xlabel('Excitation (nm)')
            ax.set_ylabel('Emission (nm)')
        else:
            im = ax.imshow(maps_dict[t], cmap='hot', aspect='auto', origin='lower')
        
        ax.set_title(f'{tp_labels[t]}')
        plt.colorbar(im, ax=ax, fraction=0.046)
    
    plt.suptitle(title, fontsize=14, fontweight='bold')
    plt.tight_layout()
    plt.savefig(output_path, dpi=150, bbox_inches='tight')
    plt.close()


def plot_mse_maps(mse_dict, output_path, emission_axis=None, excitation_axis=None,
                 threshold=0.05, title="Reconstruction Error (MSE)"):
    """Plot MSE maps with threshold line for all timepoints."""
    tp_labels = ['0h', '6h', '24h']
    
    fig, axes = plt.subplots(2, 3, figsize=(15, 10))
    
    for t in range(3):
        # Top row: MSE maps
        ax = axes[0, t]
        
        if emission_axis is not None and excitation_axis is not None:
            extent = [excitation_axis[0], excitation_axis[-1],
                     emission_axis[0], emission_axis[-1]]
            im = ax.imshow(mse_dict[t], cmap='viridis', aspect='auto',
                          origin='lower', extent=extent, vmin=0, vmax=0.1)
            ax.set_xlabel('Excitation (nm)')
            ax.set_ylabel('Emission (nm)')
        else:
            im = ax.imshow(mse_dict[t], cmap='viridis', aspect='auto', origin='lower',
                          vmin=0, vmax=0.1)
        
        ax.set_title(f'{tp_labels[t]} - MSE')
        plt.colorbar(im, ax=ax, fraction=0.046)
        
        # Bottom row: Binary mask (MSE < threshold)
        ax2 = axes[1, t]
        mask = mse_dict[t] < threshold
        
        if emission_axis is not None and excitation_axis is not None:
            im2 = ax2.imshow(mask.astype(float), cmap='Greens', aspect='auto',
                            origin='lower', extent=extent, vmin=0, vmax=1)
            ax2.set_xlabel('Excitation (nm)')
            ax2.set_ylabel('Emission (nm)')
        else:
            im2 = ax2.imshow(mask.astype(float), cmap='Greens', aspect='auto', 
                            origin='lower', vmin=0, vmax=1)
        
        pct_valid = 100 * mask.sum() / mask.size
        ax2.set_title(f'{tp_labels[t]} - Valid Region ({pct_valid:.1f}%)')
    
    plt.suptitle(f'{title}\nThreshold: MSE < {threshold}', fontsize=14, fontweight='bold')
    plt.tight_layout()
    plt.savefig(output_path, dpi=150, bbox_inches='tight')
    plt.close()


def plot_masked_importance(importance_map, mse_map, mse_mask, rois, output_path,
                          emission_axis=None, excitation_axis=None, 
                          title="Masked Importance with ROIs"):
    """Plot importance map with MSE mask overlay and ROI boxes."""
    fig, axes = plt.subplots(1, 3, figsize=(18, 6))
    
    extent = None
    if emission_axis is not None and excitation_axis is not None:
        extent = [excitation_axis[0], excitation_axis[-1],
                 emission_axis[0], emission_axis[-1]]
    
    # Plot 1: Original importance
    ax1 = axes[0]
    if extent:
        im1 = ax1.imshow(importance_map, cmap='hot', aspect='auto', origin='lower', extent=extent)
    else:
        im1 = ax1.imshow(importance_map, cmap='hot', aspect='auto', origin='lower')
    ax1.set_title('Original Importance')
    plt.colorbar(im1, ax=ax1, fraction=0.046)
    
    # Plot 2: MSE map
    ax2 = axes[1]
    if extent:
        im2 = ax2.imshow(mse_map, cmap='viridis', aspect='auto', origin='lower', extent=extent,
                        vmin=0, vmax=0.1)
    else:
        im2 = ax2.imshow(mse_map, cmap='viridis', aspect='auto', origin='lower', vmin=0, vmax=0.1)
    ax2.set_title('Reconstruction Error (MSE)')
    plt.colorbar(im2, ax=ax2, fraction=0.046)
    
    # Plot 3: Masked importance with ROIs
    ax3 = axes[2]
    masked_imp = importance_map.copy()
    masked_imp[~mse_mask] = 0
    
    if extent:
        im3 = ax3.imshow(masked_imp, cmap='hot', aspect='auto', origin='lower', extent=extent)
        ax3.set_xlabel('Excitation (nm)')
        ax3.set_ylabel('Emission (nm)')
        
        # Draw ROI boxes
        for i, roi in enumerate(rois[:10]):
            if 'exc_min_nm' in roi:
                rect_x = [roi['exc_min_nm'], roi['exc_max_nm'], roi['exc_max_nm'],
                         roi['exc_min_nm'], roi['exc_min_nm']]
                rect_y = [roi['em_min_nm'], roi['em_min_nm'], roi['em_max_nm'],
                         roi['em_max_nm'], roi['em_min_nm']]
                ax3.plot(rect_x, rect_y, 'c-', linewidth=2)
                ax3.text(roi['exc_centroid_nm'], roi['em_centroid_nm'], 
                        str(i+1), color='cyan', fontweight='bold', 
                        ha='center', va='center', fontsize=10)
    else:
        im3 = ax3.imshow(masked_imp, cmap='hot', aspect='auto', origin='lower')
    
    ax3.set_title(f'Masked Importance ({len(rois)} ROIs)')
    plt.colorbar(im3, ax=ax3, fraction=0.046)
    
    plt.suptitle(title, fontsize=14, fontweight='bold')
    plt.tight_layout()
    plt.savefig(output_path, dpi=150, bbox_inches='tight')
    plt.close()


# ==============================================================================
# MAIN PIPELINE
# ==============================================================================

def main():
    parser = argparse.ArgumentParser(
        description='Timepoint importance with reconstruction error filtering',
        formatter_class=argparse.RawDescriptionHelpFormatter
    )
    
    parser.add_argument('--model_dir', type=str, required=True,
                       help='Directory containing best_model_fold*.pth files')
    parser.add_argument('--tp_dirs', type=str, required=True,
                       help='Comma-separated paths to 0h,6h,24h directories')
    parser.add_argument('--labels_csv', type=str, required=True,
                       help='CSV with columns: code, group (ALS/CTRL)')
    parser.add_argument('--output_dir', type=str, default='./importance_results')
    parser.add_argument('--n_folds', type=int, default=5)
    parser.add_argument('--top_dims', type=int, default=20)
    parser.add_argument('--latent_dim', type=int, default=512)
    parser.add_argument('--batch_size', type=int, default=4)
    parser.add_argument('--roi_percentile', type=float, default=90)
    parser.add_argument('--mse_threshold', type=float, default=0.05,
                       help='MSE threshold for filtering (default: 0.05). '
                            'Only regions with MSE < threshold are considered reliable.')
    
    args = parser.parse_args()
    
    # Setup
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    
    tp_dirs = args.tp_dirs.split(',')
    assert len(tp_dirs) == 3, "Must provide exactly 3 timepoint directories (0h, 6h, 24h)"
    
    device = torch.device('cuda' if torch.cuda.is_available() else 'cpu')
    print(f"Device: {device}")
    print(f"Model directory: {args.model_dir}")
    print(f"Output directory: {output_dir}")
    print(f"MSE threshold: {args.mse_threshold}")
    
    # Verify models exist
    model_files = []
    for fold in range(1, args.n_folds + 1):
        model_path = os.path.join(args.model_dir, f'best_model_fold{fold}.pth')
        if not os.path.exists(model_path):
            print(f"ERROR: Model not found: {model_path}")
            sys.exit(1)
        model_files.append(model_path)
    print(f"Found {len(model_files)} pre-trained models")
    
    # Load labels
    labels_df = pd.read_csv(args.labels_csv)
    if 'sample_id' in labels_df.columns and 'code' not in labels_df.columns:
        labels_df['code'] = labels_df['sample_id']
    
    codes = labels_df['code'].astype(str).tolist()
    groups = labels_df['group'].tolist()
    labels = np.array([1 if g == 'ALS' else 0 for g in groups])
    
    print(f"\nLoaded {len(codes)} samples")
    print(f"  ALS: {sum(labels)}, CTRL: {len(labels) - sum(labels)}")
    
    # Create dataset
    dataset = ExcelImageDataset(codes, labels.tolist(), tp_dirs)
    loader = DataLoader(dataset, batch_size=args.batch_size, shuffle=False, num_workers=0)
    
    # Get spectral axes from first file
    first_file = dataset.find_file(tp_dirs[0], codes[0])
    emission_axis, excitation_axis = load_spectral_axes(first_file)
    print(f"Spectral grid: {len(emission_axis)} emission × {len(excitation_axis)} excitation")
    
    # Storage
    tp_labels = ['0h', '6h', '24h']
    all_fold_consensus = {t: [] for t in range(3)}
    all_fold_gradient = {t: [] for t in range(3)}
    all_fold_linear = {t: [] for t in range(3)}
    all_fold_mse = {t: [] for t in range(3)}  # NEW: MSE per fold
    all_fold_top_dims = []
    
    print(f"\n{'='*70}")
    print(f"PROCESSING {args.n_folds} PRE-TRAINED MODELS")
    print(f"{'='*70}")
    
    for fold_idx in range(args.n_folds):
        print(f"\n{'='*70}")
        print(f"MODEL {fold_idx + 1}/{args.n_folds}")
        print(f"{'='*70}")
        
        # Initialize model
        model = ConvAutoencoderWithAttention(
            in_channels=1,
            latent_dim=args.latent_dim,
            num_classes=2
        ).to(device)
        
        # Do a forward pass to initialize dynamic layers
        sample_batch, _, _ = next(iter(loader))
        with torch.no_grad():
            _ = model.encoder(sample_batch[0, 0:1].to(device))
        
        # Initialize decoder
        if model.encoder.encoded_shape is not None:
            target_shape = (sample_batch.shape[3], sample_batch.shape[4])
            model.init_decoder(model.encoder.encoded_shape, target_shape)
            model = model.to(device)
        
        # Load pre-trained weights
        model_path = model_files[fold_idx]
        print(f"  Loading: {model_path}")
        state_dict = torch.load(model_path, map_location=device)
        model.load_state_dict(state_dict)
        model.eval()
        print(f"  ✓ Model loaded successfully")
        
        # NEW: Compute reconstruction error per timepoint
        print(f"  Computing reconstruction error per timepoint...")
        mse_per_tp, sample_mse = compute_per_sample_reconstruction_error(model, loader, device)
        
        for t in range(3):
            mean_mse = mse_per_tp[t].mean()
            pct_below_thr = 100 * (mse_per_tp[t] < args.mse_threshold).sum() / mse_per_tp[t].size
            print(f"    {tp_labels[t]}: mean MSE={mean_mse:.4f}, {pct_below_thr:.1f}% below {args.mse_threshold}")
            all_fold_mse[t].append(mse_per_tp[t])
        
        # Step 1: Rank latent dimensions
        print(f"  Ranking latent dimensions...")
        top_dims, signs, z_all, y_all = rank_latent_dimensions(
            model, loader, device, top_k=args.top_dims
        )
        all_fold_top_dims.append(list(top_dims))
        print(f"    Top {args.top_dims} dims: {list(top_dims[:5])}... (showing first 5)")
        
        # Step 2: Gradient Saliency per timepoint
        print(f"  Computing Gradient Saliency...")
        gradient_maps = compute_gradient_saliency_per_timepoint(
            model, loader, top_dims, signs, device
        )
        
        # Step 3: Linear Probing per timepoint
        print(f"  Computing Linear Probing...")
        linear_maps = compute_linear_probe_per_timepoint(
            model, loader, top_dims, z_all, device
        )
        
        # Step 4: Consensus per timepoint
        print(f"  Creating consensus maps...")
        consensus_maps = aggregate_importance_per_timepoint(
            gradient_maps, linear_maps, top_dims, T=3
        )
        
        # Store
        for t in range(3):
            all_fold_consensus[t].append(consensus_maps[t])
            g_avg = np.mean([gradient_maps[int(d)][t] for d in top_dims], axis=0)
            l_avg = np.mean([linear_maps[int(d)][t] for d in top_dims], axis=0)
            all_fold_gradient[t].append(g_avg)
            all_fold_linear[t].append(l_avg)
        
        # Save fold-specific results
        fold_dir = output_dir / f'fold{fold_idx + 1}'
        fold_dir.mkdir(exist_ok=True)
        
        for t in range(3):
            np.save(fold_dir / f'consensus_{tp_labels[t]}.npy', consensus_maps[t])
            np.save(fold_dir / f'gradient_{tp_labels[t]}.npy', 
                   np.mean([gradient_maps[int(d)][t] for d in top_dims], axis=0))
            np.save(fold_dir / f'linear_{tp_labels[t]}.npy',
                   np.mean([linear_maps[int(d)][t] for d in top_dims], axis=0))
            np.save(fold_dir / f'mse_{tp_labels[t]}.npy', mse_per_tp[t])
        
        np.save(fold_dir / 'top_dims.npy', top_dims)
        
        plot_importance_maps(
            consensus_maps, fold_dir / 'consensus_maps.png',
            emission_axis, excitation_axis,
            title=f'Fold {fold_idx+1} Consensus (Gradient + Linear)'
        )
        
        plot_mse_maps(
            mse_per_tp, fold_dir / 'mse_maps.png',
            emission_axis, excitation_axis,
            threshold=args.mse_threshold,
            title=f'Fold {fold_idx+1} Reconstruction Error'
        )
        
        print(f"  ✓ Saved fold results to {fold_dir}")
    
    # ==============================================================================
    # CROSS-MODEL CONSENSUS
    # ==============================================================================
    
    print(f"\n{'='*70}")
    print("COMPUTING CROSS-MODEL CONSENSUS")
    print(f"{'='*70}")
    
    final_consensus = {}
    final_std = {}
    final_stability = {}
    final_mse = {}  # NEW: Cross-fold MSE
    final_mse_mask = {}  # NEW: Valid region mask
    all_rois = {}
    all_rois_unfiltered = {}  # For comparison
    
    for t in range(3):
        tp = tp_labels[t]
        
        # Importance consensus
        fold_maps = np.stack(all_fold_consensus[t], axis=0)
        mean_map = np.mean(fold_maps, axis=0)
        std_map = np.std(fold_maps, axis=0)
        stability_map = mean_map / (std_map + 1e-10)
        
        final_consensus[t] = mean_map
        final_std[t] = std_map
        final_stability[t] = stability_map
        
        # MSE consensus (NEW)
        fold_mse_maps = np.stack(all_fold_mse[t], axis=0)
        mean_mse_map = np.mean(fold_mse_maps, axis=0)
        final_mse[t] = mean_mse_map
        
        # Create mask: valid where MSE < threshold (NEW)
        mse_mask = mean_mse_map < args.mse_threshold
        final_mse_mask[t] = mse_mask
        
        pct_valid = 100 * mse_mask.sum() / mse_mask.size
        print(f"\n  {tp}: MSE < {args.mse_threshold} in {pct_valid:.1f}% of pixels")
        
        # Save
        np.save(output_dir / f'final_consensus_{tp}.npy', mean_map)
        np.save(output_dir / f'fold_std_{tp}.npy', std_map)
        np.save(output_dir / f'stability_{tp}.npy', stability_map)
        np.save(output_dir / f'final_mse_{tp}.npy', mean_mse_map)
        np.save(output_dir / f'mse_mask_{tp}.npy', mse_mask)
        
        # Extract ROIs from MASKED importance (NEW)
        rois = extract_rois_from_masked_importance(
            mean_map, mse_mask, args.roi_percentile, 
            emission_axis, excitation_axis
        )
        
        # Also extract unfiltered ROIs for comparison
        rois_unfiltered = extract_rois_at_percentile(
            mean_map, args.roi_percentile, emission_axis, excitation_axis
        )
        
        # Add stability and MSE info
        for roi in rois:
            y0, y1 = roi['y0'], roi['y1']
            x0, x1 = roi['x0'], roi['x1']
            roi['stability'] = float(stability_map[y0:y1+1, x0:x1+1].mean())
            roi['fold_std'] = float(std_map[y0:y1+1, x0:x1+1].mean())
            roi['mean_mse_in_roi'] = float(mean_mse_map[y0:y1+1, x0:x1+1].mean())
            roi['timepoint'] = tp
        
        for roi in rois_unfiltered:
            y0, y1 = roi['y0'], roi['y1']
            x0, x1 = roi['x0'], roi['x1']
            roi['stability'] = float(stability_map[y0:y1+1, x0:x1+1].mean())
            roi['fold_std'] = float(std_map[y0:y1+1, x0:x1+1].mean())
            roi['mean_mse_in_roi'] = float(mean_mse_map[y0:y1+1, x0:x1+1].mean())
            roi['timepoint'] = tp
        
        all_rois[tp] = rois
        all_rois_unfiltered[tp] = rois_unfiltered
        
        print(f"  {tp}: {len(rois)} ROIs (filtered) vs {len(rois_unfiltered)} (unfiltered)")
        for i, roi in enumerate(rois[:5]):
            em_str = f"{roi.get('em_centroid_nm', roi['cy']):.0f}nm"
            exc_str = f"{roi.get('exc_centroid_nm', roi['cx']):.0f}nm"
            print(f"    ROI {i+1}: em~{em_str}, exc~{exc_str}, "
                  f"imp={roi['mean_importance']:.4f}, mse={roi['mean_mse_in_roi']:.4f}, "
                  f"stability={roi['stability']:.2f}")
        
        # Visualizations
        plot_masked_importance(
            mean_map, mean_mse_map, mse_mask, rois,
            output_dir / f'masked_importance_{tp}.png',
            emission_axis, excitation_axis,
            title=f'{tp} - Masked by Reconstruction Quality (MSE < {args.mse_threshold})'
        )
    
    # Combined visualizations
    plot_importance_maps(
        final_consensus,
        output_dir / 'final_consensus_all_timepoints.png',
        emission_axis, excitation_axis,
        title=f'Final Consensus Importance Maps (avg of {args.n_folds} models)'
    )
    
    plot_mse_maps(
        final_mse,
        output_dir / 'final_mse_all_timepoints.png',
        emission_axis, excitation_axis,
        threshold=args.mse_threshold,
        title=f'Final Reconstruction Error (avg of {args.n_folds} models)'
    )
    
    # Save ROIs
    rois_flat = []
    for tp, rois in all_rois.items():
        rois_flat.extend(rois)
    pd.DataFrame(rois_flat).to_csv(output_dir / 'rois_filtered_by_mse.csv', index=False)
    
    rois_flat_unfiltered = []
    for tp, rois in all_rois_unfiltered.items():
        rois_flat_unfiltered.extend(rois)
    pd.DataFrame(rois_flat_unfiltered).to_csv(output_dir / 'rois_unfiltered.csv', index=False)
    
    # Summary
    print(f"\n{'='*70}")
    print("SUMMARY")
    print(f"{'='*70}")
    
    summary = [
        "TIMEPOINT-SPECIFIC FEATURE IMPORTANCE WITH RECONSTRUCTION FILTERING",
        "=" * 70,
        f"\nModels loaded from: {args.model_dir}",
        f"Number of folds: {args.n_folds}",
        f"Top dimensions per model: {args.top_dims}",
        f"Latent dimension: {args.latent_dim}",
        f"ROI percentile: {args.roi_percentile}",
        f"MSE threshold: {args.mse_threshold}",
        f"\nSamples: {len(codes)} ({sum(labels)} ALS, {len(labels)-sum(labels)} CTRL)",
        f"Spectral grid: {len(emission_axis)} em × {len(excitation_axis)} exc",
        f"\nReconstruction Quality by Timepoint:",
    ]
    
    for tp in tp_labels:
        t = tp_labels.index(tp)
        mean_mse = final_mse[t].mean()
        pct_valid = 100 * final_mse_mask[t].sum() / final_mse_mask[t].size
        summary.append(f"  - {tp}: mean MSE = {mean_mse:.4f}, {pct_valid:.1f}% valid (< {args.mse_threshold})")
    
    summary.append(f"\nROIs per timepoint (filtered / unfiltered):")
    for tp in tp_labels:
        summary.append(f"  - {tp}: {len(all_rois[tp])} / {len(all_rois_unfiltered[tp])}")
    
    summary.append(f"\nTop Filtered ROIs by importance:")
    for tp in tp_labels:
        summary.append(f"\n  {tp}:")
        for i, roi in enumerate(all_rois[tp][:3]):
            em = roi.get('em_centroid_nm', roi['cy'])
            exc = roi.get('exc_centroid_nm', roi['cx'])
            summary.append(f"    {i+1}. em~{em:.0f}nm, exc~{exc:.0f}nm, "
                          f"imp={roi['mean_importance']:.4f}, "
                          f"mse={roi['mean_mse_in_roi']:.4f}, "
                          f"stability={roi['stability']:.2f}")
    
    summary_text = '\n'.join(summary)
    print(summary_text)
    
    with open(output_dir / 'SUMMARY.txt', 'w') as f:
        f.write(summary_text)
    
    # Config
    config = {
        'model_dir': args.model_dir,
        'n_folds': args.n_folds,
        'top_dims': args.top_dims,
        'latent_dim': args.latent_dim,
        'roi_percentile': args.roi_percentile,
        'mse_threshold': args.mse_threshold,
        'n_samples': len(codes),
        'fold_top_dims': [[int(d) for d in dims] for dims in all_fold_top_dims],
        'mse_statistics': {
            tp: {
                'mean': float(final_mse[t].mean()),
                'std': float(final_mse[t].std()),
                'pct_below_threshold': float(100 * final_mse_mask[t].sum() / final_mse_mask[t].size)
            }
            for t, tp in enumerate(tp_labels)
        }
    }
    with open(output_dir / 'config.json', 'w') as f:
        json.dump(config, f, indent=2)
    
    print(f"\n{'='*70}")
    print(f"Results saved to: {output_dir}")
    print(f"{'='*70}")
    print(f"\nKey outputs:")
    print(f"  - rois_filtered_by_mse.csv  : ROIs from regions with low reconstruction error")
    print(f"  - rois_unfiltered.csv       : ROIs without MSE filtering (for comparison)")
    print(f"  - final_mse_*.npy           : Consensus reconstruction error maps")
    print(f"  - mse_mask_*.npy            : Boolean masks (True where MSE < {args.mse_threshold})")
    print(f"  - masked_importance_*.png   : Visualizations with MSE filtering")


if __name__ == '__main__':
    main()
