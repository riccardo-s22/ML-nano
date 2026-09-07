#!/usr/bin/env python3
"""
TIMEPOINT-SPECIFIC FEATURE IMPORTANCE PIPELINE (Using Pre-trained Models)
=========================================================================

This script loads pre-trained 5-fold CV models and computes feature importance
for each timepoint using gradient saliency and linear probing.

WORKFLOW:
---------
1. Load 5 pre-trained models from disk
2. For EACH model:
   a) Extract latent features → rank by classification importance → TOP 20
   b) For each top-20 dimension × each timepoint (0h, 6h, 24h):
      - Gradient Saliency: |∂z_d/∂x_t|
      - Linear Probing: Ridge weights
   c) Consensus = average(gradient, linear) per timepoint
3. Cross-model consensus = average across 5 models per timepoint
4. Extract ROIs at 95th percentile

USAGE:
------
python timepoint_importance_pretrained.py \
    --model_dir C:/Users/.../5_fold_models_original \
    --tp_dirs C:/Users/.../out_0h,C:/Users/.../out_6h,C:/Users/.../out_24h \
    --labels_csv C:/Users/.../sample_labels.csv \
    --output_dir ./importance_results
"""

import os
import sys
import argparse
import json
from pathlib import Path
from typing import List, Tuple, Dict, Optional
from collections import defaultdict

import numpy as np
import pandas as pd
import torch
import torch.nn as nn
from torch.utils.data import Dataset, DataLoader
from sklearn.linear_model import Ridge, LogisticRegression
from sklearn.preprocessing import StandardScaler
from sklearn.model_selection import StratifiedKFold
from scipy.stats import rankdata
from scipy import ndimage
from scipy.ndimage import gaussian_filter
import matplotlib.pyplot as plt
import seaborn as sns
from openpyxl import load_workbook
import warnings
warnings.filterwarnings('ignore')

# Reproducibility
SEED = 42
np.random.seed(SEED)
torch.manual_seed(SEED)


# ==============================================================================
# MODEL ARCHITECTURE (must match the saved models exactly)
# ==============================================================================

class ConvEncoder(nn.Module):
    """Convolutional encoder - MUST match original architecture."""
    
    def __init__(self, in_channels: int, latent_dim: int):
        super(ConvEncoder, self).__init__()
        
        self.conv1 = nn.Sequential(
            nn.Conv2d(in_channels, 32, kernel_size=3, padding=1),
            nn.BatchNorm2d(32),
            nn.ReLU(),
            nn.MaxPool2d(2, 2)
        )
        self.conv2 = nn.Sequential(
            nn.Conv2d(32, 64, kernel_size=3, padding=1),
            nn.BatchNorm2d(64),
            nn.ReLU(),
            nn.MaxPool2d(2, 2)
        )
        self.conv3 = nn.Sequential(
            nn.Conv2d(64, 128, kernel_size=3, padding=1),
            nn.BatchNorm2d(128),
            nn.ReLU(),
            nn.MaxPool2d(2, 2)
        )
        self.conv4 = nn.Sequential(
            nn.Conv2d(128, 256, kernel_size=3, padding=1),
            nn.BatchNorm2d(256),
            nn.ReLU()
        )
        
        self.flatten_size = None
        self.fc = None
        self.latent_dim = latent_dim
        self.encoded_shape = None
        
    def forward(self, x):
        x = self.conv1(x)
        x = self.conv2(x)
        x = self.conv3(x)
        x = self.conv4(x)
        
        if self.encoded_shape is None:
            self.encoded_shape = (x.shape[1], x.shape[2], x.shape[3])
        
        if self.fc is None:
            self.flatten_size = x.shape[1] * x.shape[2] * x.shape[3]
            self.fc = nn.Linear(self.flatten_size, self.latent_dim).to(x.device)
        
        x = x.view(x.size(0), -1)
        x = self.fc(x)
        return x


class ConvDecoder(nn.Module):
    """Convolutional decoder - MUST match original architecture."""
    
    def __init__(self, latent_dim: int, out_channels: int, 
                 encoded_shape: Tuple[int, int, int], target_shape: Tuple[int, int]):
        super(ConvDecoder, self).__init__()
        
        self.encoded_channels, self.encoded_h, self.encoded_w = encoded_shape
        self.target_h, self.target_w = target_shape
        
        self.fc = nn.Linear(latent_dim, self.encoded_channels * self.encoded_h * self.encoded_w)
        
        self.deconv1 = nn.Sequential(
            nn.ConvTranspose2d(256, 128, kernel_size=3, stride=2, padding=1, output_padding=1),
            nn.BatchNorm2d(128),
            nn.ReLU()
        )
        self.deconv2 = nn.Sequential(
            nn.ConvTranspose2d(128, 64, kernel_size=3, stride=2, padding=1, output_padding=1),
            nn.BatchNorm2d(64),
            nn.ReLU()
        )
        self.deconv3 = nn.Sequential(
            nn.ConvTranspose2d(64, 32, kernel_size=3, stride=2, padding=1, output_padding=1),
            nn.BatchNorm2d(32),
            nn.ReLU()
        )
        self.final_conv = nn.Conv2d(32, out_channels, kernel_size=3, padding=1)
        self.sigmoid = nn.Sigmoid()
        
    def forward(self, z):
        x = self.fc(z)
        x = x.view(x.size(0), self.encoded_channels, self.encoded_h, self.encoded_w)
        x = self.deconv1(x)
        x = self.deconv2(x)
        x = self.deconv3(x)
        x = self.final_conv(x)
        
        current_h, current_w = x.shape[2], x.shape[3]
        if current_h != self.target_h or current_w != self.target_w:
            x = torch.nn.functional.interpolate(
                x, size=(self.target_h, self.target_w), 
                mode='bilinear', align_corners=False
            )
        
        x = self.sigmoid(x)
        return x


class AttentionAggregation(nn.Module):
    """Attention mechanism - MUST match original architecture."""
    
    def __init__(self, latent_dim: int):
        super(AttentionAggregation, self).__init__()
        self.attention_net = nn.Sequential(
            nn.Linear(latent_dim, 64),
            nn.Tanh(),
            nn.Linear(64, 1)
        )
        
    def forward(self, latent_vectors):
        attention_scores = self.attention_net(latent_vectors)
        attention_scores = attention_scores.squeeze(-1)
        attention_weights = torch.softmax(attention_scores, dim=1)
        attention_weights_expanded = attention_weights.unsqueeze(-1)
        aggregated = torch.sum(latent_vectors * attention_weights_expanded, dim=1)
        return aggregated, attention_weights


class Classifier(nn.Module):
    """Classification head - MUST match original architecture."""
    
    def __init__(self, latent_dim: int, num_classes: int, dropout_rate: float = 0.3):
        super(Classifier, self).__init__()
        self.classifier = nn.Sequential(
            nn.Linear(latent_dim, 128),
            nn.ReLU(),
            nn.Dropout(dropout_rate),
            nn.Linear(128, num_classes)
        )
        
    def forward(self, z_agg):
        return self.classifier(z_agg)


class ConvAutoencoderWithAttention(nn.Module):
    """Complete model - MUST match original architecture."""
    
    def __init__(self, in_channels: int, latent_dim: int, num_classes: int):
        super(ConvAutoencoderWithAttention, self).__init__()
        
        self.encoder = ConvEncoder(in_channels, latent_dim)
        self.attention = AttentionAggregation(latent_dim)
        self.classifier = Classifier(latent_dim, num_classes)
        self.decoder = None
        self.latent_dim = latent_dim
        self.in_channels = in_channels
        
    def init_decoder(self, encoded_shape: tuple, target_shape: tuple):
        self.decoder = ConvDecoder(self.latent_dim, self.in_channels, 
                                   encoded_shape, target_shape)


# ==============================================================================
# DATA LOADING
# ==============================================================================

class ExcelImageDataset(Dataset):
    """Dataset for loading Excel spectral data."""
    
    def __init__(self, codes: List[str], labels: List[int], 
                 timepoint_dirs: List[str], file_suffix: str = ".xlsx"):
        self.codes = codes
        self.labels = labels
        self.timepoint_dirs = timepoint_dirs
        self.file_suffix = file_suffix
        self._cache = {}
        
    def __len__(self):
        return len(self.codes)
    
    def load_excel_as_image(self, filepath: str) -> np.ndarray:
        """Load Excel and convert to normalized numpy array."""
        if filepath in self._cache:
            return self._cache[filepath]
            
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active
        
        data = []
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if row_idx == 0:
                continue
            row_data = []
            for col_idx, cell in enumerate(row):
                if col_idx == 0:
                    continue
                if isinstance(cell, (int, float)) and cell is not None:
                    row_data.append(float(cell))
                else:
                    row_data.append(0.0)
            if len(row_data) > 0:
                data.append(row_data)
        
        img_array = np.array(data, dtype=np.float32)
        
        min_val, max_val = img_array.min(), img_array.max()
        if max_val > min_val:
            img_array = (img_array - min_val) / (max_val - min_val)
        
        self._cache[filepath] = img_array
        return img_array
    
    def find_file(self, tp_dir: str, code: str) -> str:
        """Find the Excel file for a given code."""
        # Try direct match first
        direct_path = os.path.join(tp_dir, f"{code}{self.file_suffix}")
        if os.path.exists(direct_path):
            return direct_path
        
        # Try common patterns
        patterns = [
            f"{code}.xlsx",
            f"{code}_with_emission.xlsx",
            f"{code}.3d__570.opj_with_emission.xlsx",
        ]
        
        for pattern in patterns:
            test_path = os.path.join(tp_dir, pattern)
            if os.path.exists(test_path):
                return test_path
        
        # Search directory
        for f in os.listdir(tp_dir):
            if f.startswith(code) and f.endswith('.xlsx'):
                return os.path.join(tp_dir, f)
        
        raise FileNotFoundError(f"Could not find file for code {code} in {tp_dir}")
    
    def __getitem__(self, idx):
        code = self.codes[idx]
        label = self.labels[idx]
        
        timepoint_images = []
        for tp_dir in self.timepoint_dirs:
            filepath = self.find_file(tp_dir, code)
            img = self.load_excel_as_image(filepath)
            if len(img.shape) == 2:
                img = img[np.newaxis, :, :]
            timepoint_images.append(torch.FloatTensor(img))
        
        images = torch.stack(timepoint_images)
        return images, torch.tensor(label, dtype=torch.long), code


def load_spectral_axes(filepath: str) -> Tuple[np.ndarray, np.ndarray]:
    """Load emission and excitation axes from Excel file."""
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    
    def parse_num(x):
        if x is None:
            return np.nan
        if isinstance(x, (int, float)):
            return float(x)
        s = str(x).replace('Excitation_', '').replace('Emission_', '').replace('nm', '').strip()
        try:
            return float(s)
        except:
            return np.nan
    
    header = rows[0]
    excitation = np.array([parse_num(v) for v in header[1:] if v is not None])
    emission = np.array([parse_num(r[0]) for r in rows[1:] if r[0] is not None])
    
    excitation = excitation[~np.isnan(excitation)]
    emission = emission[~np.isnan(emission)]
    
    return emission, excitation


# ==============================================================================
# IMPORTANCE COMPUTATION
# ==============================================================================

def rank_latent_dimensions(model, loader, device, top_k: int = 20):
    """Rank latent dimensions by classification importance using logistic regression."""
    z_list, y_list = [], []
    
    model.eval()
    with torch.no_grad():
        for images, labels, _ in loader:
            images = images.to(device)
            
            latents = [model.encoder(images[:, t]) for t in range(images.shape[1])]
            z_agg, _ = model.attention(torch.stack(latents, dim=1))
            
            z_list.append(z_agg.cpu().numpy())
            y_list.append(labels.numpy())
    
    z_all = np.concatenate(z_list, axis=0)
    y_all = np.concatenate(y_list, axis=0)
    
    scaler = StandardScaler()
    z_scaled = scaler.fit_transform(z_all)
    
    clf = LogisticRegression(class_weight='balanced', max_iter=5000, solver='liblinear')
    clf.fit(z_scaled, y_all)
    
    weights = clf.coef_[0]
    top_dims = np.argsort(np.abs(weights))[-top_k:][::-1]
    signs = {int(d): np.sign(weights[d]) if weights[d] != 0 else 1.0 for d in top_dims}
    
    return top_dims, signs, z_all, y_all


def compute_gradient_saliency_per_timepoint(model, loader, dims, signs, device):
    """Compute gradient saliency |∂z_d/∂x_t| for each dimension and timepoint."""
    sample_batch, _, _ = next(iter(loader))
    _, T, _, H, W = sample_batch.shape
    
    grad_maps = {int(d): {t: np.zeros((H, W), dtype=np.float64) for t in range(T)} for d in dims}
    counts = {int(d): {t: 0 for t in range(T)} for d in dims}
    
    model.eval()
    
    for images, _, _ in loader:
        images = images.to(device).requires_grad_(True)
        
        latents = [model.encoder(images[:, t]) for t in range(images.shape[1])]
        z_agg, _ = model.attention(torch.stack(latents, dim=1))
        
        for d in dims:
            d = int(d)
            s = float(signs.get(d, 1.0))
            target = (z_agg[:, d] * s).sum()
            
            model.zero_grad()
            if images.grad is not None:
                images.grad.zero_()
            
            target.backward(retain_graph=True)
            
            grads = images.grad.detach().abs()
            
            for t in range(T):
                grad_t = grads[:, t].mean(dim=(0, 1)).cpu().numpy()
                grad_maps[d][t] += grad_t
                counts[d][t] += 1
        
        images = images.detach()
    
    for d in grad_maps:
        for t in grad_maps[d]:
            if counts[d][t] > 0:
                grad_maps[d][t] /= counts[d][t]
    
    return grad_maps


def compute_linear_probe_per_timepoint(model, loader, dims, z_all, device):
    """Compute linear probing importance: Ridge(x_t → z_d)."""
    sample_batch, _, _ = next(iter(loader))
    _, T, C, H, W = sample_batch.shape
    
    inputs_per_tp = {t: [] for t in range(T)}
    
    model.eval()
    with torch.no_grad():
        for images, _, _ in loader:
            for t in range(T):
                flat = images[:, t].numpy().reshape(images.shape[0], -1)
                inputs_per_tp[t].append(flat)
    
    for t in range(T):
        inputs_per_tp[t] = np.vstack(inputs_per_tp[t])[:z_all.shape[0]]
    
    linear_maps = {int(d): {} for d in dims}
    
    for d in dims:
        d = int(d)
        y_target = z_all[:, d]
        
        for t in range(T):
            X = inputs_per_tp[t]
            
            scaler = StandardScaler()
            X_scaled = scaler.fit_transform(X)
            
            ridge = Ridge(alpha=1.0)
            ridge.fit(X_scaled, y_target)
            
            coef = np.abs(ridge.coef_).reshape(C, H, W).mean(axis=0)
            linear_maps[d][t] = coef
    
    return linear_maps


def create_consensus_map(gradient_map: np.ndarray, linear_map: np.ndarray) -> np.ndarray:
    """Create consensus by averaging normalized gradient and linear maps."""
    def normalize(x):
        mn, mx = x.min(), x.max()
        return (x - mn) / (mx - mn + 1e-10)
    
    return (normalize(gradient_map) + normalize(linear_map)) / 2.0


def aggregate_importance_per_timepoint(gradient_maps, linear_maps, dims, T=3):
    """Aggregate consensus per timepoint across dimensions."""
    consensus_per_tp = {t: [] for t in range(T)}
    
    for d in dims:
        d = int(d)
        for t in range(T):
            cons = create_consensus_map(gradient_maps[d][t], linear_maps[d][t])
            consensus_per_tp[t].append(cons)
    
    final = {}
    for t in range(T):
        final[t] = np.mean(consensus_per_tp[t], axis=0)
    
    return final


# ==============================================================================
# ROI EXTRACTION
# ==============================================================================

def extract_rois_at_percentile(importance_map: np.ndarray, percentile: float = 95,
                               emission_axis: np.ndarray = None, 
                               excitation_axis: np.ndarray = None,
                               min_area: int = 3) -> List[Dict]:
    """Extract ROIs from importance map at given percentile threshold."""
    threshold = np.percentile(importance_map, percentile)
    mask = importance_map >= threshold
    
    mask_smooth = gaussian_filter(mask.astype(float), sigma=0.5) > 0.5
    labeled, n_components = ndimage.label(mask_smooth)
    
    rois = []
    for comp_id in range(1, n_components + 1):
        comp_mask = labeled == comp_id
        ys, xs = np.where(comp_mask)
        
        if len(xs) < min_area:
            continue
        
        y0, y1 = int(ys.min()), int(ys.max())
        x0, x1 = int(xs.min()), int(xs.max())
        
        roi = {
            'component_id': comp_id,
            'y0': y0, 'y1': y1, 'x0': x0, 'x1': x1,
            'cy': float(ys.mean()), 'cx': float(xs.mean()),
            'area_px': len(xs),
            'mean_importance': float(importance_map[comp_mask].mean()),
            'max_importance': float(importance_map[comp_mask].max()),
            'sum_importance': float(importance_map[comp_mask].sum()),
        }
        
        if emission_axis is not None and len(emission_axis) > y1:
            roi['em_min_nm'] = float(emission_axis[y0])
            roi['em_max_nm'] = float(emission_axis[min(y1, len(emission_axis)-1)])
            roi['em_centroid_nm'] = float(emission_axis[min(int(roi['cy']), len(emission_axis)-1)])
        
        if excitation_axis is not None and len(excitation_axis) > x1:
            roi['exc_min_nm'] = float(excitation_axis[x0])
            roi['exc_max_nm'] = float(excitation_axis[min(x1, len(excitation_axis)-1)])
            roi['exc_centroid_nm'] = float(excitation_axis[min(int(roi['cx']), len(excitation_axis)-1)])
        
        rois.append(roi)
    
    return sorted(rois, key=lambda r: r['mean_importance'], reverse=True)


# ==============================================================================
# VISUALIZATION
# ==============================================================================

def plot_importance_maps(maps_dict, output_path, emission_axis=None, excitation_axis=None,
                        title="Importance Maps"):
    """Plot importance maps for all timepoints."""
    tp_labels = ['0h', '6h', '24h']
    
    fig, axes = plt.subplots(1, 3, figsize=(15, 5))
    
    for t in range(3):
        ax = axes[t]
        
        if emission_axis is not None and excitation_axis is not None:
            extent = [excitation_axis[0], excitation_axis[-1],
                     emission_axis[0], emission_axis[-1]]
            im = ax.imshow(maps_dict[t], cmap='hot', aspect='auto',
                          origin='lower', extent=extent)
            ax.set_xlabel('Excitation (nm)')
            ax.set_ylabel('Emission (nm)')
        else:
            im = ax.imshow(maps_dict[t], cmap='hot', aspect='auto', origin='lower')
        
        ax.set_title(f'{tp_labels[t]}')
        plt.colorbar(im, ax=ax, fraction=0.046)
    
    plt.suptitle(title, fontsize=14, fontweight='bold')
    plt.tight_layout()
    plt.savefig(output_path, dpi=150, bbox_inches='tight')
    plt.close()


def plot_rois_on_map(importance_map, rois, output_path, 
                    emission_axis=None, excitation_axis=None, title="ROIs"):
    """Plot importance map with ROI boxes."""
    fig, ax = plt.subplots(figsize=(10, 8))
    
    if emission_axis is not None and excitation_axis is not None:
        extent = [excitation_axis[0], excitation_axis[-1],
                 emission_axis[0], emission_axis[-1]]
        im = ax.imshow(importance_map, cmap='hot', aspect='auto',
                      origin='lower', extent=extent)
        ax.set_xlabel('Excitation (nm)')
        ax.set_ylabel('Emission (nm)')
        
        for i, roi in enumerate(rois[:10]):  # Show top 10
            if 'exc_min_nm' in roi:
                rect_x = [roi['exc_min_nm'], roi['exc_max_nm'], roi['exc_max_nm'],
                         roi['exc_min_nm'], roi['exc_min_nm']]
                rect_y = [roi['em_min_nm'], roi['em_min_nm'], roi['em_max_nm'],
                         roi['em_max_nm'], roi['em_min_nm']]
                ax.plot(rect_x, rect_y, 'c-', linewidth=2)
                ax.text(roi['exc_centroid_nm'], roi['em_centroid_nm'], 
                       str(i+1), color='cyan', fontweight='bold', 
                       ha='center', va='center', fontsize=10)
    else:
        im = ax.imshow(importance_map, cmap='hot', aspect='auto', origin='lower')
    
    ax.set_title(title)
    plt.colorbar(im, ax=ax, fraction=0.046, label='Importance')
    plt.tight_layout()
    plt.savefig(output_path, dpi=150, bbox_inches='tight')
    plt.close()


# ==============================================================================
# MAIN PIPELINE
# ==============================================================================

def main():
    parser = argparse.ArgumentParser(
        description='Timepoint importance using pre-trained models',
        formatter_class=argparse.RawDescriptionHelpFormatter
    )
    
    parser.add_argument('--model_dir', type=str, required=True,
                       help='Directory containing best_model_fold*.pth files')
    parser.add_argument('--tp_dirs', type=str, required=True,
                       help='Comma-separated paths to 0h,6h,24h directories')
    parser.add_argument('--labels_csv', type=str, required=True,
                       help='CSV with columns: code, group (ALS/CTRL)')
    parser.add_argument('--output_dir', type=str, default='./importance_results')
    parser.add_argument('--n_folds', type=int, default=5)
    parser.add_argument('--top_dims', type=int, default=20)
    parser.add_argument('--latent_dim', type=int, default=512)
    parser.add_argument('--batch_size', type=int, default=4)
    parser.add_argument('--roi_percentile', type=float, default=95)
    
    args = parser.parse_args()
    
    # Setup
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    
    tp_dirs = args.tp_dirs.split(',')
    assert len(tp_dirs) == 3, "Must provide exactly 3 timepoint directories (0h, 6h, 24h)"
    
    device = torch.device('cuda' if torch.cuda.is_available() else 'cpu')
    print(f"Device: {device}")
    print(f"Model directory: {args.model_dir}")
    print(f"Output directory: {output_dir}")
    
    # Verify models exist
    model_files = []
    for fold in range(1, args.n_folds + 1):
        model_path = os.path.join(args.model_dir, f'best_model_fold{fold}.pth')
        if not os.path.exists(model_path):
            print(f"ERROR: Model not found: {model_path}")
            sys.exit(1)
        model_files.append(model_path)
    print(f"Found {len(model_files)} pre-trained models")
    
    # Load labels
    labels_df = pd.read_csv(args.labels_csv)
    if 'sample_id' in labels_df.columns and 'code' not in labels_df.columns:
        labels_df['code'] = labels_df['sample_id']
    
    codes = labels_df['code'].astype(str).tolist()
    groups = labels_df['group'].tolist()
    labels = np.array([1 if g == 'ALS' else 0 for g in groups])
    
    print(f"\nLoaded {len(codes)} samples")
    print(f"  ALS: {sum(labels)}, CTRL: {len(labels) - sum(labels)}")
    
    # Create dataset
    dataset = ExcelImageDataset(codes, labels.tolist(), tp_dirs)
    loader = DataLoader(dataset, batch_size=args.batch_size, shuffle=False, num_workers=0)
    
    # Get spectral axes from first file
    first_file = dataset.find_file(tp_dirs[0], codes[0])
    emission_axis, excitation_axis = load_spectral_axes(first_file)
    print(f"Spectral grid: {len(emission_axis)} emission × {len(excitation_axis)} excitation")
    
    # Storage
    tp_labels = ['0h', '6h', '24h']
    all_fold_consensus = {t: [] for t in range(3)}
    all_fold_gradient = {t: [] for t in range(3)}
    all_fold_linear = {t: [] for t in range(3)}
    all_fold_top_dims = []
    
    print(f"\n{'='*70}")
    print(f"PROCESSING {args.n_folds} PRE-TRAINED MODELS")
    print(f"{'='*70}")
    
    for fold_idx in range(args.n_folds):
        print(f"\n{'='*70}")
        print(f"MODEL {fold_idx + 1}/{args.n_folds}")
        print(f"{'='*70}")
        
        # Initialize model
        model = ConvAutoencoderWithAttention(
            in_channels=1,
            latent_dim=args.latent_dim,
            num_classes=2
        ).to(device)
        
        # Do a forward pass to initialize dynamic layers
        sample_batch, _, _ = next(iter(loader))
        with torch.no_grad():
            _ = model.encoder(sample_batch[0, 0:1].to(device))
        
        # Initialize decoder
        if model.encoder.encoded_shape is not None:
            target_shape = (sample_batch.shape[3], sample_batch.shape[4])
            model.init_decoder(model.encoder.encoded_shape, target_shape)
            model = model.to(device)
        
        # Load pre-trained weights
        model_path = model_files[fold_idx]
        print(f"  Loading: {model_path}")
        state_dict = torch.load(model_path, map_location=device)
        model.load_state_dict(state_dict)
        model.eval()
        print(f"  ✓ Model loaded successfully")
        
        # Step 1: Rank latent dimensions
        print(f"  Ranking latent dimensions...")
        top_dims, signs, z_all, y_all = rank_latent_dimensions(
            model, loader, device, top_k=args.top_dims
        )
        all_fold_top_dims.append(list(top_dims))
        print(f"    Top {args.top_dims} dims: {list(top_dims[:5])}... (showing first 5)")
        
        # Step 2: Gradient Saliency per timepoint
        print(f"  Computing Gradient Saliency...")
        gradient_maps = compute_gradient_saliency_per_timepoint(
            model, loader, top_dims, signs, device
        )
        
        # Step 3: Linear Probing per timepoint
        print(f"  Computing Linear Probing...")
        linear_maps = compute_linear_probe_per_timepoint(
            model, loader, top_dims, z_all, device
        )
        
        # Step 4: Consensus per timepoint
        print(f"  Creating consensus maps...")
        consensus_maps = aggregate_importance_per_timepoint(
            gradient_maps, linear_maps, top_dims, T=3
        )
        
        # Store
        for t in range(3):
            all_fold_consensus[t].append(consensus_maps[t])
            g_avg = np.mean([gradient_maps[int(d)][t] for d in top_dims], axis=0)
            l_avg = np.mean([linear_maps[int(d)][t] for d in top_dims], axis=0)
            all_fold_gradient[t].append(g_avg)
            all_fold_linear[t].append(l_avg)
        
        # Save fold-specific results
        fold_dir = output_dir / f'fold{fold_idx + 1}'
        fold_dir.mkdir(exist_ok=True)
        
        for t in range(3):
            np.save(fold_dir / f'consensus_{tp_labels[t]}.npy', consensus_maps[t])
            np.save(fold_dir / f'gradient_{tp_labels[t]}.npy', 
                   np.mean([gradient_maps[int(d)][t] for d in top_dims], axis=0))
            np.save(fold_dir / f'linear_{tp_labels[t]}.npy',
                   np.mean([linear_maps[int(d)][t] for d in top_dims], axis=0))
        
        np.save(fold_dir / 'top_dims.npy', top_dims)
        
        plot_importance_maps(
            consensus_maps, fold_dir / 'consensus_maps.png',
            emission_axis, excitation_axis,
            title=f'Fold {fold_idx+1} Consensus (Gradient + Linear)'
        )
        
        print(f"  ✓ Saved fold results to {fold_dir}")
    
    # ==============================================================================
    # CROSS-MODEL CONSENSUS
    # ==============================================================================
    
    print(f"\n{'='*70}")
    print("COMPUTING CROSS-MODEL CONSENSUS")
    print(f"{'='*70}")
    
    final_consensus = {}
    final_std = {}
    final_stability = {}
    all_rois = {}
    
    for t in range(3):
        tp = tp_labels[t]
        
        fold_maps = np.stack(all_fold_consensus[t], axis=0)
        
        mean_map = np.mean(fold_maps, axis=0)
        std_map = np.std(fold_maps, axis=0)
        stability_map = mean_map / (std_map + 1e-10)
        
        final_consensus[t] = mean_map
        final_std[t] = std_map
        final_stability[t] = stability_map
        
        np.save(output_dir / f'final_consensus_{tp}.npy', mean_map)
        np.save(output_dir / f'fold_std_{tp}.npy', std_map)
        np.save(output_dir / f'stability_{tp}.npy', stability_map)
        
        # Extract ROIs at 95th percentile
        rois = extract_rois_at_percentile(
            mean_map, args.roi_percentile, emission_axis, excitation_axis
        )
        
        # Add stability info
        for roi in rois:
            y0, y1 = roi['y0'], roi['y1']
            x0, x1 = roi['x0'], roi['x1']
            roi['stability'] = float(stability_map[y0:y1+1, x0:x1+1].mean())
            roi['fold_std'] = float(std_map[y0:y1+1, x0:x1+1].mean())
            roi['timepoint'] = tp
        
        all_rois[tp] = rois
        
        print(f"\n  {tp}: {len(rois)} ROIs at {args.roi_percentile}th percentile")
        for i, roi in enumerate(rois[:5]):
            em_str = f"{roi.get('em_centroid_nm', roi['cy']):.0f}nm"
            exc_str = f"{roi.get('exc_centroid_nm', roi['cx']):.0f}nm"
            print(f"    ROI {i+1}: em~{em_str}, exc~{exc_str}, "
                  f"imp={roi['mean_importance']:.4f}, stability={roi['stability']:.2f}")
        
        plot_rois_on_map(
            mean_map, rois,
            output_dir / f'final_consensus_{tp}_with_rois.png',
            emission_axis, excitation_axis,
            title=f'{tp} Final Consensus (avg of {args.n_folds} models) - {args.roi_percentile}th percentile'
        )
    
    # Combined visualization
    plot_importance_maps(
        final_consensus,
        output_dir / 'final_consensus_all_timepoints.png',
        emission_axis, excitation_axis,
        title=f'Final Consensus Importance Maps (avg of {args.n_folds} models)'
    )
    
    # Save ROIs
    rois_flat = []
    for tp, rois in all_rois.items():
        rois_flat.extend(rois)
    pd.DataFrame(rois_flat).to_csv(output_dir / 'rois_95th_percentile.csv', index=False)
    
    # Summary
    print(f"\n{'='*70}")
    print("SUMMARY")
    print(f"{'='*70}")
    
    summary = [
        "TIMEPOINT-SPECIFIC FEATURE IMPORTANCE (Pre-trained Models)",
        "=" * 60,
        f"\nModels loaded from: {args.model_dir}",
        f"Number of folds: {args.n_folds}",
        f"Top dimensions per model: {args.top_dims}",
        f"Latent dimension: {args.latent_dim}",
        f"ROI percentile: {args.roi_percentile}",
        f"\nSamples: {len(codes)} ({sum(labels)} ALS, {len(labels)-sum(labels)} CTRL)",
        f"Spectral grid: {len(emission_axis)} em × {len(excitation_axis)} exc",
        f"\nROIs per timepoint:",
    ]
    for tp in tp_labels:
        summary.append(f"  - {tp}: {len(all_rois[tp])} ROIs")
    
    summary.append(f"\nTop ROIs by importance:")
    for tp in tp_labels:
        summary.append(f"\n  {tp}:")
        for i, roi in enumerate(all_rois[tp][:3]):
            em = roi.get('em_centroid_nm', roi['cy'])
            exc = roi.get('exc_centroid_nm', roi['cx'])
            summary.append(f"    {i+1}. em~{em:.0f}nm, exc~{exc:.0f}nm, "
                          f"imp={roi['mean_importance']:.4f}, stability={roi['stability']:.2f}")
    
    summary_text = '\n'.join(summary)
    print(summary_text)
    
    with open(output_dir / 'SUMMARY.txt', 'w') as f:
        f.write(summary_text)
    
    # Config
    config = {
        'model_dir': args.model_dir,
        'n_folds': args.n_folds,
        'top_dims': args.top_dims,
        'latent_dim': args.latent_dim,
        'roi_percentile': args.roi_percentile,
        'n_samples': len(codes),
        'fold_top_dims': [[int(d) for d in dims] for dims in all_fold_top_dims],
    }
    with open(output_dir / 'config.json', 'w') as f:
        json.dump(config, f, indent=2)
    
    print(f"\n{'='*70}")
    print(f"Results saved to: {output_dir}")
    print(f"{'='*70}")


if __name__ == '__main__':
    main()
