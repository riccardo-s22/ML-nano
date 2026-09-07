"""
Convolutional Autoencoder with Attention-Based Temporal Aggregation

This architecture implements a dual-objective deep learning model for binary classification
of subjects with multiple timepoint observations. The model simultaneously:
1. Learns to reconstruct input images (unsupervised feature learning)
2. Classifies subjects into two groups (supervised classification)

Key Design Principles:
- Shared encoder extracts features useful for both reconstruction and classification
- Attention mechanism learns which timepoints are most discriminative
- Reconstruction acts as regularizer to prevent overfitting on limited data
- Subject-level cross-validation prevents data leakage between timepoints
"""

import torch
import torch.nn as nn
import torch.optim as optim
from torch.utils.data import Dataset, DataLoader
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from sklearn.model_selection import StratifiedKFold
from sklearn.preprocessing import LabelEncoder
import os
from typing import Tuple, List
import matplotlib.pyplot as plt

import random
import argparse
from pathlib import Path
from sklearn.metrics import confusion_matrix, roc_auc_score, accuracy_score, precision_score, recall_score, f1_score, balanced_accuracy_score


# Set random seeds for reproducibility
torch.manual_seed(42)
np.random.seed(42)
class ExcelImageDataset(Dataset):
    """
    Custom dataset for loading Excel files as image-like arrays with multiple timepoints.
    
    Design rationale:
    - Each subject has 3 Excel files (one per timepoint: 0h, 6h, 24h)
    - Excel files are treated as 2D numerical arrays (analogous to images)
    - All 3 timepoints are loaded together to maintain subject-level structure
    - First row (headers) and first column (emission labels) are excluded from data
    """
    
    def __init__(self, codes: List[str], labels: List[int], 
                 timepoint_dirs: List[str], transform=None):
        """
        Args:
            codes: Subject identifiers (filenames without extension)
            labels: Binary group labels for classification
            timepoint_dirs: List of 3 directories containing timepoint data
            transform: Optional data augmentation (not currently used)
        """
        self.codes = codes
        self.labels = labels
        self.timepoint_dirs = timepoint_dirs
        self.transform = transform
        
    def __len__(self):
        return len(self.codes)
    
    def load_excel_as_image(self, filepath: str) -> np.ndarray:
        """
        Load Excel file and convert to numpy array, treating it as image data.
        
        Processing steps:
        1. Skip first row (column headers like "Emission", "Excitation_...")
        2. Skip first column (row labels like emission wavelengths)
        3. Extract only the numeric data matrix
        4. Normalize to [0, 1] range for neural network input
        
        Returns:
            2D numpy array of shape (n_rows, n_cols) normalized to [0, 1]
        """
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active
        
        # Extract numerical data from worksheet, skipping first row (headers)
        data = []
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            # Skip first row (column headers)
            if row_idx == 0:
                continue
            
            row_data = []
            for col_idx, cell in enumerate(row):
                # Skip first column (row labels/emission values)
                if col_idx == 0:
                    continue
                
                # Convert to float, handling None and non-numeric values
                if isinstance(cell, (int, float)) and cell is not None:
                    row_data.append(float(cell))
                else:
                    row_data.append(0.0)
            
            # Only add row if it has data
            if len(row_data) > 0:
                data.append(row_data)
        
        # Convert to numpy array
        img_array = np.array(data, dtype=np.float32)
        
        # Check if we got valid data
        if img_array.size == 0:
            raise ValueError(f"No valid data found in {filepath}")
        
        # Normalize to [0, 1] range for stable neural network training
        min_val = img_array.min()
        max_val = img_array.max()
        if max_val > min_val:
            img_array = (img_array - min_val) / (max_val - min_val)
        else:
            # If all values are the same, just return zeros
            img_array = np.zeros_like(img_array)
        
        return img_array
    
    def __getitem__(self, idx):
        """
        Load all 3 timepoint images for a single subject.
        
        Returns:
            images: Tensor of shape (3, 1, H, W) - 3 timepoints, 1 channel, H rows, W columns
            label: Integer label (0 or 1) for binary classification
        """
        code = self.codes[idx]
        label = self.labels[idx]
        
        # Load images from all three timepoints (0h, 6h, 24h)
        timepoint_images = []
        for tp_dir in self.timepoint_dirs:
            filepath = os.path.join(tp_dir, f"{code}.xlsx")
            img = self.load_excel_as_image(filepath)
            
            # Add channel dimension (grayscale image has 1 channel)
            # Shape: (H, W) -> (1, H, W)
            if len(img.shape) == 2:
                img = img[np.newaxis, :, :]
            
            timepoint_images.append(torch.FloatTensor(img))
        
        # Stack timepoints: shape becomes (3, C, H, W)
        # This maintains the temporal structure for the attention mechanism
        images = torch.stack(timepoint_images)
        
        if self.transform:
            # Apply same transform to all timepoints to maintain consistency
            images = torch.stack([self.transform(img) for img in images])
        
        return images, torch.tensor(label, dtype=torch.long)


class ConvEncoder(nn.Module):
    """
    Convolutional encoder that extracts hierarchical features from images.
    
    Architecture Philosophy:
    - Convolutional layers learn spatial patterns (edges, textures, structures)
    - Each layer learns increasingly abstract features through the hierarchy
    - MaxPooling progressively reduces spatial dimensions (downsampling)
    - BatchNorm stabilizes training and acts as regularization
    - ReLU introduces non-linearity for learning complex patterns
    
    The encoder is SHARED across all timepoints, ensuring consistent feature extraction.
    This is critical for the attention mechanism to meaningfully compare timepoints.
    """
    
    def __init__(self, in_channels: int, latent_dim: int):
        """
        Args:
            in_channels: Number of input channels (1 for grayscale data)
            latent_dim: Dimension of the compressed latent representation
        """
        super(ConvEncoder, self).__init__()
        
        # First convolutional block
        # Purpose: Learn low-level features (edges, gradients, basic textures)
        # Conv2d: Applies learnable filters that detect spatial patterns
        # BatchNorm2d: Normalizes activations for stable training
        # ReLU: Non-linear activation (outputs max(0, x))
        # MaxPool2d: Downsamples by taking maximum in 2x2 windows
        self.conv1 = nn.Sequential(
            nn.Conv2d(in_channels, 32, kernel_size=3, padding=1),  # 32 filters
            nn.BatchNorm2d(32),
            nn.ReLU(),
            nn.MaxPool2d(2, 2)  # Reduces spatial dimensions by half
        )
        
        # Second convolutional block
        # Purpose: Combine low-level features into mid-level patterns
        self.conv2 = nn.Sequential(
            nn.Conv2d(32, 64, kernel_size=3, padding=1),  # 64 filters
            nn.BatchNorm2d(64),
            nn.ReLU(),
            nn.MaxPool2d(2, 2)  # Further spatial reduction
        )
        
        # Third convolutional block
        # Purpose: Learn high-level abstract features
        self.conv3 = nn.Sequential(
            nn.Conv2d(64, 128, kernel_size=3, padding=1),  # 128 filters
            nn.BatchNorm2d(128),
            nn.ReLU(),
            nn.MaxPool2d(2, 2)  # Another spatial reduction
        )
        
        # Fourth convolutional block
        # Purpose: Learn highest-level abstract representations
        # Note: No pooling here - maintains spatial structure before flattening
        self.conv4 = nn.Sequential(
            nn.Conv2d(128, 256, kernel_size=3, padding=1),  # 256 filters
            nn.BatchNorm2d(256),
            nn.ReLU()
        )
        
        # These will be initialized dynamically based on input dimensions
        # This allows the model to work with any input size
        self.flatten_size = None
        self.fc = None  # Fully connected layer to compress to latent_dim
        self.latent_dim = latent_dim
        self.encoded_shape = None  # Store shape for decoder initialization
        self.input_shape = None  # Store original input shape for exact reconstruction
        
    def forward(self, x):
        """
        Forward pass through the encoder.
        
        Args:
            x: Input tensor of shape (batch, channels, height, width)
        
        Returns:
            Latent vector of shape (batch, latent_dim)
        
        Process:
        1. Store input dimensions for decoder
        2. Apply convolutional blocks sequentially (creates feature hierarchy)
        3. Flatten spatial features into 1D vector
        4. Project to latent space via fully connected layer
        """
        # Store original input shape for decoder to match exactly
        if self.input_shape is None:
            self.input_shape = (x.shape[2], x.shape[3])
        
        # Apply convolutional hierarchy
        # Each layer extracts increasingly abstract features
        x = self.conv1(x)  # Low-level features
        x = self.conv2(x)  # Mid-level patterns
        x = self.conv3(x)  # High-level structures
        x = self.conv4(x)  # Abstract representations
        
        # Store the encoded shape (for decoder initialization)
        if self.encoded_shape is None:
            self.encoded_shape = (x.shape[1], x.shape[2], x.shape[3])
        
        # Initialize FC layer on first forward pass (now we know dimensions)
        # This dynamic initialization allows flexibility with input sizes
        if self.fc is None:
            self.flatten_size = x.shape[1] * x.shape[2] * x.shape[3]
            self.fc = nn.Linear(self.flatten_size, self.latent_dim).to(x.device)
        
        # Flatten spatial dimensions: (batch, channels, H, W) -> (batch, channels*H*W)
        x = x.view(x.size(0), -1)
        
        # Project to latent space: (batch, flatten_size) -> (batch, latent_dim)
        # This is the compressed representation used by both decoder and classifier
        x = self.fc(x)
        
        return x


class ConvDecoder(nn.Module):
    """
    Convolutional decoder that reconstructs images from latent representations.
    
    Architecture Philosophy:
    - Mirror of the encoder, but using transpose convolutions for upsampling
    - ConvTranspose2d "reverses" the pooling operations to increase spatial dimensions
    - Final interpolation ensures exact dimension matching with input
    
    Purpose in Dual-Objective Framework:
    - Reconstruction loss forces encoder to preserve information
    - Acts as regularization to prevent overfitting on limited subjects
    - Ensures learned features are information-rich, not just discriminative
    """
    
    def __init__(self, latent_dim: int, out_channels: int, encoded_shape: tuple, target_shape: tuple):
        """
        Args:
            latent_dim: Dimension of latent representation
            out_channels: Number of output channels (1 for grayscale)
            encoded_shape: Shape after encoding (channels, height, width)
            target_shape: Original input dimensions (height, width) for exact reconstruction
        """
        super(ConvDecoder, self).__init__()
        
        # Store shapes for reconstruction
        self.encoded_channels, self.encoded_h, self.encoded_w = encoded_shape
        self.target_h, self.target_w = target_shape
        
        # Project latent vector back to spatial feature maps
        self.fc = nn.Linear(latent_dim, self.encoded_channels * self.encoded_h * self.encoded_w)
        
        # First transpose convolutional block
        # ConvTranspose2d with stride=2 doubles spatial dimensions
        # This "reverses" the MaxPool operations from encoding
        self.deconv1 = nn.Sequential(
            nn.ConvTranspose2d(256, 128, kernel_size=3, stride=2, padding=1, output_padding=1),
            nn.BatchNorm2d(128),
            nn.ReLU()
        )
        
        # Second transpose convolutional block
        self.deconv2 = nn.Sequential(
            nn.ConvTranspose2d(128, 64, kernel_size=3, stride=2, padding=1, output_padding=1),
            nn.BatchNorm2d(64),
            nn.ReLU()
        )
        
        # Third transpose convolutional block
        self.deconv3 = nn.Sequential(
            nn.ConvTranspose2d(64, 32, kernel_size=3, stride=2, padding=1, output_padding=1),
            nn.BatchNorm2d(32),
            nn.ReLU()
        )
        
        # Final convolution to get output channels
        self.final_conv = nn.Conv2d(32, out_channels, kernel_size=3, padding=1)
        
        # Sigmoid ensures output is in [0, 1] range (matching normalized input)
        self.sigmoid = nn.Sigmoid()
        
    def forward(self, z):
        """
        Reconstruct image from latent representation.
        
        Args:
            z: Latent vector of shape (batch, latent_dim)
        
        Returns:
            Reconstructed image of shape (batch, channels, height, width)
        
        Process:
        1. Project latent to spatial features
        2. Apply transpose convolutions to upsample
        3. Interpolate to exact input dimensions
        4. Apply sigmoid for [0, 1] output range
        """
        # Project latent back to feature maps
        x = self.fc(z)
        x = x.view(x.size(0), self.encoded_channels, self.encoded_h, self.encoded_w)
        
        # Apply upsampling hierarchy (mirror of encoder)
        x = self.deconv1(x)
        x = self.deconv2(x)
        x = self.deconv3(x)
        x = self.final_conv(x)
        
        # Adaptive resizing to match exact target dimensions
        # This handles cases where transpose convolutions don't perfectly reverse pooling
        # (e.g., 71 columns -> after pooling/unpooling might become 64 or 72)
        current_h, current_w = x.shape[2], x.shape[3]
        
        if current_h != self.target_h or current_w != self.target_w:
            # Use bilinear interpolation to resize to exact dimensions
            x = torch.nn.functional.interpolate(
                x, 
                size=(self.target_h, self.target_w), 
                mode='bilinear', 
                align_corners=False
            )
        
        # Apply sigmoid to constrain output to [0, 1]
        x = self.sigmoid(x)
        return x


class AttentionAggregation(nn.Module):
    """
    Attention mechanism for aggregating multiple timepoint latent vectors.
    
    Design Philosophy - Approach 3 from our discussion:
    - Learns which timepoints are most important for classification
    - Weights (α1, α2, α3) are LEARNED during training, not fixed
    - Allows model to discover temporal patterns in the data
    - Provides interpretability: we can see which timepoints matter most
    
    How it works:
    1. Small neural network scores each timepoint's importance
    2. Softmax ensures weights sum to 1 (valid probability distribution)
    3. Weighted sum creates aggregated representation for classification
    
    Example learned weights:
    - α1=0.7, α2=0.2, α3=0.1: Early timepoint (0h) most discriminative
    - α1=0.1, α2=0.2, α3=0.7: Late timepoint (24h) most discriminative
    - α1=0.33, α2=0.33, α3=0.34: All timepoints equally important
    """
    
    def __init__(self, latent_dim: int):
        """
        Args:
            latent_dim: Dimension of each latent vector
        """
        super(AttentionAggregation, self).__init__()
        
        # Small neural network that learns to score timepoint importance
        # Input: latent vector (latent_dim)
        # Output: single score indicating importance
        self.attention_net = nn.Sequential(
            nn.Linear(latent_dim, 64),  # Project to smaller space
            nn.Tanh(),  # Non-linear activation
            nn.Linear(64, 1)  # Output single importance score
        )
        
    def forward(self, latent_vectors):
        """
        Compute attention-weighted aggregation of timepoint latent vectors.
        
        Args:
            latent_vectors: Tensor of shape (batch, num_timepoints, latent_dim)
                           e.g., (batch, 3, 512) for 3 timepoints
        
        Returns:
            aggregated: Weighted sum of shape (batch, latent_dim)
            attention_weights: Learned weights of shape (batch, num_timepoints)
                              These tell us which timepoints the model considers important
        
        Process:
        1. Compute importance score for each timepoint
        2. Normalize scores with softmax (makes them sum to 1)
        3. Weighted sum: z_agg = α1·z1 + α2·z2 + α3·z3
        """
        # Compute attention scores for each timepoint
        # Shape: (batch, num_timepoints, 1)
        attention_scores = self.attention_net(latent_vectors)
        attention_scores = attention_scores.squeeze(-1)  # Shape: (batch, num_timepoints)
        
        # Apply softmax to get weights that sum to 1
        # This ensures weights form a valid probability distribution
        attention_weights = torch.softmax(attention_scores, dim=1)  # Shape: (batch, num_timepoints)
        
        # Weighted sum of latent vectors
        # Expand weights for broadcasting: (batch, num_timepoints) -> (batch, num_timepoints, 1)
        attention_weights_expanded = attention_weights.unsqueeze(-1)
        
        # Weighted sum: (batch, num_timepoints, latent_dim) -> (batch, latent_dim)
        aggregated = torch.sum(latent_vectors * attention_weights_expanded, dim=1)
        
        return aggregated, attention_weights


class Classifier(nn.Module):
    """
    Classification head for binary group prediction.
    
    Takes the aggregated latent representation and predicts group membership.
    
    Design elements:
    - Fully connected layers for classification
    - Dropout for regularization (prevents overfitting on limited data)
    - Final layer outputs logits for 2 classes
    - Softmax (applied during loss calculation) provides confidence scores
    """
    
    def __init__(self, latent_dim: int, num_classes: int, dropout_rate: float = 0.3):
        """
        Args:
            latent_dim: Dimension of aggregated latent vector
            num_classes: Number of output classes (2 for binary classification)
            dropout_rate: Probability of dropping neurons during training
        """
        super(Classifier, self).__init__()
        
        self.classifier = nn.Sequential(
            nn.Linear(latent_dim, 128),  # Project to smaller space
            nn.ReLU(),
            nn.Dropout(dropout_rate),  # Regularization to prevent overfitting
            nn.Linear(128, num_classes)  # Output logits for each class
        )
        
    def forward(self, z_agg):
        """
        Predict class from aggregated latent representation.
        
        Args:
            z_agg: Aggregated latent vector of shape (batch, latent_dim)
        
        Returns:
            Class logits of shape (batch, num_classes)
        """
        return self.classifier(z_agg)


class ConvAutoencoderWithAttention(nn.Module):
    """
    Complete dual-objective model integrating all components.
    
    ARCHITECTURE OVERVIEW:
    
    Input: 3 images per subject (timepoints t1, t2, t3)
           ↓
    Shared Encoder (processes each timepoint independently)
           ↓
    3 Latent Vectors (z1, z2, z3)
           ↓
           ├─→ Attention Aggregation → z_agg → Classifier → Group Prediction
           │                                                  (SUPERVISED)
           │
           └─→ Individual Decoders → Reconstructed Images
                                      (UNSUPERVISED)
    
    KEY DESIGN DECISIONS:
    
    1. SHARED ENCODER:
       - Same weights process all timepoints
       - Ensures features are comparable across time
       - Critical for attention mechanism to work
    
    2. DUAL OBJECTIVES:
       - Reconstruction: Forces features to be information-rich
       - Classification: Forces features to be discriminative
       - Both objectives share the encoder (joint optimization)
    
    3. ATTENTION AGGREGATION:
       - Collapses 3 latent vectors into 1 for classification
       - Learns which timepoints matter most
       - Provides interpretability
    
    4. SEPARATE RECONSTRUCTION:
       - Each latent vector reconstructs its own timepoint
       - Maintains correspondence: z1→t1, z2→t2, z3→t3
       - Cleaner than trying to reconstruct all from aggregated vector
    
    WHY THIS DESIGN?
    - With limited data (20 subjects), reconstruction provides additional supervision
    - Attention mechanism handles non-independence between timepoints
    - Dual objectives act as mutual regularizers
    - Architecture respects the temporal structure of the data
    """
    
    def __init__(self, in_channels: int, latent_dim: int, num_classes: int):
        """
        Args:
            in_channels: Number of input channels (1 for grayscale)
            latent_dim: Dimension of latent space
            num_classes: Number of output classes (2 for binary)
        """
        super(ConvAutoencoderWithAttention, self).__init__()
        
        self.encoder = ConvEncoder(in_channels, latent_dim)
        self.attention = AttentionAggregation(latent_dim)
        self.decoder = None  # Will be initialized after first forward pass
        self.classifier = Classifier(latent_dim, num_classes)
        self.in_channels = in_channels
        self.latent_dim = latent_dim
        
    def forward(self, timepoint_images):
        """
        Forward pass through the complete architecture.
        
        Args:
            timepoint_images: Tensor of shape (batch, num_timepoints, channels, height, width)
                             e.g., (4, 3, 1, 512, 71) for batch of 4 subjects
        
        Returns:
            reconstructions: Reconstructed images (batch, num_timepoints, channels, height, width)
            class_logits: Classification predictions (batch, num_classes)
            attention_weights: Learned timepoint importance (batch, num_timepoints)
        
        STEP-BY-STEP PROCESS:
        
        1. ENCODING (creating latent representations):
           - Process each timepoint through shared encoder
           - Produces 3 separate latent vectors per subject
           - Each latent is a compressed representation of one timepoint
        
        2. ATTENTION AGGREGATION (for classification):
           - Combine 3 latent vectors using learned weights
           - Results in single aggregated representation per subject
           - This aggregated vector is used for classification
        
        3. RECONSTRUCTION (for regularization):
           - Each individual latent vector reconstructs its timepoint
           - Ensures encoder preserves information
           - Acts as unsupervised regularization
        
        4. CLASSIFICATION (supervised prediction):
           - Aggregated latent predicts group membership
           - Outputs logits that become probabilities via softmax
        """
        batch_size, num_timepoints = timepoint_images.shape[0], timepoint_images.shape[1]
        
        # STEP 1: ENCODE EACH TIMEPOINT INDEPENDENTLY
        # This maintains consistency: same encoder for all timepoints
        # Critical for attention mechanism to compare timepoints meaningfully
        latent_vectors = []
        for t in range(num_timepoints):
            z_t = self.encoder(timepoint_images[:, t])  # Shape: (batch, latent_dim)
            latent_vectors.append(z_t)
        
        # Initialize decoder after first encoding (now we know dimensions)
        # Dynamic initialization allows model to work with any input size
        if self.decoder is None and self.encoder.encoded_shape is not None:
            self.decoder = ConvDecoder(
                self.latent_dim, 
                self.in_channels, 
                self.encoder.encoded_shape,
                self.encoder.input_shape  # Pass target shape for exact reconstruction
            ).to(timepoint_images.device)
        
        # Stack latent vectors: (batch, num_timepoints, latent_dim)
        # This groups all timepoints for a subject together
        latent_vectors = torch.stack(latent_vectors, dim=1)
        
        # STEP 2: AGGREGATE WITH ATTENTION (for classification branch)
        # Learns which timepoints are most important for distinguishing groups
        # Output is single vector per subject
        z_agg, attention_weights = self.attention(latent_vectors)
        
        # STEP 3: RECONSTRUCT FROM INDIVIDUAL LATENT VECTORS
        # Each timepoint is reconstructed from its own latent representation
        # This is cleaner than reconstructing all from aggregated vector
        # Maintains correspondence: z1→image1, z2→image2, z3→image3
        reconstructions = []
        for t in range(num_timepoints):
            recon_t = self.decoder(latent_vectors[:, t])
            reconstructions.append(recon_t)
        
        # Stack reconstructions: (batch, num_timepoints, channels, height, width)
        reconstructions = torch.stack(reconstructions, dim=1)
        
        # STEP 4: CLASSIFY FROM AGGREGATED REPRESENTATION
        # Uses the attention-weighted combination of all timepoints
        # Outputs logits that become probabilities via softmax
        class_logits = self.classifier(z_agg)
        
        return reconstructions, class_logits, attention_weights


def train_epoch(model, dataloader, optimizer, device, alpha=1.0, beta=1.0):
    """
    One training epoch for the dual-objective convolutional autoencoder.

    Parameters
    ----------
    model : nn.Module
        Autoencoder + classifier model.
    dataloader : DataLoader
        Training dataloader.
    optimizer : torch.optim.Optimizer
        Optimizer for model parameters.
    device : torch.device
        CPU or CUDA device.
    alpha : float
        Weight for reconstruction loss.
    beta : float
        Weight for classification loss.

    Returns
    -------
    avg_loss : float
        Average total loss over the epoch.
    avg_recon_loss : float
        Average reconstruction (MSE) loss.
    avg_class_loss : float
        Average classification (cross-entropy) loss.
    accuracy : float
        Classification accuracy (percent) over the epoch.
    attention_weights_all : np.ndarray
        Concatenated attention weights for all batches.
    """

    model.train()

    total_loss = 0.0
    total_recon_loss = 0.0
    total_class_loss = 0.0
    correct = 0
    total = 0

    reconstruction_criterion = nn.MSELoss()
    classification_criterion = nn.CrossEntropyLoss()

    all_attention_weights = []

    for images, labels in dataloader:
        images = images.to(device)
        labels = labels.to(device)

        optimizer.zero_grad()

        # Forward pass
        reconstructions, class_logits, attention_weights = model(images)

        # Reconstruction loss: how well we reconstruct the input image
        recon_loss = reconstruction_criterion(reconstructions, images)

        # Classification loss: how well we predict ALS vs control
        class_loss = classification_criterion(class_logits, labels)

        # Combined loss
        loss = alpha * recon_loss + beta * class_loss

        # Backpropagation and parameter update
        loss.backward()
        optimizer.step()

        # Accumulate scalar metrics
        total_loss += loss.item()
        total_recon_loss += recon_loss.item()
        total_class_loss += class_loss.item()

        # Classification accuracy
        _, predicted = torch.max(class_logits, 1)
        total += labels.size(0)
        correct += (predicted == labels).sum().item()

        # Collect attention weights for analysis
        all_attention_weights.append(attention_weights.detach().cpu().numpy())

    avg_loss = total_loss / len(dataloader)
    avg_recon_loss = total_recon_loss / len(dataloader)
    avg_class_loss = total_class_loss / len(dataloader)
    accuracy = 100.0 * correct / total if total > 0 else 0.0

    attention_weights_all = np.concatenate(all_attention_weights, axis=0) if all_attention_weights else np.array([])

    return avg_loss, avg_recon_loss, avg_class_loss, accuracy, attention_weights_all



def set_seed(seed: int = 42) -> None:
    """Set random seeds for reproducibility."""
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    if torch.cuda.is_available():
        torch.cuda.manual_seed_all(seed)


def _infer_label_columns(labels_df: pd.DataFrame):
    """
    Infer code/subject and label columns from a labels CSV.

    Supported common schemas:
      - code,label
      - filename,label (filename may include .xlsx)
      - subject,label
      - sample,label
    """
    cols = set([c.lower() for c in labels_df.columns])

    # label column
    label_col = None
    for cand in ["label", "class", "group", "y"]:
        if cand in cols:
            label_col = next(c for c in labels_df.columns if c.lower() == cand)
            break
    if label_col is None:
        raise ValueError("labels_csv must contain a label column (e.g., 'label' or 'class').")

    # code/subject column
    code_col = None
    for cand in ["code", "subject", "sample", "filename", "id"]:
        if cand in cols:
            code_col = next(c for c in labels_df.columns if c.lower() == cand)
            break
    if code_col is None:
        raise ValueError("labels_csv must contain a sample identifier column (e.g., 'code', 'subject', 'sample', or 'filename').")

    return code_col, label_col


def _normalize_code(code: str) -> str:
    """Normalize a sample code coming from CSV (e.g., strip extension)."""
    code = str(code).strip()
    if code.lower().endswith(".xlsx"):
        code = code[:-5]
    if code.lower().endswith(".xls"):
        code = code[:-4]
    return code


def build_subject_splits(labels_df: pd.DataFrame, code_col: str, label_col: str, num_folds: int, seed: int):
    """
    Subject-level stratified CV splits.

    If the CSV contains a distinct 'subject' column and separate 'filename'/timepoint rows, you can adapt this
    function; by default, we treat code as the subject identifier.
    """
    # Treat code as subject identifier by default
    labels_df = labels_df.copy()
    labels_df["_code"] = labels_df[code_col].apply(_normalize_code)

    # If multiple rows per code exist, we require consistent labels and take the first
    per_code = labels_df.groupby("_code")[label_col].agg(lambda s: s.iloc[0]).reset_index()
    y = per_code[label_col].values
    codes = per_code["_code"].values

    skf = StratifiedKFold(n_splits=num_folds, shuffle=True, random_state=seed)
    splits = []
    for train_idx, val_idx in skf.split(codes, y):
        train_codes = codes[train_idx].tolist()
        val_codes = codes[val_idx].tolist()
        splits.append((train_codes, val_codes))

    return splits, per_code


def compute_binary_metrics(y_true, y_prob, threshold: float = 0.5):
    """
    Compute confusion matrix + common classification metrics for the positive class = 1.
    """
    y_true = np.asarray(y_true).astype(int)
    y_prob = np.asarray(y_prob).astype(float)
    y_pred = (y_prob >= threshold).astype(int)

    cm = confusion_matrix(y_true, y_pred, labels=[0, 1])
    tn, fp, fn, tp = cm.ravel() if cm.size == 4 else (0, 0, 0, 0)

    # Core metrics
    acc = accuracy_score(y_true, y_pred) if len(y_true) else float("nan")
    bal_acc = balanced_accuracy_score(y_true, y_pred) if len(np.unique(y_true)) > 1 else float("nan")
    prec = precision_score(y_true, y_pred, zero_division=0)
    rec = recall_score(y_true, y_pred, zero_division=0)  # sensitivity
    f1 = f1_score(y_true, y_pred, zero_division=0)

    # Specificity
    spec = (tn / (tn + fp)) if (tn + fp) > 0 else float("nan")

    # AUC
    try:
        auc = roc_auc_score(y_true, y_prob) if len(np.unique(y_true)) > 1 else float("nan")
    except ValueError:
        auc = float("nan")

    return {
        "tn": int(tn), "fp": int(fp), "fn": int(fn), "tp": int(tp),
        "accuracy": float(acc),
        "balanced_accuracy": float(bal_acc),
        "precision": float(prec),
        "recall_sensitivity": float(rec),
        "specificity": float(spec),
        "f1": float(f1),
        "auc": float(auc),
        "threshold": float(threshold),
    }


def main():
    parser = argparse.ArgumentParser(
        description="Train ConvAutoencoderWithAttention on Excel EEM matrices with subject-level CV and per-fold metrics."
    )
    parser.add_argument("--labels_csv", type=str, required=True, help="CSV containing at least [code/subject/filename] and [label].")
    parser.add_argument("--data_dir", type=str, required=True, help="Root directory containing timepoint subfolders with .xlsx files.")
    parser.add_argument("--tp_dirs", type=str, required=True, help="Comma-separated subfolder paths for timepoints (e.g., out_0h,out_6h,out_24h).")
    parser.add_argument("--output_dir", type=str, default="./results", help="Directory to write models and metrics.")
    parser.add_argument("--latent_dim", type=int, default=128, help="Latent dimension.")
    parser.add_argument("--batch_size", type=int, default=4, help="Batch size (subjects).")
    parser.add_argument("--epochs", type=int, default=100, help="Max epochs per fold.")
    parser.add_argument("--learning_rate", type=float, default=1e-3, help="Learning rate.")
    parser.add_argument("--alpha", type=float, default=1.0, help="Weight for reconstruction loss.")
    parser.add_argument("--beta", type=float, default=1.0, help="Weight for classification loss.")
    parser.add_argument("--num_folds", type=int, default=5, help="Number of subject-level folds.")
    parser.add_argument("--patience", type=int, default=15, help="Early stopping patience (epochs).")
    parser.add_argument("--min_delta", type=float, default=1e-4, help="Minimum improvement in val loss to reset patience.")
    parser.add_argument("--seed", type=int, default=42, help="Random seed.")
    parser.add_argument("--threshold", type=float, default=0.5, help="Probability threshold for confusion matrix-based metrics.")
    args = parser.parse_args()

    set_seed(args.seed)
    os.makedirs(args.output_dir, exist_ok=True)

    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    print(f"Using device: {device}")

    labels_df = pd.read_csv(args.labels_csv)
    code_col, label_col = _infer_label_columns(labels_df)

    # Convert labels to 0/1 if they are strings
    # Keep stable ordering for binary classification
    if labels_df[label_col].dtype == object:
        le = LabelEncoder()
        labels_df[label_col] = le.fit_transform(labels_df[label_col].astype(str))
        print(f"Label mapping: {dict(zip(le.classes_, le.transform(le.classes_)))}")

    # Prepare timepoint directories
    tp_dirs = [d.strip() for d in args.tp_dirs.split(",") if d.strip()]
    if len(tp_dirs) < 1:
        raise ValueError("tp_dirs must contain at least one directory.")
    tp_dirs = [os.path.join(args.data_dir, d) if not os.path.isabs(d) else d for d in tp_dirs]

    for d in tp_dirs:
        if not os.path.isdir(d):
            raise FileNotFoundError(f"Timepoint directory not found: {d}")

    # Subject-level splits
    splits, per_code = build_subject_splits(labels_df, code_col, label_col, args.num_folds, args.seed)
    print(f"Total subjects: {len(per_code)}")

    fold_metrics = []
    attention_summary = []

    for fold_idx, (train_codes, val_codes) in enumerate(splits, start=1):
        print(f"\n===== FOLD {fold_idx}/{args.num_folds} =====")
        train_labels = per_code[per_code["_code"].isin(train_codes)][label_col].tolist()
        val_labels = per_code[per_code["_code"].isin(val_codes)][label_col].tolist()

        train_dataset = ExcelImageDataset(train_codes, train_labels, tp_dirs)
        val_dataset = ExcelImageDataset(val_codes, val_labels, tp_dirs)

        train_loader = DataLoader(train_dataset, batch_size=args.batch_size, shuffle=True, num_workers=0)
        val_loader = DataLoader(val_dataset, batch_size=args.batch_size, shuffle=False, num_workers=0)

        model = ConvAutoencoderWithAttention(in_channels=1, latent_dim=args.latent_dim, num_classes=2).to(device)
        optimizer = torch.optim.Adam(model.parameters(), lr=args.learning_rate)

        best_val_loss = float("inf")
        best_epoch = -1
        best_state = None
        no_improve = 0

        for epoch in range(1, args.epochs + 1):
            train_loss, train_recon, train_class, train_acc, _ = train_epoch(
                model, train_loader, optimizer, device, alpha=args.alpha, beta=args.beta
            )
            val_loss, val_recon, val_class, val_acc, attn_val = validate(
                model, val_loader, device, alpha=args.alpha, beta=args.beta
            )

            print(
                f"Epoch {epoch:03d} | "
                f"train loss={train_loss:.4f} (recon={train_recon:.4f}, class={train_class:.4f}) acc={train_acc:.1f}% | "
                f"val loss={val_loss:.4f} (recon={val_recon:.4f}, class={val_class:.4f}) acc={val_acc:.1f}%"
            )

            if val_loss < best_val_loss - args.min_delta:
                best_val_loss = val_loss
                best_epoch = epoch
                best_state = {k: v.cpu().clone() for k, v in model.state_dict().items()}
                no_improve = 0
            else:
                no_improve += 1
                if no_improve >= args.patience:
                    print(f"Early stopping at epoch {epoch} (best epoch {best_epoch}, best val loss {best_val_loss:.4f})")
                    break

        # Restore best model
        if best_state is not None:
            model.load_state_dict(best_state)
            model_path = os.path.join(args.output_dir, f"best_model_fold{fold_idx}.pth")
            torch.save(best_state, model_path)
            print(f"Saved best model: {model_path}")

        # Evaluate best model on validation fold
        model.eval()
        y_true, y_prob = [], []
        attn_all = []

        with torch.no_grad():
            for images, labels in val_loader:
                images = images.to(device)
                labels = labels.to(device)

                _, logits, attn = model(images)
                prob = torch.softmax(logits, dim=1)[:, 1]

                y_true.extend(labels.detach().cpu().numpy().tolist())
                y_prob.extend(prob.detach().cpu().numpy().tolist())

                if attn is not None:
                    attn_all.append(attn.detach().cpu().numpy())

        metrics = compute_binary_metrics(y_true, y_prob, threshold=args.threshold)
        metrics.update({
            "fold": fold_idx,
            "best_epoch": best_epoch,
            "best_val_loss": float(best_val_loss),
            "n_val": int(len(y_true)),
        })
        fold_metrics.append(metrics)

        # Attention summary (mean weights per timepoint over val subjects)
        if attn_all:
            attn_all = np.concatenate(attn_all, axis=0)
            attn_mean = attn_all.mean(axis=0).tolist()
            attn_std = attn_all.std(axis=0).tolist()
        else:
            attn_mean, attn_std = [], []
        attention_summary.append({
            "fold": fold_idx,
            "attention_mean": attn_mean,
            "attention_std": attn_std,
        })

        print(
            f"[FOLD {fold_idx}] AUC={metrics['auc']:.4f} | "
            f"acc={metrics['accuracy']:.3f} | sens={metrics['recall_sensitivity']:.3f} | "
            f"spec={metrics['specificity']:.3f} | f1={metrics['f1']:.3f} | "
            f"CM=[[{metrics['tn']},{metrics['fp']}],[{metrics['fn']},{metrics['tp']}]]"
        )

    metrics_df = pd.DataFrame(fold_metrics)
    metrics_path = os.path.join(args.output_dir, "cv_fold_metrics.csv")
    metrics_df.to_csv(metrics_path, index=False)
    print(f"\nSaved fold metrics: {metrics_path}")

    attn_df = pd.DataFrame(attention_summary)
    attn_path = os.path.join(args.output_dir, "cv_attention_summary.csv")
    attn_df.to_csv(attn_path, index=False)
    print(f"Saved attention summary: {attn_path}")

    # Print summary across folds
    valid_auc = metrics_df["auc"].replace([np.inf, -np.inf], np.nan).dropna()
    if len(valid_auc) > 0:
        print(f"\nAUC mean ± SD across folds: {valid_auc.mean():.4f} ± {valid_auc.std(ddof=0):.4f}")
    else:
        print("\nAUC could not be computed across folds (likely due to single-class folds).")


if __name__ == "__main__":
    main()
