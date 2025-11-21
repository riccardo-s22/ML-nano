"""
Convolutional Autoencoder for Single Timepoint Classification

This script adapts the previous multi-timepoint architecture to handle a
single measurement per subject. The model retains the dual-objective design
(reconstruction + classification) but removes the attention-based temporal
aggregation. Each subject now provides one Excel file that is treated as a
2D image for feature learning.
"""

import os
from typing import List, Tuple

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import torch
import torch.nn as nn
import torch.optim as optim
from openpyxl import load_workbook
from sklearn.model_selection import StratifiedKFold
from sklearn.preprocessing import LabelEncoder
from torch.utils.data import DataLoader, Dataset

# Set random seeds for reproducibility
torch.manual_seed(42)
np.random.seed(42)


class ExcelImageDataset(Dataset):
    """Dataset for loading single-timepoint Excel files as image tensors."""

    def __init__(self, codes: List[str], labels: List[int], data_dir: str, transform=None):
        self.codes = codes
        self.labels = labels
        self.data_dir = data_dir
        self.transform = transform

    def __len__(self):
        return len(self.codes)

    def load_excel_as_image(self, filepath: str) -> np.ndarray:
        """Load an Excel file, drop headers/labels, and normalize to [0, 1]."""
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active

        data = []
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if row_idx == 0:
                continue
            row_data = []
            for col_idx, cell in enumerate(row):
                if col_idx == 0:
                    continue
                row_data.append(float(cell) if isinstance(cell, (int, float)) and cell is not None else 0.0)
            if row_data:
                data.append(row_data)

        img_array = np.array(data, dtype=np.float32)
        if img_array.size == 0:
            raise ValueError(f"No valid data found in {filepath}")

        min_val, max_val = img_array.min(), img_array.max()
        if max_val > min_val:
            img_array = (img_array - min_val) / (max_val - min_val)
        else:
            img_array = np.zeros_like(img_array)

        return img_array

    def __getitem__(self, idx):
        code = self.codes[idx]
        label = self.labels[idx]
        filepath = os.path.join(self.data_dir, f"{code}.xlsx")
        img = self.load_excel_as_image(filepath)

        if len(img.shape) == 2:
            img = img[np.newaxis, :, :]

        image_tensor = torch.FloatTensor(img)
        if self.transform:
            image_tensor = self.transform(image_tensor)

        return image_tensor, torch.tensor(label, dtype=torch.long)


class ConvEncoder(nn.Module):
    def __init__(self, in_channels: int, latent_dim: int):
        super().__init__()
        self.conv1 = nn.Sequential(
            nn.Conv2d(in_channels, 32, kernel_size=3, padding=1),
            nn.BatchNorm2d(32),
            nn.ReLU(),
            nn.MaxPool2d(2, 2),
        )
        self.conv2 = nn.Sequential(
            nn.Conv2d(32, 64, kernel_size=3, padding=1),
            nn.BatchNorm2d(64),
            nn.ReLU(),
            nn.MaxPool2d(2, 2),
        )
        self.conv3 = nn.Sequential(
            nn.Conv2d(64, 128, kernel_size=3, padding=1),
            nn.BatchNorm2d(128),
            nn.ReLU(),
            nn.MaxPool2d(2, 2),
        )
        self.conv4 = nn.Sequential(
            nn.Conv2d(128, 256, kernel_size=3, padding=1),
            nn.BatchNorm2d(256),
            nn.ReLU(),
        )

        self.flatten_size = None
        self.fc = None
        self.latent_dim = latent_dim
        self.encoded_shape = None
        self.input_shape = None

    def forward(self, x: torch.Tensor) -> torch.Tensor:
        if self.input_shape is None:
            self.input_shape = (x.shape[2], x.shape[3])

        x = self.conv1(x)
        x = self.conv2(x)
        x = self.conv3(x)
        x = self.conv4(x)

        if self.encoded_shape is None:
            self.encoded_shape = (x.shape[1], x.shape[2], x.shape[3])

        if self.fc is None:
            self.flatten_size = x.shape[1] * x.shape[2] * x.shape[3]
            self.fc = nn.Linear(self.flatten_size, self.latent_dim).to(x.device)

        x = x.view(x.size(0), -1)
        x = self.fc(x)
        return x


class ConvDecoder(nn.Module):
    def __init__(self, latent_dim: int, out_channels: int, encoded_shape: tuple, target_shape: tuple):
        super().__init__()
        self.encoded_channels, self.encoded_h, self.encoded_w = encoded_shape
        self.target_h, self.target_w = target_shape

        self.fc = nn.Linear(latent_dim, self.encoded_channels * self.encoded_h * self.encoded_w)

        self.deconv1 = nn.Sequential(
            nn.ConvTranspose2d(256, 128, kernel_size=3, stride=2, padding=1, output_padding=1),
            nn.BatchNorm2d(128),
            nn.ReLU(),
        )
        self.deconv2 = nn.Sequential(
            nn.ConvTranspose2d(128, 64, kernel_size=3, stride=2, padding=1, output_padding=1),
            nn.BatchNorm2d(64),
            nn.ReLU(),
        )
        self.deconv3 = nn.Sequential(
            nn.ConvTranspose2d(64, 32, kernel_size=3, stride=2, padding=1, output_padding=1),
            nn.BatchNorm2d(32),
            nn.ReLU(),
        )
        self.final_conv = nn.Conv2d(32, out_channels, kernel_size=3, padding=1)
        self.sigmoid = nn.Sigmoid()

    def forward(self, z: torch.Tensor) -> torch.Tensor:
        x = self.fc(z)
        x = x.view(x.size(0), self.encoded_channels, self.encoded_h, self.encoded_w)

        x = self.deconv1(x)
        x = self.deconv2(x)
        x = self.deconv3(x)
        x = self.final_conv(x)

        current_h, current_w = x.shape[2], x.shape[3]
        if current_h != self.target_h or current_w != self.target_w:
            x = torch.nn.functional.interpolate(
                x, size=(self.target_h, self.target_w), mode="bilinear", align_corners=False
            )

        return self.sigmoid(x)


class Classifier(nn.Module):
    def __init__(self, latent_dim: int, num_classes: int, dropout_rate: float = 0.3):
        super().__init__()
        self.classifier = nn.Sequential(
            nn.Linear(latent_dim, 128),
            nn.ReLU(),
            nn.Dropout(dropout_rate),
            nn.Linear(128, num_classes),
        )

    def forward(self, z: torch.Tensor) -> torch.Tensor:
        return self.classifier(z)


class ConvAutoencoder(nn.Module):
    """Dual-objective autoencoder for single timepoint classification."""

    def __init__(self, in_channels: int, latent_dim: int, num_classes: int):
        super().__init__()
        self.encoder = ConvEncoder(in_channels, latent_dim)
        self.decoder = None
        self.classifier = Classifier(latent_dim, num_classes)
        self.in_channels = in_channels
        self.latent_dim = latent_dim

    def forward(self, images: torch.Tensor) -> Tuple[torch.Tensor, torch.Tensor]:
        z = self.encoder(images)

        if self.decoder is None and self.encoder.encoded_shape is not None:
            self.decoder = ConvDecoder(
                self.latent_dim, self.in_channels, self.encoder.encoded_shape, self.encoder.input_shape
            ).to(images.device)

        reconstructions = self.decoder(z)
        class_logits = self.classifier(z)
        return reconstructions, class_logits


def train_epoch(model, dataloader, optimizer, device, alpha: float = 1.0, beta: float = 1.0):
    model.train()
    total_loss = total_recon_loss = total_class_loss = 0.0
    correct = total = 0

    reconstruction_criterion = nn.MSELoss()
    classification_criterion = nn.CrossEntropyLoss()

    for images, labels in dataloader:
        images = images.to(device)
        labels = labels.to(device)

        optimizer.zero_grad()
        reconstructions, class_logits = model(images)

        recon_loss = reconstruction_criterion(reconstructions, images)
        class_loss = classification_criterion(class_logits, labels)
        loss = alpha * recon_loss + beta * class_loss

        loss.backward()
        optimizer.step()

        total_loss += loss.item()
        total_recon_loss += recon_loss.item()
        total_class_loss += class_loss.item()

        _, predicted = torch.max(class_logits, 1)
        total += labels.size(0)
        correct += (predicted == labels).sum().item()

    avg_loss = total_loss / len(dataloader)
    avg_recon_loss = total_recon_loss / len(dataloader)
    avg_class_loss = total_class_loss / len(dataloader)
    accuracy = 100 * correct / total
    return avg_loss, avg_recon_loss, avg_class_loss, accuracy


def validate(model, dataloader, device, alpha: float = 1.0, beta: float = 1.0):
    model.eval()
    total_loss = total_recon_loss = total_class_loss = 0.0
    correct = total = 0

    reconstruction_criterion = nn.MSELoss()
    classification_criterion = nn.CrossEntropyLoss()

    with torch.no_grad():
        for images, labels in dataloader:
            images = images.to(device)
            labels = labels.to(device)

            reconstructions, class_logits = model(images)
            recon_loss = reconstruction_criterion(reconstructions, images)
            class_loss = classification_criterion(class_logits, labels)
            loss = alpha * recon_loss + beta * class_loss

            total_loss += loss.item()
            total_recon_loss += recon_loss.item()
            total_class_loss += class_loss.item()

            _, predicted = torch.max(class_logits, 1)
            total += labels.size(0)
            correct += (predicted == labels).sum().item()

    avg_loss = total_loss / len(dataloader)
    avg_recon_loss = total_recon_loss / len(dataloader)
    avg_class_loss = total_class_loss / len(dataloader)
    accuracy = 100 * correct / total
    return avg_loss, avg_recon_loss, avg_class_loss, accuracy


def plot_history(train_history, val_history, fold: int):
    epochs = range(1, len(train_history["loss"]) + 1)
    plt.figure(figsize=(12, 5))

    plt.subplot(1, 2, 1)
    plt.plot(epochs, train_history["loss"], label="Train Loss")
    plt.plot(epochs, val_history["loss"], label="Val Loss")
    plt.xlabel("Epoch")
    plt.ylabel("Loss")
    plt.title(f"Fold {fold} - Total Loss")
    plt.legend()

    plt.subplot(1, 2, 2)
    plt.plot(epochs, train_history["accuracy"], label="Train Acc")
    plt.plot(epochs, val_history["accuracy"], label="Val Acc")
    plt.xlabel("Epoch")
    plt.ylabel("Accuracy (%)")
    plt.title(f"Fold {fold} - Accuracy")
    plt.legend()

    plt.tight_layout()
    plt.savefig(f"training_curves_fold{fold}.png")
    plt.close()


def main():
    LABELS_CSV = "sample_labels.csv"
    DATA_DIR = "out_0h"
    LATENT_DIM = 512
    BATCH_SIZE = 4
    NUM_EPOCHS = 100
    LEARNING_RATE = 0.001
    ALPHA = 1.0
    BETA = 1.0
    N_SPLITS = 5

    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    print(f"Using device: {device}")

    labels_df = pd.read_csv(LABELS_CSV)
    print(f"Loaded {len(labels_df)} samples")
    print(f"Groups: {labels_df['group'].unique()}")

    label_encoder = LabelEncoder()
    labels_df["encoded_group"] = label_encoder.fit_transform(labels_df["group"])
    num_classes = len(label_encoder.classes_)
    print(f"Number of classes: {num_classes}")

    codes = labels_df["code"].astype(str).values
    labels = labels_df["encoded_group"].values

    skf = StratifiedKFold(n_splits=N_SPLITS, shuffle=True, random_state=42)
    fold_results = []

    print("\n" + "=" * 70)
    print("STARTING SUBJECT-LEVEL CROSS-VALIDATION")
    print(f"Total subjects: {len(codes)}")
    print("Timepoints per subject: 1")
    print(f"Total images: {len(codes)}")
    print("=" * 70)

    for fold, (train_idx, val_idx) in enumerate(skf.split(codes, labels)):
        print(f"\n{'=' * 70}")
        print(f"FOLD {fold + 1}/{N_SPLITS}")
        print("=" * 70)

        train_codes, train_labels = codes[train_idx], labels[train_idx]
        val_codes, val_labels = codes[val_idx], labels[val_idx]

        print(f"Training subjects: {len(train_codes)}")
        print(f"Validation subjects: {len(val_codes)}")
        print(f"Training images: {len(train_codes)}")
        print(f"Validation images: {len(val_codes)}")

        train_dataset = ExcelImageDataset(train_codes, train_labels, DATA_DIR)
        val_dataset = ExcelImageDataset(val_codes, val_labels, DATA_DIR)

        train_loader = DataLoader(train_dataset, batch_size=BATCH_SIZE, shuffle=True)
        val_loader = DataLoader(val_dataset, batch_size=BATCH_SIZE, shuffle=False)

        sample_images, _ = next(iter(train_loader))
        in_channels, height, width = sample_images.shape[1:]
        print("\nData dimensions:")
        print(f"  Channels: {in_channels}")
        print(f"  Height: {height}")
        print(f"  Width: {width}")

        model = ConvAutoencoder(in_channels=in_channels, latent_dim=LATENT_DIM, num_classes=num_classes).to(device)

        total_params = sum(p.numel() for p in model.parameters())
        trainable_params = sum(p.numel() for p in model.parameters() if p.requires_grad)
        print("\nModel parameters:")
        print(f"  Total: {total_params:,}")
        print(f"  Trainable: {trainable_params:,}")

        optimizer = optim.Adam(model.parameters(), lr=LEARNING_RATE)

        best_val_acc = 0
        train_history = {"loss": [], "recon_loss": [], "class_loss": [], "accuracy": []}
        val_history = {"loss": [], "recon_loss": [], "class_loss": [], "accuracy": []}

        print(f"\nStarting training for {NUM_EPOCHS} epochs...")
        print(f"Loss weights: α={ALPHA} (reconstruction), β={BETA} (classification)")

        for epoch in range(NUM_EPOCHS):
            train_loss, train_recon, train_class, train_acc = train_epoch(
                model, train_loader, optimizer, device, ALPHA, BETA
            )
            val_loss, val_recon, val_class, val_acc = validate(
                model, val_loader, device, ALPHA, BETA
            )

            train_history["loss"].append(train_loss)
            train_history["recon_loss"].append(train_recon)
            train_history["class_loss"].append(train_class)
            train_history["accuracy"].append(train_acc)

            val_history["loss"].append(val_loss)
            val_history["recon_loss"].append(val_recon)
            val_history["class_loss"].append(val_class)
            val_history["accuracy"].append(val_acc)

            if (epoch + 1) % 10 == 0:
                print(f"\nEpoch [{epoch + 1}/{NUM_EPOCHS}]")
                print(
                    f"  Train - Loss: {train_loss:.4f}, Recon: {train_recon:.4f}, "
                    f"Class: {train_class:.4f}, Acc: {train_acc:.2f}%"
                )
                print(
                    f"  Val   - Loss: {val_loss:.4f}, Recon: {val_recon:.4f}, "
                    f"Class: {val_class:.4f}, Acc: {val_acc:.2f}%"
                )

            if val_acc > best_val_acc:
                best_val_acc = val_acc
                torch.save(model.state_dict(), f"best_model_fold{fold + 1}.pth")
                print(f"  ✓ New best model saved (val acc: {val_acc:.2f}%)")

        fold_results.append(
            {
                "fold": fold + 1,
                "best_val_accuracy": best_val_acc,
                "train_history": train_history,
                "val_history": val_history,
            }
        )

        plot_history(train_history, val_history, fold + 1)

    all_accuracies = [r["best_val_accuracy"] for r in fold_results]
    avg_acc = np.mean(all_accuracies)
    std_acc = np.std(all_accuracies)

    print("\nOverall Performance:")
    print(f"  Average Accuracy: {avg_acc:.2f}% ± {std_acc:.2f}%")
    print("\nIndividual Fold Results:")
    for r in fold_results:
        print(f"  Fold {r['fold']}: {r['best_val_accuracy']:.2f}%")

    print("\n" + "=" * 70)
    print("TRAINING COMPLETE")
    print("=" * 70)


if __name__ == "__main__":
    main()
